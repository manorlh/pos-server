"""
"סימוני תזונה" and the product description (docs/SPEC_PRODUCT_DIETARY.md): the tags are
validated, stored in one order without repeats, kept free of contradictions, edited from
the product form and the Excel sheet, and reach the till on both catalog serialisers —
as does the description, and the allergens beside them.
"""
from __future__ import annotations

import io
import uuid
from types import SimpleNamespace

import openpyxl
import pytest
from pydantic import ValidationError

from app.models.product import Product
from app.schemas.product import ProductCreate, ProductResponse, ProductUpdate
from app.services import catalog_sheet as S
from app.services import dietary
from app.services.sync import _serialize_merged_product, _serialize_product
from test_catalog_import import commit, preview, products_named, template, w, workbook  # noqa: F401
from test_product_weighed import _product


# ── The rules ─────────────────────────────────────────────────────────────────


class TestClean:
    def test_known_codes_in_the_fixed_order_each_once(self) -> None:
        assert dietary.clean(["spicy", "gluten_free", "spicy", "meat"]) == ["meat", "gluten_free", "spicy"]

    def test_order_does_not_count(self) -> None:
        assert dietary.clean(["dairy", "spicy"]) == dietary.clean(["spicy", "dairy"]) == ["dairy", "spicy"]

    def test_case_and_spaces_are_read_as_the_code(self) -> None:
        assert dietary.clean([" Vegetarian ", "GLUTEN_FREE"]) == ["vegetarian", "gluten_free"]

    def test_none_and_empty_are_no_tags(self) -> None:
        assert dietary.clean(None) == [] and dietary.clean([]) == [] and dietary.clean(["", " "]) == []
        assert dietary.to_column([]) is None

    def test_an_unknown_code_is_refused(self) -> None:
        with pytest.raises(dietary.DietaryTagError, match="unknown dietary tag: kosher"):
            dietary.clean(["vegan", "kosher"])

    def test_not_a_list_of_text_is_refused(self) -> None:
        with pytest.raises(dietary.DietaryTagError):
            dietary.clean([1, 2])
        with pytest.raises(dietary.DietaryTagError):
            dietary.clean({"vegan": True})

    def test_vegan_is_also_vegetarian(self) -> None:
        assert dietary.clean(["vegan"]) == ["vegan", "vegetarian"]
        assert dietary.clean(["spicy", "vegan", "vegetarian"]) == ["vegan", "vegetarian", "spicy"]

    @pytest.mark.parametrize(
        "tags",
        [["meat", "dairy"], ["vegan", "meat"], ["vegan", "dairy"], ["vegetarian", "meat"]],
    )
    def test_contradictions_are_refused(self, tags) -> None:
        with pytest.raises(dietary.DietaryTagError, match="cannot go together"):
            dietary.clean(tags)

    def test_dairy_vegetarian_and_the_rest_go_together(self) -> None:
        assert dietary.clean(["dairy", "vegetarian", "gluten_free", "spicy"]) == [
            "vegetarian", "dairy", "gluten_free", "spicy",
        ]

    def test_reading_back_never_raises(self) -> None:
        # A row stored before a rule (or by hand): its known codes, in order.
        assert dietary.tags_out(["spicy", "kosher", "VEGAN"]) == ["vegan", "spicy"]
        assert dietary.tags_out(None) == [] and dietary.tags_out("vegan") == [] and dietary.tags_out(7) == []

    def test_hebrew_labels_and_aliases(self) -> None:
        assert dietary.labels(["dairy", "vegetarian"]) == ["צמחוני", "חלבי"]
        assert [dietary.code_of(w) for w in ("טבעוני", "ללא גלוטן", "GF", "חלבית", "meat", "כשר")] == [
            "vegan", "gluten_free", "gluten_free", "dairy", "meat", None,
        ]
        assert set(dietary.DIETARY_LABELS_HE) == set(dietary.DIETARY_TAGS)


# ── The API's schemas ─────────────────────────────────────────────────────────

_BASE = {"name": "סלט", "price": "42.00", "sku": "S-1", "categoryId": str(uuid.uuid4())}


class TestSchemas:
    def test_create_cleans_the_tags(self) -> None:
        body = ProductCreate.model_validate({**_BASE, "dietaryTags": ["spicy", "vegan", "spicy"]})
        assert body.dietary_tags == ["vegan", "vegetarian", "spicy"]
        assert ProductCreate.model_validate(_BASE).dietary_tags is None

    def test_create_refuses_unknown_and_contradicting_tags(self) -> None:
        with pytest.raises(ValidationError):
            ProductCreate.model_validate({**_BASE, "dietaryTags": ["halal"]})
        with pytest.raises(ValidationError):
            ProductCreate.model_validate({**_BASE, "dietaryTags": ["meat", "dairy"]})

    def test_update_leaves_them_alone_unless_sent(self) -> None:
        assert "dietary_tags" not in ProductUpdate.model_validate({"price": "7"}).model_dump(exclude_unset=True)
        cleared = ProductUpdate.model_validate({"dietaryTags": []}).model_dump(exclude_unset=True)
        assert cleared == {"dietary_tags": []}
        assert ProductUpdate.model_validate({"dietaryTags": None}).model_dump(exclude_unset=True) == {"dietary_tags": None}

    def test_description_up_to_1000(self) -> None:
        assert ProductCreate.model_validate({**_BASE, "description": "א" * 1000}).description == "א" * 1000
        with pytest.raises(ValidationError):
            ProductCreate.model_validate({**_BASE, "description": "א" * 1001})
        with pytest.raises(ValidationError):
            ProductUpdate.model_validate({"description": "א" * 1001})

    def test_the_response_always_has_a_list(self) -> None:
        def row(tags):
            from datetime import datetime, timezone

            now = datetime.now(timezone.utc)
            return SimpleNamespace(
                id=uuid.uuid4(), tenant_id=None, company_id=None, shop_id=None, pos_machine_id=None,
                global_product_id=None, catalog_level="global", is_local_override=False, name="x",
                description="תיאור", price=1, sku="1", global_sku=None, sku_auto_assigned=False,
                category_id=uuid.uuid4(), image_url=None, in_stock=True, is_available=True,
                stock_quantity=0, barcode=None, tax_rate=None, voucher_id=None, dietary_tags=tags,
                alerts=None, companions=None, allergen_alert=None, allergen_alert_require_ack=None,
                created_at=now, updated_at=now,
            )

        assert ProductResponse.model_validate(row(None)).model_dump(by_alias=True)["dietaryTags"] == []
        out = ProductResponse.model_validate(row(["spicy", "dairy"])).model_dump(by_alias=True)
        assert out["dietaryTags"] == ["dairy", "spicy"] and out["description"] == "תיאור"


# ── The till's catalog ────────────────────────────────────────────────────────


class TestSyncPayload:
    def test_both_serialisers_ship_tags_description_and_allergens(self) -> None:
        p = _product(is_weighed=False, unit_label=None)
        p.dietary_tags = ["spicy", "vegan", "vegetarian"]
        p.description = "חומוס, טחינה ושמן זית"
        p.allergens = ["sesame"]
        for row in (_serialize_product(p), _serialize_merged_product(p, None, None, uuid.uuid4(), None)):
            assert row["dietaryTags"] == ["vegan", "vegetarian", "spicy"]
            assert row["description"] == "חומוס, טחינה ושמן זית"
            assert row["allergens"] == ["sesame"]

    def test_none_is_an_empty_list(self) -> None:
        p = _product(is_weighed=False, unit_label=None)
        p.dietary_tags = None
        assert _serialize_product(p)["dietaryTags"] == []
        assert _serialize_merged_product(p, None, None, uuid.uuid4(), None)["dietaryTags"] == []

    def test_they_are_taken_from_the_global_row(self) -> None:
        global_p = _product(is_weighed=False, unit_label=None)
        global_p.dietary_tags = ["dairy"]
        global_p.description = "של הרשת"
        local = _product(is_weighed=False, unit_label=None)
        local.dietary_tags = ["meat"]
        local.description = "מקומי"
        row = _serialize_merged_product(global_p, local, None, uuid.uuid4(), None)
        assert row["dietaryTags"] == ["dairy"] and row["description"] == "של הרשת"


# ── Round trip through the products router ────────────────────────────────────


def test_the_router_stores_edits_and_ships_them(monkeypatch) -> None:
    from shift_world import accept_str_uuids, make_world

    from app.models.category import Category
    from app.routers import products as R

    accept_str_uuids(monkeypatch)
    world = make_world()
    monkeypatch.setattr(R, "_trigger_catalog_notify", lambda *a, **k: None)
    monkeypatch.setattr(R, "allocate_global_sku", lambda *a, **k: "100001")
    monkeypatch.setattr(R, "resolve_sku_for_create", lambda db, tenant, sku: (sku, False))
    cat = Category(id=uuid.uuid4(), tenant_id=world.tenant.id, company_id=world.company.id, name="סלטים")
    world.db.add(cat)
    world.db.commit()

    created = R.create_product(
        ProductCreate.model_validate({
            "name": "סלט ירוק", "price": "39.00", "sku": "SL-1", "categoryId": str(cat.id),
            "companyId": str(world.company.id), "description": "חסה, מלפפון, עשבי תיבול",
            "dietaryTags": ["gluten_free", "vegan"],
        }),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    stored = world.db.get(Product, created.id)
    assert stored.dietary_tags == ["vegan", "vegetarian", "gluten_free"]
    assert ProductResponse.model_validate(stored).model_dump(by_alias=True)["dietaryTags"] == [
        "vegan", "vegetarian", "gluten_free",
    ]

    # Omitted: left alone.
    R.update_product(
        str(created.id), ProductUpdate.model_validate({"price": "41.00"}),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    assert world.db.get(Product, created.id).dietary_tags == ["vegan", "vegetarian", "gluten_free"]

    # Changed, then the payload a till pulls.
    R.update_product(
        str(created.id), ProductUpdate.model_validate({"dietaryTags": ["dairy", "spicy"], "description": "עם פטה"}),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    row = _serialize_product(world.db.get(Product, created.id))
    assert row["dietaryTags"] == ["dairy", "spicy"] and row["description"] == "עם פטה"

    # Cleared: stored as NULL, shipped as [].
    R.update_product(
        str(created.id), ProductUpdate.model_validate({"dietaryTags": []}),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    assert world.db.get(Product, created.id).dietary_tags is None
    assert _serialize_product(world.db.get(Product, created.id))["dietaryTags"] == []


# ── The Excel sheet ───────────────────────────────────────────────────────────


class TestSheetCell:
    def test_hebrew_codes_and_a_mix(self) -> None:
        assert S.parse_dietary("חריף, טבעוני").value == ("vegan", "vegetarian", "spicy")
        assert S.parse_dietary("dairy;gluten_free").value == ("dairy", "gluten_free")
        assert S.parse_dietary("בשרי, spicy, בשרי").value == ("meat", "spicy")

    def test_empty_is_no_change_and_none_clears(self) -> None:
        assert S.parse_dietary(None).value is None and S.parse_dietary("  ").value is None
        assert S.parse_dietary("ללא").value == ()

    def test_unknown_and_contradicting_are_row_errors(self) -> None:
        assert "סימון תזונה לא מוכר ('כשר')" in S.parse_dietary("טבעוני, כשר").error
        assert S.parse_dietary("בשרי, חלבי").error == "סימוני תזונה סותרים: בשרי וחלבי"
        assert S.parse_dietary("טבעוני, חלבי").error is not None


class TestImportExport:
    def test_a_new_product_gets_its_tags_and_description(self, w) -> None:
        commit(w, workbook([{
            "name": "שקשוקה", "category": "אוכל", "price": 48, "dietary": "צמחוני, חריף",
            "description": "ביצים ברוטב עגבניות",
        }]))
        made = products_named(w, "שקשוקה")[0]
        assert made.dietary_tags == ["vegetarian", "spicy"] and made.description == "ביצים ברוטב עגבניות"

    def test_an_update_shows_the_change_in_hebrew_and_ללא_clears(self, w) -> None:
        out = preview(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "dietary": "vegan"}]))
        change = next(c for c in out["products"][0]["changes"] if c["field"] == "dietary")
        assert change["label"] == "סימוני תזונה" and change["after"] == "טבעוני, צמחוני"
        commit(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "dietary": "vegan"}]))
        w.db.refresh(w.cola)
        assert w.cola.dietary_tags == ["vegan", "vegetarian"]
        # Empty: no change.
        commit(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12}]))
        w.db.refresh(w.cola)
        assert w.cola.dietary_tags == ["vegan", "vegetarian"]
        commit(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "dietary": "ללא"}]))
        w.db.refresh(w.cola)
        assert w.cola.dietary_tags is None

    def test_a_contradiction_is_an_error_row(self, w) -> None:
        out = preview(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "dietary": "בשרי, חלבי"}]))
        assert out["products"][0]["status"] == "error"

    def test_the_export_writes_hebrew_and_reads_back_as_no_change(self, w) -> None:
        w.cola.dietary_tags = ["vegan", "vegetarian", "gluten_free"]
        w.db.commit()
        data = template(w, with_data=True)
        ws = openpyxl.load_workbook(io.BytesIO(data))[S.SHEET_PRODUCTS]
        col = [c.key for c in S.PRODUCT_COLUMNS].index("dietary") + 1
        assert ws.cell(1, col).value == "סימוני תזונה"
        cells = {ws.cell(r, 2).value: ws.cell(r, col).value for r in range(2, 8) if ws.cell(r, 2).value}
        assert cells["קולה"] == "טבעוני, צמחוני, ללא גלוטן"
        out = preview(w, data)
        assert out["summary"]["productsUnchanged"] == 3 and out["summary"]["errors"] == 0


# ── The menu-broadcast review ─────────────────────────────────────────────────


def test_the_broadcast_review_names_a_tag_change_in_hebrew_but_not_an_older_snapshot() -> None:
    from app.services import menu_broadcast as MB

    def snap(**fields):
        return {"products": {"p1": {"name": "סלט", "price": 30, "shopListed": True, **fields}}}

    changed = MB.diff(snap(dietaryTags=[]), snap(dietaryTags=["vegan", "vegetarian"]))["products"]
    assert changed[0]["changes"] == [{"field": "dietaryTags", "before": [], "after": ["טבעוני", "צמחוני"]}]
    # A publication made before the field existed carries none: not a change.
    assert MB.diff(snap(), snap(dietaryTags=[]))["products"] == []
    assert MB.diff(snap(dietaryTags=["dairy"]), snap(dietaryTags=["dairy"]))["products"] == []


# ── The till parameter ────────────────────────────────────────────────────────


def test_the_till_parameter_is_built_in_and_off() -> None:
    from app.services.till_parameters import BUILTIN_PARAMETERS

    spec = next(p for p in BUILTIN_PARAMETERS if p.key == "showDietaryMarks")
    assert spec.label == "הצג סימוני תזונה בקופה"
    assert spec.value_type == "boolean" and spec.default_value is False
