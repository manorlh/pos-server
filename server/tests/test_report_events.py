"""
Temporary events ("אירועים", docs/SPEC_EVENTS.md): tills grouped at report level only.

* the overlap rule (409 naming the other event), the window (date + hour, both required);
* the window filter, the KPIs on a small fixture;
* each insight rule: weak till, idle till, high tip, refunds, exceptions, open shift;
* confirm: the snapshot is frozen, the tills are released, the event is read-only, and a
  later sale does not change the confirmed report;
* the reconciliation against the Z (match, mismatch, missing Z) and against the card
  transmissions (match, mismatch, failed, not yet transmitted);
* the Excel export has the expected sheets.

Runs on the in-memory SQLite world of tests/shift_world.py.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.models.audit_exception import AuditException
from app.models.card_transmission import CardTransmission, CardTransmissionItem
from app.models.pos_machine import PairingStatus, POSMachine
from app.models.report_event import ReportEvent, ReportEventMachine
from app.models.shift import ShiftStatus
from app.models.transaction import Transaction
from app.models.transaction_payment import TransactionPayment
from app.models.z_report import ZOrigin, ZReport
from app.routers import report_events as R
from app.schemas.report_event import ReportEventConfirm, ReportEventCreate, ReportEventUpdate
from app.services.report_events import crud as C
from app.services.report_events.report import build_report, report_for
from app.services.report_events.rules import weighted_median
from app.services.shift_totals import compute_totals
from app.services.z_builder import machine_section
from shift_world import accept_str_uuids, make_world

# 2026-09-27 is IDT (UTC+3): 18:00 local = 15:00 UTC.
DAY = "2026-09-27"
START = datetime(2026, 9, 27, 15, 0, tzinfo=timezone.utc)   # 18:00 local
END = datetime(2026, 9, 27, 20, 0, tzinfo=timezone.utc)     # 23:00 local
AFTER = datetime(2026, 9, 28, 6, 0, tzinfo=timezone.utc)


@pytest.fixture
def w(monkeypatch):
    accept_str_uuids(monkeypatch)
    return make_world()


def at(minutes: float) -> datetime:
    """Minutes after the event's start (18:00 local)."""
    return START + timedelta(minutes=minutes)


def sale(w, till, minutes, total, *, shift=None, method="cash", tip="0", discount="0", credit_note=False, legs=None):
    tx = w.doc(till, shift, total, method=method, tip=tip, discount=discount, credit_note=credit_note, legs=legs)
    tx.created_at = at(minutes)
    w.db.flush()
    return tx


def add_till(w, name):
    m = POSMachine(
        id=uuid.uuid4(), tenant_id=w.tenant.id, shop_id=w.shop.id, distributor_id=w.admin.id, name=name,
        machine_code=f"M-{name}", pos_number=str(len(w.tills) + 10), is_active=True,
        pairing_status=PairingStatus.ASSIGNED, last_heartbeat_at=AFTER,
    )
    w.db.add(m)
    w.db.flush()
    w.tills.append(m)
    return m


def create(w, name="במה ראשית", tills=None, start=("18:00", "23:00"), dates=(DAY, DAY), **extra):
    body = ReportEventCreate(
        shopId=w.shop.id, name=name, startDate=dates[0], startTime=start[0], endDate=dates[1], endTime=start[1],
        machineIds=[t.id for t in (tills if tills is not None else w.tills)], **extra,
    )
    out = R.create_report_event(body, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    return w.db.query(ReportEvent).filter(ReportEvent.id == uuid.UUID(out["id"])).one()


def report(w, event, now=AFTER):
    w.db.expire_all()
    return build_report(w.db, event, now=now)


def codes(rep, level=None):
    return [i["code"] for i in rep["insights"] if level is None or i["level"] == level]


def till_row(rep, till):
    return next(t for t in rep["tills"] if t["machineId"] == str(till.id))


# ── The event: window, overlap ────────────────────────────────────────────────


def test_the_window_is_a_local_date_and_hour_and_both_are_required(w):
    event = create(w)
    assert event.starts_at.replace(tzinfo=timezone.utc) == START
    assert event.ends_at.replace(tzinfo=timezone.utc) == END
    out = R.get_report_event(event.id, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert (out["startDate"], out["startTime"], out["endDate"], out["endTime"]) == (DAY, "18:00", DAY, "23:00")

    with pytest.raises(HTTPException) as e:
        create(w, name="x", start=("18:00", ""), tills=[])
    assert e.value.status_code == 422 and e.value.detail["code"] == "invalid_window"
    with pytest.raises(HTTPException) as e:
        create(w, name="x", start=("23:00", "18:00"), tills=[])
    assert e.value.detail["code"] == "invalid_window"
    with pytest.raises(HTTPException) as e:
        create(w, name="x", dates=(DAY, "2026-10-15"), tills=[])
    assert e.value.detail["code"] == "invalid_window"


def test_a_till_cannot_be_in_two_events_whose_windows_overlap(w):
    t1, t2 = w.tills
    first = create(w, name="במה א", tills=[t1])
    with pytest.raises(HTTPException) as e:
        create(w, name="במה ב", tills=[t1, t2], start=("22:00", "23:30"))
    assert e.value.status_code == 409
    assert e.value.detail["code"] == "till_in_overlapping_event"
    assert e.value.detail["eventId"] == str(first.id)
    assert e.value.detail["eventName"] == "במה א"
    assert "במה א" in e.value.detail["message"]
    # Back to back is not an overlap; another till is free.
    create(w, name="אחרי", tills=[t1], start=("23:00", "23:59"))
    create(w, name="במה ב", tills=[t2], start=("22:00", "23:30"))
    # The till list marks the busy till with the event it is in.
    view = R.get_report_event_tills(
        shop_id=w.shop.id, start_date=DAY, start_time="19:00", end_date=DAY, end_time="20:00",
        exclude_event_id=None, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db,
    )
    busy = {t["name"]: t["busy"] for t in view["tills"]}
    assert busy["Till 1"]["eventName"] == "במה א"
    # Editing an event onto a busy till is refused too.
    with pytest.raises(HTTPException) as e:
        R.update_report_event(first.id, ReportEventUpdate(machineIds=[t1.id, t2.id]),
                              current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert e.value.status_code == 409


def test_only_the_shops_own_tills_may_be_assigned(w):
    with pytest.raises(HTTPException) as e:
        create(w, tills=[w.other_till])
    assert e.value.status_code == 422 and e.value.detail["code"] == "till_not_in_shop"


# ── The report ────────────────────────────────────────────────────────────────


def test_only_documents_inside_the_window_count(w):
    t1, _t2 = w.tills
    event = create(w)
    sale(w, t1, -1, "11.00")      # 17:59 — before
    sale(w, t1, 0, "20.00")       # 18:00 — in
    sale(w, t1, 299, "30.00")     # 22:59 — in
    sale(w, t1, 300, "40.00")     # 23:00 — after (the end is exclusive)
    sale(w, w.other_till, 60, "99.00")   # another shop's till
    rep = report(w, event)
    assert rep["kpis"]["salesCount"] == 2
    assert rep["kpis"]["net"] == 50.0


def test_the_kpis_on_a_small_fixture(w):
    t1, t2 = w.tills
    event = create(w)
    sale(w, t1, 10, "100.00", method="cash")
    sale(w, t1, 20, "50.00", discount="10.00", method="card", tip="5.00")
    sale(w, t1, 30, "20.00", method="cash", credit_note=True)
    sale(w, t2, 40, "60.00", method="card", tip="30.00")
    k = report(w, event)["kpis"]
    assert k["gross"] == 210.0
    assert k["discounts"] == 10.0
    assert k["sales"] == 200.0
    assert k["refunds"] == 20.0 and k["refundsCount"] == 1
    assert k["net"] == 180.0
    assert k["salesCount"] == 3
    assert k["avgTicket"] == 66.67
    assert k["tips"] == 35.0 and k["tipPct"] == 17.5
    assert (k["cash"], k["card"], k["other"]) == (80.0, 100.0, 0.0)
    assert k["peakHour"]["net"] == 180.0
    assert k["activeTills"] == 2


def test_weighted_median_is_the_midpoint_on_an_even_split():
    assert weighted_median([(10, 1), (30, 1)]) == 20
    assert weighted_median([(10, 1), (30, 3)]) == 30
    assert weighted_median([]) is None


def test_a_weak_till_is_flagged_against_the_median(w):
    t1, t2 = w.tills
    t3 = add_till(w, "Till 3")
    event = create(w)
    for minute in range(0, 180, 10):
        sale(w, t1, minute, "100.00")
        sale(w, t3, minute + 2, "90.00")
        if minute % 60 == 0:
            sale(w, t2, minute + 5, "20.00")
    sale(w, t2, 175, "20.00")
    rep = report(w, event)
    assert till_row(rep, t2)["weak"] is True
    assert till_row(rep, t1)["weak"] is False
    weak = [i for i in rep["insights"] if i["code"] == "weak_till"]
    assert len(weak) == 1 and weak[0]["ref"]["id"] == str(t2.id)


def test_an_idle_gap_is_flagged(w):
    t1, t2 = w.tills
    event = create(w)
    sale(w, t1, 0, "10.00")
    sale(w, t1, 10, "10.00")
    sale(w, t1, 100, "10.00")     # 90 minutes of nothing
    for minute in range(0, 110, 15):
        sale(w, t2, minute, "10.00")
    rep = report(w, event)
    gaps = till_row(rep, t1)["idleGaps"]
    assert [g["kind"] for g in gaps] == ["gap"] and gaps[0]["minutes"] == 90
    assert till_row(rep, t2)["idle"] is False
    assert "idle_till" in codes(rep, "warning")


def test_a_till_with_no_sales_is_an_alert(w):
    t1, _t2 = w.tills
    event = create(w)
    sale(w, t1, 30, "10.00")
    assert "no_sales_till" in codes(report(w, event), "alert")


def test_a_high_tip_is_flagged_with_its_document(w):
    t1, _t2 = w.tills
    event = create(w, thresholds={"highTipPct": 20})
    sale(w, t1, 10, "100.00", tip="10.00")   # 10% — fine
    big = sale(w, t1, 20, "100.00", tip="30.00")
    rep = report(w, event)
    tips = [i for i in rep["insights"] if i["code"] == "high_tip"]
    assert len(tips) == 1 and tips[0]["ref"] == {"kind": "document", "id": str(big.id), "name": big.transaction_number}
    # An absolute threshold catches a tip under the percentage too.
    R.update_report_event(event.id, ReportEventUpdate(thresholds={"highTipAmount": 8}),
                          current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    rep = report(w, event)
    assert len([i for i in rep["insights"] if i["code"] == "high_tip"]) == 2


def test_refunds_are_reported_and_an_outlier_till_flagged(w):
    t1, t2 = w.tills
    event = create(w)
    for minute in range(0, 100, 10):
        sale(w, t1, minute, "50.00")
        sale(w, t2, minute + 1, "50.00")
    sale(w, t2, 30, "40.00", credit_note=True)
    sale(w, t2, 40, "40.00", credit_note=True)
    sale(w, t1, 50, "5.00", credit_note=True)
    rep = report(w, event)
    assert rep["kpis"]["refundsCount"] == 3
    assert "refunds" in codes(rep)
    outliers = [i for i in rep["insights"] if i["code"] == "refund_outlier"]
    assert any(i["ref"]["id"] == str(t2.id) for i in outliers)


def test_exceptions_are_counted_per_till_and_reported(w):
    t1, _t2 = w.tills
    event = create(w)
    sale(w, t1, 5, "10.00")
    for n in range(3):
        w.db.add(AuditException(
            id=uuid.uuid4(), tenant_id=w.tenant.id, shop_id=w.shop.id, machine_id=t1.id,
            exception_type="drawer_open", severity="medium", dedupe_key=f"drawer_open:{n}",
            occurred_at=at(30 + n),
        ))
    w.db.add(AuditException(
        id=uuid.uuid4(), tenant_id=w.tenant.id, shop_id=w.shop.id, machine_id=t1.id,
        exception_type="line_void", severity="low", dedupe_key="line_void:x", occurred_at=at(400),
    ))  # after the event
    w.db.flush()
    rep = report(w, event)
    assert rep["kpis"]["exceptionsCount"] == 3
    assert till_row(rep, t1)["exceptions"] == 3
    assert "exception_drawer_open" in codes(rep, "warning")


def test_an_open_shift_is_flagged_and_warned_before_confirming(w):
    t1, _t2 = w.tills
    event = create(w)
    shift = w.shift(t1, 1, status=ShiftStatus.OPEN, opened_at=at(-30))
    sale(w, t1, 15, "25.00", shift=shift)
    rep = report(w, event)
    assert "open_shift" in codes(rep, "alert")       # the event has ended and it is still open
    assert any(x["code"] == "open_shift" for x in rep["readiness"]["warnings"])
    shift_row = rep["shifts"][0]
    assert shift_row["status"] == "open" and shift_row["documentsInWindow"] == 1


# ── Confirmation ("נותן תוקף") ────────────────────────────────────────────────


def test_confirm_freezes_the_report_and_releases_the_tills(w, monkeypatch):
    t1, t2 = w.tills
    event = create(w)
    sale(w, t1, 10, "100.00")
    sale(w, t2, 20, "50.00")

    # Before the end: blocked unless the user insists.
    with pytest.raises(HTTPException) as e:
        C.confirm_event(w.db, w.admin, w.tenant.id, event, force=False, note=None, now=at(60))
    assert e.value.status_code == 409 and e.value.detail["code"] == "event_not_ready"
    assert e.value.detail["readiness"]["blocking"][0]["code"] == "not_ended"

    snap = C.confirm_event(w.db, w.admin, w.tenant.id, event, force=False, note="מאושר למפיק", now=AFTER)
    w.db.commit()
    assert snap["frozen"] is True and snap["kpis"]["net"] == 150.0
    assert event.status == "confirmed" and event.snapshot["kpis"]["net"] == 150.0
    rows = w.db.query(ReportEventMachine).filter(ReportEventMachine.event_id == event.id).all()
    assert len(rows) == 2 and all(r.released_at is not None for r in rows)
    assert [m["name"] for m in snap["event"]["machines"]] == ["Till 1", "Till 2"]

    # A sale that lands later, inside the window, changes nothing that was confirmed.
    sale(w, t1, 30, "999.00")
    w.db.commit()
    again = report_for(w.db, event)
    assert again["frozen"] is True and again["kpis"]["net"] == 150.0
    assert build_report(w.db, event, now=AFTER)["kpis"]["net"] == 1149.0  # the live one would

    # Read-only now.
    with pytest.raises(HTTPException) as e:
        R.update_report_event(event.id, ReportEventUpdate(name="x"), current_user=w.admin,
                              active_tenant_id=w.tenant.id, db=w.db)
    assert e.value.status_code == 409 and e.value.detail["code"] == "event_confirmed"
    with pytest.raises(HTTPException) as e:
        R.delete_report_event(event.id, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert e.value.status_code == 409
    with pytest.raises(HTTPException):
        R.confirm_report_event(event.id, ReportEventConfirm(), current_user=w.admin,
                               active_tenant_id=w.tenant.id, db=w.db)

    # The released tills may join a new event over the same hours.
    create(w, name="שוב", tills=[t1, t2])


def test_confirm_before_the_end_with_force_records_it(w):
    t1, _t2 = w.tills
    event = create(w)
    sale(w, t1, 10, "10.00")
    snap = C.confirm_event(w.db, w.admin, w.tenant.id, event, force=True, note=None, now=at(60))
    assert snap["frozen"] is True
    assert event.confirm_note == "אושר לפני סוף האירוע"


def test_a_draft_can_be_deleted_and_edited(w):
    t1, t2 = w.tills
    event = create(w, tills=[t1])
    out = R.update_report_event(
        event.id, ReportEventUpdate(name="חדש", machineIds=[t1.id, t2.id], endTime="23:30"),
        current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db,
    )
    assert out["name"] == "חדש" and out["endTime"] == "23:30" and len(out["machineIds"]) == 2
    R.delete_report_event(event.id, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert w.db.query(ReportEvent).count() == 0
    assert w.db.query(ReportEventMachine).count() == 0


# ── Reconciliation: Z ─────────────────────────────────────────────────────────


def build_z(w, till, shifts, number=1, tweak=None):
    w.db.flush()
    section = machine_section(till, shifts, compute_totals(w.db, [s.id for s in shifts]))
    if tweak:
        tweak(section)
    z = ZReport(
        id=uuid.uuid4(), tenant_id=w.tenant.id, shop_id=w.shop.id, origin=ZOrigin.CLOUD,
        business_date=date(2026, 9, 27), closed_at=AFTER, shop_sequence_number=number,
        per_machine=[section], period_start=shifts[0].opened_at, period_end=AFTER,
    )
    w.db.add(z)
    w.db.flush()
    for s in shifts:
        s.z_report_id = z.id
    w.db.flush()
    return z


def test_the_z_matches_its_documents_and_says_it_covers_more_than_the_window(w):
    t1, _t2 = w.tills
    event = create(w, tills=[t1])
    shift = w.shift(t1, 1, opened_at=at(-120))
    sale(w, t1, -60, "40.00", shift=shift)                 # before the event, same shift
    sale(w, t1, 30, "100.00", shift=shift, method="card")
    sale(w, t1, 60, "10.00", shift=shift, credit_note=True)
    build_z(w, t1, [shift])
    rep = report(w, event)
    z = rep["reconciliation"]["z"]
    assert z["status"] == "match"
    row = z["rows"][0]
    assert row["status"] == "match"
    assert row["coversOutside"] is True and row["outsideNet"] == 40.0
    assert row["inWindow"]["sales"] == 100.0 and row["inWindow"]["refunds"] == 10.0
    assert z["tills"][0]["pendingZ"]["count"] == 0
    assert not [c for c in codes(rep) if c.startswith("z_")]


def test_a_z_that_does_not_add_up_to_its_documents_is_an_alert(w):
    t1, _t2 = w.tills
    event = create(w, tills=[t1])
    shift = w.shift(t1, 1, opened_at=at(-10))
    sale(w, t1, 30, "100.00", shift=shift)

    def off(section):
        section["totalSales"] = "110.00"
        section["totalCash"] = "110.00"

    build_z(w, t1, [shift], tweak=off)
    rep = report(w, event)
    row = rep["reconciliation"]["z"]["rows"][0]
    assert row["status"] == "mismatch" and row["diff"]["sales"] == 10.0
    assert rep["reconciliation"]["z"]["status"] == "mismatch"
    assert "z_mismatch" in codes(rep, "alert")


def test_documents_in_no_z_yet_are_flagged(w):
    t1, _t2 = w.tills
    event = create(w, tills=[t1])
    shift = w.shift(t1, 1, opened_at=at(-10))
    sale(w, t1, 30, "100.00", shift=shift)
    sale(w, t1, 40, "20.00")                       # no shift at all
    rep = report(w, event)
    z = rep["reconciliation"]["z"]
    assert z["status"] == "pending"
    assert z["tills"][0]["pendingZ"] == {"count": 2, "net": 120.0}
    assert "z_pending" in codes(rep, "warning")
    assert any(x["code"] == "pending_z" for x in rep["readiness"]["warnings"])


# ── Reconciliation: card transmissions ────────────────────────────────────────


def card_sale(w, till, minutes, total, uid):
    tx = sale(w, till, minutes, total, method="card")
    leg = w.db.query(TransactionPayment).filter(TransactionPayment.transaction_id == tx.id).one()
    leg.terminal_uid = uid
    w.db.flush()
    return tx, leg


def batch(w, till, legs, *, status="success", count=None, amount=None, minutes=320, with_items=True):
    b = CardTransmission(
        id=uuid.uuid4(), tenant_id=w.tenant.id, machine_id=till.id, shop_id=w.shop.id, trigger="shift_close",
        started_at=at(minutes), finished_at=at(minutes + 1), status=status, batch_number="B-17",
        transaction_count=count if count is not None else len(legs),
        amount=Decimal(amount) if amount is not None else sum((l.amount for l in legs), Decimal("0")),
        terminal_transaction_count=len(legs) if with_items else 0,
    )
    w.db.add(b)
    w.db.flush()
    if status == "success":
        for leg in legs:
            if with_items:
                w.db.add(CardTransmissionItem(id=uuid.uuid4(), transmission_id=b.id, machine_id=till.id,
                                              terminal_uid=leg.terminal_uid))
            leg.transmission_id = b.id
            leg.transmitted_batch = b.batch_number
    w.db.flush()
    return b


def test_transmitted_card_sales_match_their_batch(w):
    t1, _t2 = w.tills
    t1.transmission_tracking_started_at = at(-600)
    event = create(w, tills=[t1])
    _tx, a = card_sale(w, t1, 10, "40.00", "u1")
    _tx, b = card_sale(w, t1, 20, "60.00", "u2")
    batch(w, t1, [a, b])
    rep = report(w, event)
    tx = rep["reconciliation"]["transmissions"]
    till = tx["tills"][0]
    assert tx["status"] == "match"
    assert till["transmitted"] == {"count": 2, "amount": 100.0}
    assert till["batches"][0]["compare"] == "match" and till["batches"][0]["level"] == "transaction"
    assert not [c for c in codes(rep) if c.startswith("tx_")]


def test_a_batch_whose_count_differs_is_an_alert(w):
    t1, _t2 = w.tills
    t1.transmission_tracking_started_at = at(-600)
    event = create(w, tills=[t1])
    _tx, a = card_sale(w, t1, 10, "40.00", "u1")
    batch(w, t1, [a], count=2, amount="80.00")
    rep = report(w, event)
    assert rep["reconciliation"]["transmissions"]["tills"][0]["batches"][0]["compare"] == "mismatch"
    assert "tx_mismatch" in codes(rep, "alert")


def test_a_batch_without_uids_is_compared_at_the_level_of_amounts(w):
    t1, _t2 = w.tills
    t1.transmission_tracking_started_at = at(-600)
    event = create(w, tills=[t1])
    _tx, a = card_sale(w, t1, 10, "40.00", "u1")
    _tx, b = card_sale(w, t1, 20, "60.00", "u2")
    batch(w, t1, [a, b], with_items=False)
    till = report(w, event)["reconciliation"]["transmissions"]["tills"][0]
    assert till["batches"][0]["level"] == "amounts"
    assert till["batches"][0]["compare"] == "match" and till["batches"][0]["amountCompared"] == 100.0


def test_untransmitted_card_sales_and_a_failed_batch_are_flagged(w):
    t1, _t2 = w.tills
    t1.transmission_tracking_started_at = at(-600)
    event = create(w, tills=[t1])
    card_sale(w, t1, 10, "40.00", "u1")
    batch(w, t1, [], status="failed", count=0, amount="0")
    rep = report(w, event)
    till = rep["reconciliation"]["transmissions"]["tills"][0]
    assert till["untransmitted"]["count"] == 1 and till["untransmitted"]["amount"] == 40.0
    assert "tx_untransmitted" in codes(rep, "warning")
    assert "tx_failed" in codes(rep, "alert")
    assert any(x["code"] == "untransmitted" for x in rep["readiness"]["warnings"])


# ── Export, compare ───────────────────────────────────────────────────────────


def test_the_export_is_a_workbook_with_the_expected_sheets(w):
    t1, t2 = w.tills
    event = create(w, producerName="הפקות בע״מ")
    sale(w, t1, 10, "100.00", tip="30.00")
    sale(w, t2, 20, "50.00", method="card")
    sale(w, t2, 25, "10.00", credit_note=True)
    response = R.export_report_event(event.id, bucket=60, current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert response.media_type.startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(response.body))
    assert wb.sheetnames == ["סיכום", "תובנות", "קופות", "פריטים", "פילוח", "ציר זמן", "משמרות", "התאמות"]
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)
    assert wb["סיכום"]["A1"].value == "במה ראשית"
    assert wb["קופות"]["A1"].value == "קופה"
    # 18:00–23:00 in hours: five buckets on the timeline.
    assert wb["ציר זמן"].cell(row=3 + 5, column=1).value == "סה״כ"


def test_two_events_can_be_compared(w):
    t1, t2 = w.tills
    a = create(w, name="א", tills=[t1])
    b = create(w, name="ב", tills=[t2])
    sale(w, t1, 10, "100.00")
    sale(w, t2, 10, "40.00")
    out = R.compare_report_events(ids=f"{a.id},{b.id}", current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
    assert [e["event"]["name"] for e in out["events"]] == ["א", "ב"]
    assert [e["kpis"]["net"] for e in out["events"]] == [100.0, 40.0]
    with pytest.raises(HTTPException):
        R.compare_report_events(ids=str(a.id), current_user=w.admin, active_tenant_id=w.tenant.id, db=w.db)
