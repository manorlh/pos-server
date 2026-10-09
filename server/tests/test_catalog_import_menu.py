"""
The menu sheet's restaurant layer - app/services/catalog_import_menu.py: add-on groups and
their options, which groups a product or category gets, quick notes, pictures by link and
catalog-menu prices, imported with (and like) the products and categories.

What each class pins, and how it could look fine while doing damage:

* **Template** - the new sheets with examples, star headers, dropdowns that list the
  groups typed in the groups sheet, the company's groups pre-filled, a price column per
  catalog menu; the whole menu exports and imports back as "no change".
* **Groups and options** - created, updated, an empty cell keeps, nothing absent deleted,
  the same file twice changes nothing; the group editor's rules (required / minimum,
  removal without quantities, an option's own maximum) refused in Hebrew, not written.
* **Links** - a product's / category's groups in order, "ללא", "ירושה"; an unknown group
  or one only in an example row refused with its row.
* **Quick notes** - on categories, products, every dish; a product that only inherited
  chips keeps them; idempotent; a note without a target or with an unknown one refused.
* **Pictures** - downloaded once (the network is mocked), stored where an upload is, a
  broken link reported without failing the import, internal addresses refused.
* **Menu prices** - updated, a product added to the menu, "ללא" back to the catalog's
  price, an unknown menu column warned and skipped.
* **Permissions** - a company manager cannot change the organization's group or menu.
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

import openpyxl
import pytest
from PIL import Image

from app.models.catalog_menu import CatalogMenu, CatalogMenuCategory, CatalogMenuProduct, CatalogMenuSyncState
from app.models.category import Category
from app.models.menu import MenuSyncState, ModifierGroup, ModifierLink, ModifierOption, PrepNotePreset
from app.models.product import Product
from app.services import catalog_images as IMG
from app.services import catalog_import as I
from app.services import catalog_sheet as S
from app.services import local_media
from test_catalog_import import commit, messages, preview, products_named, template, w  # noqa: F401


def _png(color=(200, 30, 30)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 30), color).save(out, format="PNG")
    return out.getvalue()


@pytest.fixture
def m(w, monkeypatch, tmp_path):  # noqa: F811 - `w` is test_catalog_import's fixture
    """The import world, with a menu layer: groups, a link, a chip and a lunch menu."""
    db = w.db
    t = w.tenant.id
    w.doneness = ModifierGroup(id=uuid.uuid4(), tenant_id=t, company_id=w.company.id, name="מידת עשייה",
                               kind="choice", min_select=1, max_select=1, sort_order=1)
    w.sauces = ModifierGroup(id=uuid.uuid4(), tenant_id=t, company_id=None, name="רטבים", kind="addon",
                             sort_order=2)
    db.add_all([w.doneness, w.sauces])
    db.flush()
    w.medium = ModifierOption(id=uuid.uuid4(), group_id=w.doneness.id, name="מדיום", price=Decimal("0"),
                              is_default=True, sort_order=0)
    w.well = ModifierOption(id=uuid.uuid4(), group_id=w.doneness.id, name="וול דאן", price=Decimal("0"), sort_order=1)
    w.chili = ModifierOption(id=uuid.uuid4(), group_id=w.sauces.id, name="צ'ילי", price=Decimal("2"), sort_order=0)
    db.add_all([w.medium, w.well, w.chili])
    db.add(ModifierLink(id=uuid.uuid4(), tenant_id=t, target_type="category", target_id=w.burgers.id,
                        group_id=w.doneness.id, sort_order=0))
    db.add(PrepNotePreset(id=uuid.uuid4(), tenant_id=t, target_type="category", target_id=w.burgers.id,
                          text="בלי בצל", sort_order=0))
    w.lunch = CatalogMenu(id=uuid.uuid4(), tenant_id=t, company_id=w.company.id, name="צהריים", channel="both",
                          always=True)
    db.add(w.lunch)
    db.flush()
    db.add(CatalogMenuCategory(id=uuid.uuid4(), menu_id=w.lunch.id, category_id=w.drinks.id, all_products=True))
    db.add(CatalogMenuProduct(id=uuid.uuid4(), menu_id=w.lunch.id, product_id=w.cola.id, sort_order=0,
                              price=Decimal("10")))
    db.commit()

    # Pictures: stored under the test's own folder, never Cloudinary, never the network.
    monkeypatch.setattr(local_media, "MEDIA_DIR", tmp_path)
    monkeypatch.setattr("app.services.cloudinary_service.cloudinary_configured", lambda settings=None: False)
    w.media = tmp_path
    w.downloads = []

    def fake_download(url):
        w.downloads.append(url)
        if "broken" in url:
            raise IMG.ImageError("השרת החזיר שגיאה (404) - בדקו שהקישור פתוח לכולם")
        return _png()

    monkeypatch.setattr(IMG, "download", fake_download)
    return w


def book(products=None, categories=None, groups=None, options=None, notes=None, menus=()) -> bytes:
    """A filled-in workbook: each sheet's rows as {column key: value}; None leaves the sheet out."""
    wb = openpyxl.Workbook()
    first = True
    product_columns = S.PRODUCT_COLUMNS + tuple(S.menu_price_column(n) for n in menus)
    for title, columns, rows in (
        (S.SHEET_PRODUCTS, product_columns, products), (S.SHEET_CATEGORIES, S.CATEGORY_COLUMNS, categories),
        (S.SHEET_GROUPS, S.GROUP_COLUMNS, groups), (S.SHEET_OPTIONS, S.OPTION_COLUMNS, options),
        (S.SHEET_NOTES, S.NOTE_COLUMNS, notes),
    ):
        if rows is None:
            continue
        ws = wb.active if first else wb.create_sheet(title)
        first = False
        ws.title = title
        ws.append([c.header for c in columns])
        for row in rows:
            ws.append([row.get(c.key) for c in columns])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def own_groups(world, target_type, target_id):
    rows = world.db.query(ModifierLink).filter(ModifierLink.target_type == target_type,
                                           ModifierLink.target_id == target_id).order_by(ModifierLink.sort_order).all()
    return [r.group_id for r in rows]


def own_notes(world, target_type, target_id):
    rows = world.db.query(PrepNotePreset).filter(PrepNotePreset.target_type == target_type,
                                             PrepNotePreset.target_id == target_id).order_by(PrepNotePreset.sort_order)
    return [(n.text, n.sort_order, bool(n.is_important)) for n in rows]


def group_named(world, name):
    return world.db.query(ModifierGroup).filter(ModifierGroup.name == name).one()


def all_messages(out):
    return [m["text"] for key in ("products", "categories", "groups", "options", "notes")
            for r in out[key] for m in r["messages"]] + [i["text"] for i in out["issues"]]


# ── Template ──────────────────────────────────────────────────────────────────


class TestTemplateSheets:
    def test_new_sheets_with_examples_star_headers_and_dropdowns(self, m):
        wb = openpyxl.load_workbook(io.BytesIO(template(m)))
        groups, options, notes = wb[S.SHEET_GROUPS], wb[S.SHEET_OPTIONS], wb[S.SHEET_NOTES]
        for ws, columns, required in ((groups, S.GROUP_COLUMNS, {"שם קבוצה*"}),
                                      (options, S.OPTION_COLUMNS, {"קבוצה*", "אפשרות*"}),
                                      (notes, S.NOTE_COLUMNS, {"הערה*"})):
            assert ws.sheet_view.rightToLeft is True
            assert [c.value for c in ws[1]] == [c.header for c in columns]
            assert {c.value for c in ws[1] if c.value.endswith("*")} == required
            assert all(c.comment is not None for c in ws[1])
            assert ws.cell(2, 1).value == "דוגמה"
        assert options.freeze_panes == "D2" and groups.freeze_panes == "C2"
        # The company's groups follow the examples even in the blank template (the dropdowns
        # list them); their options and prices only in the current menu's - the blank one is
        # what the public share link serves.
        assert groups.cell(5, 2).value == "מידת עשייה" and groups.cell(6, 2).value == "רטבים"
        assert options.max_row == 9
        current = openpyxl.load_workbook(io.BytesIO(template(m, with_data=True)))[S.SHEET_OPTIONS]
        assert [current.cell(r, 3).value for r in (2, 3, 4)] == ["מדיום", "וול דאן", "צ'ילי"]
        assert current.cell(4, 4).value == 2

        def validation(ws, key, columns):
            letter = openpyxl.utils.get_column_letter([c.key for c in columns].index(key) + 1)
            return next(dv for dv in ws.data_validations.dataValidation
                        if any(str(r).startswith(letter + "2") for r in str(dv.sqref).split()))

        products = wb[S.SHEET_PRODUCTS]
        columns = S.PRODUCT_COLUMNS + (S.menu_price_column("צהריים"),)
        assert products.cell(1, len(columns)).value == "מחיר בתפריט: צהריים"
        assert validation(products, "groups", columns).formula1.startswith(f"'{S.SHEET_GROUPS}'!$B$2")
        assert validation(options, "group", S.OPTION_COLUMNS).formula1.startswith(f"'{S.SHEET_GROUPS}'!$B$2")
        assert validation(groups, "kind", S.GROUP_COLUMNS).formula1.startswith(f"'{S.SHEET_LISTS}'!")
        assert validation(products, "image", columns).type == "custom"
        assert validation(options, "price", S.OPTION_COLUMNS).formula1 == "-1000"
        assert validation(notes, "products", S.NOTE_COLUMNS).formula1.startswith(f"'{S.SHEET_PRODUCTS}'!$B$2")

    def test_the_whole_menu_exports_and_imports_back_as_no_change(self, m):
        db = m.db
        t = m.tenant.id
        db.add(ModifierLink(id=uuid.uuid4(), tenant_id=t, target_type="product", target_id=m.classic.id,
                            group_id=m.doneness.id, sort_order=0))
        db.add(ModifierLink(id=uuid.uuid4(), tenant_id=t, target_type="product", target_id=m.classic.id,
                            group_id=m.sauces.id, sort_order=1))
        db.add(ModifierLink(id=uuid.uuid4(), tenant_id=t, target_type="category", target_id=m.drinks.id,
                            group_id=None, sort_order=0))
        db.add(PrepNotePreset(id=uuid.uuid4(), tenant_id=t, target_type="product", target_id=m.cola.id,
                              text="בלי קרח", sort_order=0))
        db.add(PrepNotePreset(id=uuid.uuid4(), tenant_id=t, company_id=m.company.id, target_type="all",
                              text="אלרגיה", is_important=True, sort_order=0))
        m.cola.image_url = "http://localhost:8001/media/pos/x/products/abc.png"
        m.classic.image_url = f"https://cdn.example.com/{IMG.marker_of('https://a.example.com/b.jpg')}.png"
        m.sauces.max_select = 2
        db.commit()
        data = template(m, with_data=True)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        products = wb[S.SHEET_PRODUCTS]
        keys = [c.key for c in S.PRODUCT_COLUMNS] + ["menu:צהריים"]
        rows = {products.cell(r, 2).value: {k: products.cell(r, i).value for i, k in enumerate(keys, start=1)}
                for r in range(2, 5)}
        assert rows["המבורגר קלאסי"]["groups"] == "מידת עשייה, רטבים"
        assert rows["קולה"]["menu:צהריים"] == 10 and rows["קולה"]["image"].endswith("abc.png")
        notes = wb[S.SHEET_NOTES]
        exported = {notes.cell(r, 2).value: [notes.cell(r, c).value for c in range(3, 7)] for r in range(2, 5)}
        assert exported == {"אלרגיה": ["כן", None, None, "כן"], "בלי בצל": [None, "המבורגרים", None, "לא"],
                            "בלי קרח": [None, None, "קולה", "לא"]}

        out = preview(m, data)
        s = out["summary"]
        assert s["errors"] == 0 and s["warnings"] == 0, all_messages(out)
        assert (s["productsUnchanged"], s["categoriesUnchanged"], s["groupsUnchanged"], s["optionsUnchanged"],
                s["notesUnchanged"]) == (3, 3, 2, 3, 3)
        assert (s["linkChanges"], s["imageChanges"], s["menuPriceChanges"]) == (0, 0, 0)
        assert out["canCommit"] is False

    def test_two_groups_by_one_name_are_shown_skipped_and_still_import_back_clean(self, m):
        twin = ModifierGroup(id=uuid.uuid4(), tenant_id=m.tenant.id, company_id=m.company.id, name="מידת עשייה",
                             kind="choice", min_select=1, max_select=1, sort_order=5)
        m.db.add(twin)
        m.db.flush()
        m.db.add(ModifierOption(id=uuid.uuid4(), group_id=twin.id, name="רייר", price=Decimal("0"), sort_order=0))
        m.db.add(ModifierLink(id=uuid.uuid4(), tenant_id=m.tenant.id, target_type="product", target_id=m.classic.id,
                              group_id=twin.id, sort_order=0))
        m.db.commit()
        data = template(m, with_data=True)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        groups = wb[S.SHEET_GROUPS]
        marked = [(groups.cell(r, 1).value, groups.cell(r, 2).value) for r in range(2, groups.max_row + 1)]
        assert marked.count(("דלג", "מידת עשייה")) == 2
        assert "יש יותר מקבוצה אחת בשם" in groups.cell(2, 1).comment.text
        products = wb[S.SHEET_PRODUCTS]
        column = [c.key for c in S.PRODUCT_COLUMNS].index("groups") + 1
        classic = next(r for r in range(2, 5) if products.cell(r, 2).value == "המבורגר קלאסי")
        assert products.cell(classic, column).value is None
        assert "שם כפול" in products.cell(classic, column).comment.text
        out = preview(m, data)
        assert out["summary"]["errors"] == 0 and out["canCommit"] is False, all_messages(out)

    def test_the_csv_carries_the_menu_price_column(self, m):
        body = template(m, with_data=True, fmt="csv").decode("utf-8-sig")
        assert body.splitlines()[0].endswith("מחיר בתפריט: צהריים")


# ── Groups and options ────────────────────────────────────────────────────────


NEW_GROUPS = [{"name": "תוספות", "kind": "תוספת", "max": 3, "free": 1, "allow_quantity": "כן"}]
NEW_OPTIONS = [
    {"group": "תוספות", "name": "גבינה", "price": 4},
    {"group": "תוספות", "name": "ביצה", "price": "₪5", "max_qty": 2, "kitchen_name": "ביצ"},
    {"group": "מידת עשייה", "name": "מדיום וול", "price": 0},
]


class TestGroups:
    def test_groups_and_options_are_created_and_attached(self, m):
        data = book(
            products=[{"name": "המבורגר קלאסי", "category": "המבורגרים", "groups": "מידת עשייה, תוספות"}],
            categories=[{"name": "שתייה", "groups": "ללא"}],
            groups=NEW_GROUPS, options=NEW_OPTIONS,
        )
        out = preview(m, data)
        assert [r["status"] for r in out["groups"]] == ["create"]
        assert [r["status"] for r in out["options"]] == ["create"] * 3
        assert out["options"][1]["price"] == "5.00"
        assert out["products"][0]["status"] == "update"
        assert out["products"][0]["changes"] == [{"field": "groups", "label": "קבוצות תוספות",
                                                  "before": "לפי המחלקה", "after": "מידת עשייה, תוספות"}]
        assert out["summary"]["linkChanges"] == 2 and out["canCommit"] is True

        result = commit(m, data, token=out["token"])
        assert (result["groupsCreated"], result["optionsCreated"], result["linksChanged"]) == (1, 3, 2)
        extras = group_named(m, "תוספות")
        assert (extras.company_id, extras.kind, extras.min_select, extras.max_select, extras.free_count,
                extras.allow_quantity, extras.sort_order) == (m.company.id, "addon", 0, 3, 1, True, 3)
        egg = m.db.query(ModifierOption).filter(ModifierOption.name == "ביצה").one()
        assert (egg.group_id, egg.price, egg.max_qty, egg.kitchen_name, egg.sort_order) == (
            extras.id, Decimal("5.00"), 2, "ביצ", 1)
        assert m.db.query(ModifierOption).filter(ModifierOption.name == "מדיום וול").one().sort_order == 2
        assert own_groups(m, "product", m.classic.id) == [m.doneness.id, extras.id]
        assert own_groups(m, "category", m.drinks.id) == [None]  # "ללא": an explicit none
        assert m.db.query(MenuSyncState).filter(MenuSyncState.tenant_id == m.tenant.id).count() == 1
        assert {machine for machine, _ in m.catalog} == {str(t.id) for t in m.tills} | {str(m.other_till.id)}

    def test_the_same_file_twice_changes_nothing(self, m):
        data = book(products=[{"name": "המבורגר קלאסי", "category": "המבורגרים", "groups": "מידת עשייה, תוספות"}],
                    groups=NEW_GROUPS, options=NEW_OPTIONS,
                    notes=[{"name": "בלי חסה", "products": "המבורגר קלאסי"}])
        commit(m, data)
        m.catalog.clear()
        out = preview(m, data)
        assert out["canCommit"] is False and out["summary"]["errors"] == 0
        assert {r["status"] for key in ("groups", "options", "notes", "products") for r in out[key]} == {"unchanged"}
        again = commit(m, data)
        assert (again["groupsCreated"], again["optionsCreated"], again["linksChanged"], again["notesCreated"]) == (0, 0, 0, 0)
        assert m.catalog == []
        assert m.db.query(ModifierGroup).filter(ModifierGroup.name == "תוספות").count() == 1

    def test_an_empty_cell_keeps_and_a_filled_one_updates(self, m):
        out = preview(m, book(groups=[{"name": "מידת עשייה", "sort": 7}],
                              options=[{"group": "מידת עשייה", "name": "מדיום", "price": 2}]))
        assert [(c["field"], c["before"], c["after"]) for c in out["groups"][0]["changes"]] == [("sort_order", 1, 7)]
        assert [(c["field"], c["before"], c["after"]) for c in out["options"][0]["changes"]] == [
            ("price", "0.00", "2.00")]
        commit(m, book(groups=[{"name": "מידת עשייה", "sort": 7}],
                       options=[{"group": "מידת עשייה", "name": "מדיום", "price": 2}]))
        group = m.db.get(ModifierGroup, m.doneness.id)
        assert (group.sort_order, group.kind, group.min_select, group.max_select) == (7, "choice", 1, 1)
        medium = m.db.get(ModifierOption, m.medium.id)
        assert (medium.price, medium.is_default) == (Decimal("2.00"), True)
        # Nothing absent is deleted.
        assert m.db.get(ModifierOption, m.well.id) is not None

    def test_required_and_the_minimum(self, m):
        out = preview(m, book(
            groups=[{"name": "גודל", "kind": "בחירה", "required": "כן", "max": 1},
                    {"name": "א", "required": "כן", "min": 0},
                    {"name": "ב", "required": "לא", "min": 2}],
            options=[{"group": "גודל", "name": "קטן"}, {"group": "גודל", "name": "גדול", "price": 6}],
        ))
        assert [r["status"] for r in out["groups"]] == ["create", "error", "error"]
        texts = all_messages(out)
        assert "שורה 3: 'חובה' = כן דורש 'מינימום בחירות' 1 לפחות" in texts
        assert any(t.startswith("שורה 4: 'מינימום בחירות' מעל 0 הופך את הקבוצה לחובה") for t in texts)
        commit(m, book(groups=[{"name": "גודל", "kind": "בחירה", "required": "כן", "max": 1}],
                       options=[{"group": "גודל", "name": "קטן"}]))
        size = group_named(m, "גודל")
        assert (size.kind, size.min_select, size.max_select) == ("choice", 1, 1)

    def test_rules_that_do_not_hold_are_refused_in_hebrew(self, m):
        out = preview(m, book(
            groups=[{"name": "בלי", "kind": "הסרה", "allow_quantity": "כן"},
                    {"name": "חובה ריקה", "required": "כן"},
                    {"name": "משהו", "kind": "קישוט"}],
            options=[{"group": "בלי", "name": "בלי בצל"},
                     {"group": "מידת עשייה", "name": "רייר", "max_qty": 3},
                     {"group": "לא קיימת", "name": "x"},
                     {"group": "מידת עשייה", "name": "זול", "price": -2000}],
        ))
        texts = all_messages(out)
        assert "שורה 2: הקבוצה 'בלי' לא תקינה: בקבוצת 'הסרה' אין 'כמות לאפשרות' ואין 'מעט / הרבה / בצד'" in texts
        assert "שורה 2: הקבוצה 'בלי' לא תקינה (שורה 2 בגיליון 'קבוצות תוספות')" in texts
        assert any("קבוצת חובה צריכה לפחות אפשרות פעילה אחת" in t for t in texts)
        assert any(t.startswith("שורה 4: סוג קבוצה לא מוכר ('קישוט')") for t in texts)
        assert "שורה 3: הקבוצה 'מידת עשייה' לא תקינה: 'מקסימום לאפשרות' מעל 1 דורש 'כמות לאפשרות' = כן בקבוצה" \
            in texts
        assert "שורה 4: קבוצת התוספות 'לא קיימת' לא קיימת - הוסיפו אותה בגיליון 'קבוצות תוספות'" in texts
        assert "שורה 5: מחיר נמוך מדי (-2000.00, לכל הפחות -1000.00)" in texts
        assert out["canCommit"] is False or out["summary"]["errors"] > 0
        assert {r["status"] for r in out["groups"]} == {"error"}

    def test_a_group_named_only_in_an_example_row(self, m):
        out = preview(m, book(products=[{"name": "קולה", "category": "שתייה", "groups": "חדשה"}],
                              groups=[{"marker": "דוגמה", "name": "חדשה"}]))
        assert out["products"][0]["status"] == "error"
        assert "מופיעה רק בשורת דוגמה" in out["products"][0]["messages"][0]["text"]

    def test_a_product_back_to_its_category_and_to_none(self, m):
        commit(m, book(products=[{"name": "המבורגר קלאסי", "category": "המבורגרים", "groups": "רטבים"}]))
        assert own_groups(m, "product", m.classic.id) == [m.sauces.id]
        commit(m, book(products=[{"name": "המבורגר קלאסי", "category": "המבורגרים", "groups": "ירושה"}]))
        assert own_groups(m, "product", m.classic.id) == []
        commit(m, book(products=[{"name": "המבורגר קלאסי", "category": "המבורגרים", "groups": "ללא"}]))
        assert own_groups(m, "product", m.classic.id) == [None]

    def test_a_company_manager_cannot_change_the_organizations_group(self, m):
        out = preview(m, book(groups=[{"name": "רטבים", "max": 9}], options=[{"group": "רטבים", "name": "איולי"}],
                              products=[{"name": "קולה", "category": "שתייה", "groups": "רטבים"}]), user=m.manager)
        assert out["groups"][0]["status"] == "error"
        assert "אין לך הרשאה לעדכן את הקבוצה 'רטבים'" in out["groups"][0]["messages"][0]["text"]
        assert out["options"][0]["status"] == "error"
        # Using the organization's group on the company's own product is allowed.
        assert out["products"][0]["status"] == "update"


# ── Quick notes ───────────────────────────────────────────────────────────────


class TestNotes:
    def test_notes_on_categories_products_and_every_dish(self, m):
        data = book(notes=[
            {"name": "רוטב בצד", "categories": "המבורגרים"},
            {"name": "בלי קרח", "products": "קולה, 7290000000017"},
            {"name": "אלרגיה לאגוזים", "all": "כן", "important": "כן"},
            {"name": "בלי בצל", "categories": "המבורגרים"},
        ])
        out = preview(m, data)
        assert [r["status"] for r in out["notes"]] == ["create", "create", "create", "unchanged"]
        assert out["notes"][1]["targets"] == "קולה"
        result = commit(m, data)
        assert result["notesCreated"] == 3
        assert own_notes(m, "category", m.burgers.id) == [("בלי בצל", 0, False), ("רוטב בצד", 1, False)]
        assert own_notes(m, "product", m.cola.id) == [("בלי קרח", 0, False)]
        every = m.db.query(PrepNotePreset).filter(PrepNotePreset.target_type == "all").one()
        assert (every.text, every.company_id, every.is_important) == ("אלרגיה לאגוזים", m.company.id, True)
        assert m.catalog  # the tills are told

    def test_a_product_that_only_inherited_chips_keeps_them(self, m):
        data = book(notes=[{"name": "בלי חסה", "products": "המבורגר קלאסי", "important": "כן"}])
        out = preview(m, data)
        assert any("יישמרו לו גם הן" in t and "בלי בצל" in t for t in all_messages(out))
        commit(m, data)
        assert own_notes(m, "product", m.classic.id) == [("בלי בצל", 0, False), ("בלי חסה", 1, True)]
        again = preview(m, data)
        assert again["notes"][0]["status"] == "unchanged" and again["canCommit"] is False

    def test_an_unchanged_chip_needs_no_permission(self, m):
        # The organization's product: a company manager may not change its chips - and
        # need not, to import back a file that names the ones it has.
        shared = Product(id=uuid.uuid4(), tenant_id=m.tenant.id, company_id=None, category_id=m.drinks.id,
                         name="סודה", price=Decimal("6"), sku="3001")
        m.db.add(shared)
        m.db.add(PrepNotePreset(id=uuid.uuid4(), tenant_id=m.tenant.id, target_type="product", target_id=shared.id,
                                text="קר", sort_order=0))
        m.db.commit()
        same = preview(m, book(notes=[{"name": "קר", "products": "סודה"}]), user=m.manager)
        assert same["notes"][0]["status"] == "unchanged" and same["summary"]["errors"] == 0
        added = preview(m, book(notes=[{"name": "חם", "products": "סודה, קולה"}]), user=m.manager)
        assert added["notes"][0]["status"] == "error"
        assert "אין לך הרשאה לשנות את ההערות של 'סודה'" in added["notes"][0]["messages"][0]["text"]
        # A refused row writes nothing, not even on the target it was allowed.
        commit(m, book(notes=[{"name": "חם", "products": "קולה, סודה"}]), user=m.manager, skip_errors=True)
        assert own_notes(m, "product", m.cola.id) == []

    def test_an_existing_chip_is_updated_not_duplicated(self, m):
        commit(m, book(notes=[{"name": "בלי בצל", "categories": "המבורגרים", "important": "כן", "sort": 4}]))
        assert own_notes(m, "category", m.burgers.id) == [("בלי בצל", 4, True)]

    def test_bad_note_rows_are_refused_in_hebrew(self, m):
        out = preview(m, book(notes=[
            {"name": "בלי כלום"},
            {"name": "x", "categories": "אין כזו"},
            {"name": "y", "products": "אין כזה"},
            {"name": "z" * 61, "all": "כן"},
        ]))
        assert [r["status"] for r in out["notes"]] == ["error"] * 4
        texts = all_messages(out)
        assert "שורה 2: ציינו על אילו מחלקות או פריטים ההערה חלה, או כתבו 'כן' בעמודה 'לכל המנות'" in texts
        assert "שורה 3: המחלקה 'אין כזו' לא נמצאה - לא בגיליון 'מחלקות' ולא במערכת" in texts
        assert "שורה 4: הפריט 'אין כזה' לא נמצא - לא בגיליון 'פריטים' ולא במערכת" in texts
        assert "שורה 5: ההערה ארוך מדי (מעל 60 תווים)" in texts

    def test_two_products_by_one_name_are_named_by_sku(self, m):
        for sku, category in (("S-1", m.drinks), ("S-2", m.food)):
            m.db.add(Product(id=uuid.uuid4(), tenant_id=m.tenant.id, company_id=m.company.id, category_id=category.id,
                             name="סלט", price=Decimal("20"), sku=sku))
        m.db.commit()
        out = preview(m, book(notes=[{"name": "בלי שמן", "products": "סלט"}]))
        assert "כתבו את המק״ט במקום השם" in out["notes"][0]["messages"][0]["text"]
        assert preview(m, book(notes=[{"name": "בלי שמן", "products": "S-2"}]))["notes"][0]["status"] == "create"

    def test_a_note_on_a_product_new_in_the_same_file(self, m):
        commit(m, book(products=[{"name": "שקשוקה", "category": "אוכל", "price": 48}],
                       notes=[{"name": "חריף", "products": "שקשוקה"}]))
        shakshuka = products_named(m, "שקשוקה")[0]
        assert own_notes(m, "product", shakshuka.id) == [("חריף", 0, False)]


# ── Pictures ──────────────────────────────────────────────────────────────────


class TestImages:
    def test_a_link_is_downloaded_stored_and_never_twice(self, m):
        data = book(products=[{"name": "המבורגר קלאסי", "category": "המבורגרים",
                               "image": "https://cdn.example.com/burger.jpg"}],
                    categories=[{"name": "שתייה", "image": "https://cdn.example.com/drinks.png"}])
        out = preview(m, data)
        assert out["products"][0]["changes"][0]["after"] == "תורד מ-cdn.example.com/burger.jpg"
        assert m.downloads == []  # the preview downloads nothing
        result = commit(m, data)
        assert result["imagesStored"] == 2 and result["imageFailures"] == []
        url = m.db.get(Product, m.classic.id).image_url
        assert url.startswith("http://localhost:8001/media/pos/") and IMG.marker_of("https://cdn.example.com/burger.jpg") in url
        assert (m.media / url.split("/media/", 1)[1]).is_file()
        assert "/categories/" in m.db.get(Category, m.drinks.id).image_url
        again = preview(m, data)
        assert again["summary"]["imageChanges"] == 0 and again["canCommit"] is False
        commit(m, data)
        assert len(m.downloads) == 2

    def test_a_broken_link_is_reported_and_the_rest_imports(self, m):
        result = commit(m, book(products=[{"name": "קולה", "category": "שתייה", "price": 13,
                                           "image": "https://cdn.example.com/broken.jpg"}]))
        assert result["productsUpdated"] == 1 and result["imagesStored"] == 0
        assert result["imageFailures"][0]["row"] == 2
        assert "השרת החזיר שגיאה (404)" in result["imageFailures"][0]["message"]
        cola = m.db.get(Product, m.cola.id)
        assert cola.price == Decimal("13.00") and cola.image_url is None

    def test_our_own_media_is_set_as_is_and_lelo_removes(self, m):
        own = "http://localhost:8001/media/pos/x/products/abc.png"
        commit(m, book(products=[{"name": "קולה", "category": "שתייה", "image": own}]))
        assert m.db.get(Product, m.cola.id).image_url == own and m.downloads == []
        out = commit(m, book(products=[{"name": "קולה", "category": "שתייה", "image": "ללא"}]))
        assert out["imagesRemoved"] == 1 and m.db.get(Product, m.cola.id).image_url is None

    def test_a_bad_link_is_a_row_error(self, m):
        out = preview(m, book(products=[{"name": "קולה", "category": "שתייה", "image": "תמונה של קולה"}]))
        assert out["products"][0]["status"] == "error"
        assert "הקישור לתמונה לא תקין" in out["products"][0]["messages"][0]["text"]

    def test_the_fetcher_refuses_internal_addresses_and_non_pictures(self):
        for url in ("http://127.0.0.1/a.png", "http://10.1.2.3/a.png", "http://169.254.169.254/latest", "ftp://x/a.png"):
            with pytest.raises(IMG.ImageError):
                IMG._check_host(url)
        with pytest.raises(IMG.ImageError):
            IMG.verify(b"<html><body>not a picture</body></html>")
        assert IMG.verify(_png()) == "PNG"
        assert IMG.direct_url("https://drive.google.com/file/d/1AbCdEfGhIjKlMn/view?usp=sharing") == \
            "https://drive.google.com/uc?export=download&id=1AbCdEfGhIjKlMn"
        assert IMG.direct_url("https://www.dropbox.com/s/abc/burger.jpg?dl=0") == \
            "https://www.dropbox.com/s/abc/burger.jpg?raw=1"


# ── Menu prices ───────────────────────────────────────────────────────────────


def menu_rows(world):
    rows = world.db.query(CatalogMenuProduct).filter(CatalogMenuProduct.menu_id == world.lunch.id).all()
    return {r.product_id: r.price for r in rows}


class TestMenuPrices:
    def test_a_menu_price_is_updated_added_and_cleared(self, m):
        data = book(products=[
            {"name": "קולה", "category": "שתייה", "menu:צהריים": 11},
            {"name": "מים", "category": "שתייה", "menu:צהריים": "7"},
            {"name": "המבורגר קלאסי", "category": "המבורגרים", "menu:צהריים": 50},
        ], menus=["צהריים"])
        out = preview(m, data)
        texts = all_messages(out)
        assert "שורה 3: הפריט יתווסף לרשימת הפריטים של התפריט 'צהריים' עם המחיר הזה" in texts
        assert "שורה 4: הפריט יתווסף לתפריט 'צהריים' וימכר בו גם אם המחלקה שלו לא בתפריט" in texts
        assert out["products"][0]["changes"] == [{"field": "menu_price", "label": "מחיר בתפריט 'צהריים'",
                                                  "before": "10.00", "after": "11.00"}]
        result = commit(m, data)
        assert result["menuPricesChanged"] == 3
        assert menu_rows(m) == {m.cola.id: Decimal("11.00"), m.water.id: Decimal("7.00"), m.classic.id: Decimal("50.00")}
        assert m.db.query(CatalogMenuSyncState).filter(CatalogMenuSyncState.tenant_id == m.tenant.id).count() == 1
        # The base price is the catalog's, untouched.
        assert m.db.get(Product, m.cola.id).price == Decimal("12.00")
        commit(m, book(products=[{"name": "קולה", "category": "שתייה", "menu:צהריים": "ללא"}], menus=["צהריים"]))
        assert menu_rows(m)[m.cola.id] is None  # back to the catalog's price, still in the menu
        assert preview(m, data)["summary"]["menuPriceChanges"] == 1

    def test_an_unknown_menu_column_is_warned_and_skipped(self, m):
        out = preview(m, book(products=[{"name": "קולה", "category": "שתייה", "menu:ערב": 9}], menus=["ערב"]))
        assert "העמודה 'מחיר בתפריט: ערב': אין תפריט בשם 'ערב' - העמודה תידלג" in [i["text"] for i in out["issues"]]
        assert out["products"][0]["status"] == "unchanged"

    def test_a_company_manager_cannot_price_the_organizations_menu(self, m):
        m.lunch.company_id = None
        m.db.commit()
        data = book(products=[{"name": "קולה", "category": "שתייה", "menu:צהריים": 11, "groups": "מידת עשייה"}],
                    menus=["צהריים"])
        out = preview(m, data, user=m.manager)
        assert out["products"][0]["status"] == "error"
        assert "אין לך הרשאה לשנות את התפריט 'צהריים'" in messages(out["products"])[-1]
        # The refused row writes nothing - not even the groups planned before the refusal.
        commit(m, data, user=m.manager, skip_errors=True)
        assert own_groups(m, "product", m.cola.id) == [] and menu_rows(m)[m.cola.id] == Decimal("10.00")


def test_the_plan_binds_the_menu_layer_to_the_token(m):
    data = book(groups=[{"name": "מידת עשייה", "sort": 3}])
    token = preview(m, data)["token"]
    m.db.get(ModifierGroup, m.doneness.id).sort_order = 3
    m.db.commit()
    with pytest.raises(Exception) as caught:
        commit(m, data, token=token)
    assert caught.value.detail["code"] == "plan_changed"
    assert I.build_plan(I.load_context(m.db, m.tenant.id, m.company, m.distributor),
                        S.read_file(data, "x.xlsx")).menu.groups[0].action == "unchanged"
