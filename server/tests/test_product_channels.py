"""
"היכן הפריט נמכר" (docs/SPEC_PRODUCT_CHANNELS.md): the code is validated on every way in,
defaults to קופות וקיוסק, is edited from the dashboard, the till and the Excel sheet, and
reaches every till and kiosk on both catalog serialisers. Also the till's picture of a
product: uploaded and removed on the till's own credentials, with a manager's authority,
like every other product edit from a till.
"""
from __future__ import annotations

import asyncio
import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile
from pydantic import ValidationError
from starlette.datastructures import Headers

from app.database import Base
from app.middleware.auth import CatalogActor, get_pos_machine_for_sync_path
from app.models.product import Product
from app.routers import images as images_router
from app.routers import sync as sync_router
from app.schemas.product import ProductCreate, ProductResponse, ProductUpdate
from app.services import catalog_sheet as S
from app.services import sales_channel as C
from app.services.permissions import Scope
from app.services.sync import _serialize_merged_product, _serialize_product
from test_catalog_import import commit, preview, products_named, template, w, workbook  # noqa: F401
from test_product_availability import world  # noqa: F401
from test_product_weighed import _product


# ── The rule ──────────────────────────────────────────────────────────────────


class TestCodes:
    def test_three_codes_all_first(self) -> None:
        assert C.SALES_CHANNELS == ("all", "kiosk_only", "pos_only")
        assert set(C.SALES_CHANNEL_LABELS_HE) == set(C.SALES_CHANNELS)
        assert C.SALES_CHANNEL_LABELS_HE == {
            "all": "קופות וקיוסק", "kiosk_only": "קיוסק בלבד", "pos_only": "קופות בלבד",
        }

    def test_who_sells_what(self) -> None:
        assert (C.on_pos("all"), C.on_kiosk("all")) == (True, True)
        assert (C.on_pos("kiosk_only"), C.on_kiosk("kiosk_only")) == (False, True)
        assert (C.on_pos("pos_only"), C.on_kiosk("pos_only")) == (True, False)

    def test_reading_back_never_hides_a_product_from_everywhere(self) -> None:
        # Missing (a row before the column) or a stray value: sold everywhere, as before.
        for stored in (None, "", "web", 7, "KIOSK_ONLY"):
            assert C.out(stored) == "all"
            assert C.on_pos(stored) and C.on_kiosk(stored)

    def test_hebrew_labels_codes_and_aliases(self) -> None:
        assert [C.code_of(v) for v in ("קיוסק בלבד", " קופות  בלבד ", "קופות וקיוסק", "Kiosk_Only", "pos")] == [
            "kiosk_only", "pos_only", "all", "kiosk_only", "pos_only",
        ]
        assert C.code_of("אונליין") is None and C.code_of("") is None and C.code_of(None) is None


# ── The API's schemas ─────────────────────────────────────────────────────────

_BASE = {"name": "במבה", "price": "1.00", "sku": "K-1", "categoryId": str(uuid.uuid4())}


class TestSchemas:
    def test_a_new_product_is_sold_everywhere_unless_told(self) -> None:
        assert ProductCreate.model_validate(_BASE).sales_channel == "all"
        for code in C.SALES_CHANNELS:
            assert ProductCreate.model_validate({**_BASE, "salesChannel": code}).sales_channel == code

    @pytest.mark.parametrize("bad", ["kiosk", "web", "", "קיוסק בלבד", 1])
    def test_anything_else_is_refused(self, bad) -> None:
        with pytest.raises(ValidationError):
            ProductCreate.model_validate({**_BASE, "salesChannel": bad})
        with pytest.raises(ValidationError):
            ProductUpdate.model_validate({"salesChannel": bad})

    def test_an_update_leaves_it_alone_unless_sent_and_null_is_not_sent(self) -> None:
        assert "sales_channel" not in ProductUpdate.model_validate({"price": "2"}).model_dump(exclude_unset=True)
        # The column is NOT NULL: null means "leave it", never a null written into it.
        assert ProductUpdate.model_validate({"salesChannel": None, "price": "2"}).model_dump(exclude_unset=True) == {
            "price": Decimal("2"),
        }
        assert ProductUpdate.model_validate({"salesChannel": "pos_only"}).model_dump(exclude_unset=True) == {
            "sales_channel": "pos_only",
        }

    def test_the_response_always_has_a_code(self) -> None:
        def row(channel):
            now = datetime.now(timezone.utc)
            return SimpleNamespace(
                id=uuid.uuid4(), tenant_id=None, company_id=None, shop_id=None, pos_machine_id=None,
                global_product_id=None, catalog_level="global", is_local_override=False, name="x",
                description=None, price=1, sku="1", global_sku=None, sku_auto_assigned=False,
                category_id=uuid.uuid4(), image_url=None, in_stock=True, is_available=True,
                stock_quantity=0, barcode=None, tax_rate=None, voucher_id=None, dietary_tags=None,
                alerts=None, companions=None, allergen_alert=None, allergen_alert_require_ack=None,
                sales_channel=channel, created_at=now, updated_at=now,
            )

        assert ProductResponse.model_validate(row(None)).model_dump(by_alias=True)["salesChannel"] == "all"
        assert ProductResponse.model_validate(row("kiosk_only")).model_dump(by_alias=True)["salesChannel"] == "kiosk_only"


# ── The tills' and kiosks' catalog ────────────────────────────────────────────


class TestSyncPayload:
    def test_both_serialisers_ship_it(self) -> None:
        p = _product(is_weighed=False, unit_label=None)
        p.sales_channel = "kiosk_only"
        for row in (_serialize_product(p), _serialize_merged_product(p, None, None, uuid.uuid4(), None)):
            assert row["salesChannel"] == "kiosk_only"

    def test_a_row_without_one_is_sold_everywhere(self) -> None:
        p = _product(is_weighed=False, unit_label=None)
        p.sales_channel = None
        assert _serialize_product(p)["salesChannel"] == "all"
        assert _serialize_merged_product(p, None, None, uuid.uuid4(), None)["salesChannel"] == "all"

    def test_it_is_taken_from_the_global_row(self) -> None:
        global_p = _product(is_weighed=False, unit_label=None)
        global_p.sales_channel = "pos_only"
        local = _product(is_weighed=False, unit_label=None)
        local.sales_channel = "kiosk_only"
        assert _serialize_merged_product(global_p, local, None, uuid.uuid4(), None)["salesChannel"] == "pos_only"


def test_the_dashboard_router_stores_edits_and_ships_them(monkeypatch) -> None:
    from shift_world import accept_str_uuids, make_world

    from app.models.category import Category
    from app.routers import products as R

    accept_str_uuids(monkeypatch)
    world = make_world()
    monkeypatch.setattr(R, "_trigger_catalog_notify", lambda *a, **k: None)
    monkeypatch.setattr(R, "allocate_global_sku", lambda *a, **k: "100001")
    monkeypatch.setattr(R, "resolve_sku_for_create", lambda db, tenant, sku: (sku, False))
    cat = Category(id=uuid.uuid4(), tenant_id=world.tenant.id, company_id=world.company.id, name="חטיפים")
    world.db.add(cat)
    world.db.commit()

    def update(body):
        R.update_product(
            str(created.id), ProductUpdate.model_validate(body),
            current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
        )
        return world.db.get(Product, created.id)

    created = R.create_product(
        ProductCreate.model_validate({
            "name": "במבה", "price": "1.00", "sku": "BM-1", "categoryId": str(cat.id),
            "companyId": str(world.company.id), "salesChannel": "kiosk_only",
        }),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    assert world.db.get(Product, created.id).sales_channel == "kiosk_only"
    assert ProductResponse.model_validate(created).model_dump(by_alias=True)["salesChannel"] == "kiosk_only"
    # Omitted, or null: left alone.
    assert update({"price": "2.00"}).sales_channel == "kiosk_only"
    assert update({"salesChannel": None}).sales_channel == "kiosk_only"
    # Changed, then the row a till pulls.
    assert update({"salesChannel": "pos_only"}).sales_channel == "pos_only"
    assert _serialize_product(world.db.get(Product, created.id))["salesChannel"] == "pos_only"

    # A product made without one is sold everywhere.
    plain = R.create_product(
        ProductCreate.model_validate({
            "name": "קולה", "price": "8.00", "sku": "CL-1", "categoryId": str(cat.id),
            "companyId": str(world.company.id),
        }),
        current_user=world.admin, active_tenant_id=world.tenant.id, db=world.db,
    )
    assert world.db.get(Product, plain.id).sales_channel == "all"


# ── From the till ─────────────────────────────────────────────────────────────


@pytest.fixture
def tw(world, monkeypatch):  # noqa: F811
    """The availability world: P is sold in h_shop and a_shop, Q in h_shop only."""
    Base.metadata.tables["sync_logs"].create(world.db.get_bind())
    world.woken = []
    monkeypatch.setattr(sync_router, "notify_all_machines_for_tenant", lambda *a, **k: world.woken.append("tenant"))
    monkeypatch.setattr(sync_router, "notify_machines_for_shop", lambda *a, **k: world.woken.append("shop"))
    skus = iter(range(1, 100))
    monkeypatch.setattr(sync_router, "resolve_sku_for_create", lambda db, tenant, sku: (sku or f"T-{next(skus)}", sku is None))
    monkeypatch.setattr(sync_router, "allocate_global_sku", lambda *a, **k: "900001")
    return world


def _actor():
    return CatalogActor(pos_user_id=uuid.uuid4())


class TestFromTheTill:
    def test_a_product_made_on_the_till_carries_its_channel(self, tw) -> None:
        body = ProductCreate.model_validate({
            "name": "מאפין", "price": "1.00", "categoryId": str(tw.P.category_id), "salesChannel": "kiosk_only",
        })
        out = sync_router.machine_create_cloud_product(str(tw.h1.id), body, machine=tw.h1, actor=_actor(), db=tw.db)
        assert tw.db.get(Product, out.id).sales_channel == "kiosk_only"
        # And without one, sold everywhere.
        body = ProductCreate.model_validate({"name": "תה", "price": "6.00", "categoryId": str(tw.P.category_id)})
        out = sync_router.machine_create_cloud_product(str(tw.h1.id), body, machine=tw.h1, actor=_actor(), db=tw.db)
        assert tw.db.get(Product, out.id).sales_channel == "all"

    def test_a_shops_own_product_is_changed_from_its_till(self, tw) -> None:
        sync_router.machine_update_cloud_product(
            str(tw.h1.id), str(tw.Q.id), ProductUpdate.model_validate({"salesChannel": "pos_only"}),
            machine=tw.h1, actor=_actor(), db=tw.db,
        )
        assert tw.db.get(Product, tw.Q.id).sales_channel == "pos_only"
        assert "tenant" in tw.woken

    def test_a_product_other_shops_sell_is_the_chains_and_refused(self, tw) -> None:
        with pytest.raises(HTTPException) as e:
            sync_router.machine_update_cloud_product(
                str(tw.h1.id), str(tw.P.id), ProductUpdate.model_validate({"salesChannel": "kiosk_only"}),
                machine=tw.h1, actor=_actor(), db=tw.db,
            )
        assert (e.value.status_code, e.value.detail) == (403, "shared_product_master_readonly")
        assert tw.db.get(Product, tw.P.id).sales_channel == "all"


def _upload(tw, product, content_type="image/jpeg", data=b"\xff\xd8\xff jpeg", keep=True):
    f = UploadFile(file=io.BytesIO(data), filename="p.jpg", headers=Headers({"content-type": content_type}))
    return asyncio.run(sync_router.machine_upload_product_image(
        str(tw.h1.id), str(product.id), file=f, keep_background=keep, machine=tw.h1, actor=_actor(), db=tw.db,
    ))


def _remove(tw, product):
    return sync_router.machine_remove_product_image(str(tw.h1.id), str(product.id), machine=tw.h1, actor=_actor(), db=tw.db)


class TestTillPicture:
    def test_the_till_uploads_with_the_dashboards_validation_and_storage(self, tw, monkeypatch) -> None:
        seen = {}

        async def fake_store(contents, tenant_id, resource, keep_background, enhance=False):
            seen.update(resource=resource, keep=keep_background, tenant=tenant_id, size=len(contents), enhance=enhance)
            return images_router.ImageUploadResponse(url="https://img/p.jpg", publicId="p1")

        monkeypatch.setattr(images_router, "store_upload", fake_store)
        out = _upload(tw, tw.Q)
        assert out == {"url": "https://img/p.jpg", "originalUrl": None, "backgroundRemoved": False, "processed": False}
        assert seen == {"resource": "products", "keep": True, "tenant": tw.tid, "size": 8, "enhance": True}
        assert tw.db.get(Product, tw.Q.id).image_url == "https://img/p.jpg"
        # Every till (and kiosk) of the tenant is told to pull.
        assert "tenant" in tw.woken

    def test_not_an_image_and_too_big_are_refused_before_storing(self, tw, monkeypatch) -> None:
        async def never(*a, **k):
            raise AssertionError("stored")

        monkeypatch.setattr(images_router, "store_upload", never)
        with pytest.raises(HTTPException) as e:
            _upload(tw, tw.Q, content_type="application/pdf")
        assert e.value.status_code == 415
        with pytest.raises(HTTPException) as e:
            _upload(tw, tw.Q, data=b"x" * (images_router._MAX_SIZE_BYTES + 1))
        assert e.value.status_code == 413
        assert tw.db.get(Product, tw.Q.id).image_url is None

    def test_removing_it_clears_the_picture_and_tells_the_tills(self, tw) -> None:
        tw.Q.image_url = "https://img/old.jpg"
        tw.db.commit()
        assert _remove(tw, tw.Q) == {"url": None}
        assert tw.db.get(Product, tw.Q.id).image_url is None
        assert tw.woken == ["tenant"]
        # Again: nothing to remove, still 200 (a queued removal replayed), nobody woken.
        assert _remove(tw, tw.Q) == {"url": None}
        assert tw.woken == ["tenant"]

    def test_a_product_other_shops_sell_keeps_its_picture(self, tw, monkeypatch) -> None:
        tw.P.image_url = "https://img/chain.jpg"
        tw.db.commit()
        for call in (lambda: _remove(tw, tw.P), lambda: _upload(tw, tw.P)):
            with pytest.raises(HTTPException) as e:
                call()
            assert (e.value.status_code, e.value.detail) == (403, "shared_product_master_readonly")
        assert tw.db.get(Product, tw.P.id).image_url == "https://img/chain.jpg"

    def test_another_shops_till_cannot_touch_it(self, tw) -> None:
        # b1 stands in another company's shop: Q is not listed there.
        with pytest.raises(HTTPException) as e:
            sync_router.machine_remove_product_image(str(tw.b1.id), str(tw.Q.id), machine=tw.b1, actor=_actor(), db=tw.db)
        assert e.value.status_code == 403


def _route(path: str, method: str):
    for route in sync_router.router.routes:
        if getattr(route, "path", None) == path and method in getattr(route, "methods", ()):
            return route
    raise AssertionError(f"{method} {path} is not a route")


def _closure_values(fn):
    return [c.cell_contents for c in (fn.__closure__ or ())]


@pytest.mark.parametrize(
    "method, path",
    [
        ("POST", "/sync/{machine_id}/products"),
        ("PUT", "/sync/{machine_id}/products/{product_id}"),
        ("POST", "/sync/{machine_id}/products/{product_id}/image"),
        ("DELETE", "/sync/{machine_id}/products/{product_id}/image"),
    ],
)
def test_every_product_write_from_a_till_needs_its_machine_token_and_a_managers_authority(method, path) -> None:
    """
    The till calls these with its own machine credentials (the machine JWT for this
    `machine_id`), and the write needs `catalog:write` — a signed-in shop manager, or a
    manager's grant; a cashier gets 401 `elevation_required` (tests/test_till_user_authority.py).
    """
    deps = [d.call for d in _route(path, method).dependant.dependencies]
    assert get_pos_machine_for_sync_path in deps
    authority = [d for d in deps if getattr(d, "__qualname__", "").startswith("require_catalog_authority")]
    assert len(authority) == 1
    assert Scope.CATALOG_WRITE in _closure_values(authority[0])


def test_a_cashier_at_the_till_is_sent_to_find_a_manager(monkeypatch) -> None:
    from test_till_user_authority import _FakeDb, _catalog_authority, _machine, _pos_user
    from app.models.pos_user import PosUserRole

    machine = _machine()
    cashier = _pos_user(machine.shop_id, role=PosUserRole.CASHIER, username="yossi")
    with pytest.raises(HTTPException) as e:
        _catalog_authority(_FakeDb(cashier), machine, operator_id=str(cashier.id))
    assert (e.value.status_code, e.value.detail) == (401, "elevation_required")
    manager = _pos_user(machine.shop_id)
    assert _catalog_authority(_FakeDb(manager), machine, operator_id=str(manager.id)).pos_user_id == manager.id


# ── The Excel sheet ───────────────────────────────────────────────────────────


class TestSheetCell:
    def test_labels_and_codes(self) -> None:
        assert S.parse_channel("קיוסק בלבד").value == "kiosk_only"
        assert S.parse_channel("pos_only").value == "pos_only"
        assert S.parse_channel(" קופות וקיוסק ").value == "all"

    def test_empty_is_no_change_and_unknown_is_a_row_error(self) -> None:
        assert S.parse_channel(None).value is None and S.parse_channel("  ").value is None
        assert "ערוץ מכירה לא מוכר ('אונליין')" in S.parse_channel("אונליין").error

    def test_the_column(self) -> None:
        col = next(c for c in S.PRODUCT_COLUMNS if c.key == "channel")
        assert col.header == "ערוץ מכירה" and not col.required


class TestImportExport:
    def test_a_new_product_gets_its_channel_and_empty_is_everywhere(self, w) -> None:
        commit(w, workbook([
            {"name": "במבה", "category": "שתייה", "price": 1, "channel": "קיוסק בלבד"},
            {"name": "תה", "category": "שתייה", "price": 6},
        ]))
        assert products_named(w, "במבה")[0].sales_channel == "kiosk_only"
        assert products_named(w, "תה")[0].sales_channel == "all"

    def test_an_update_shows_the_change_in_hebrew_and_empty_leaves_it(self, w) -> None:
        out = preview(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "channel": "קופות בלבד"}]))
        change = next(c for c in out["products"][0]["changes"] if c["field"] == "channel")
        assert change["label"] == "ערוץ מכירה"
        assert (change["before"], change["after"]) == ("קופות וקיוסק", "קופות בלבד")
        commit(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "channel": "קופות בלבד"}]))
        w.db.refresh(w.cola)
        assert w.cola.sales_channel == "pos_only"
        commit(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12}]))
        w.db.refresh(w.cola)
        assert w.cola.sales_channel == "pos_only"

    def test_an_unknown_channel_is_an_error_row(self, w) -> None:
        out = preview(w, workbook([{"name": "קולה", "category": "שתייה", "price": 12, "channel": "אונליין"}]))
        assert out["products"][0]["status"] == "error"

    def test_the_export_writes_hebrew_and_reads_back_as_no_change(self, w) -> None:
        w.cola.sales_channel = "kiosk_only"
        w.db.commit()
        data = template(w, with_data=True)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[S.SHEET_PRODUCTS]
        col = [c.key for c in S.PRODUCT_COLUMNS].index("channel") + 1
        assert ws.cell(1, col).value == "ערוץ מכירה"
        cells = {ws.cell(r, 2).value: ws.cell(r, col).value for r in range(2, 8) if ws.cell(r, 2).value}
        assert cells["קולה"] == "קיוסק בלבד" and cells["מים"] == "קופות וקיוסק"
        # The column offers the three labels as a dropdown.
        letter = openpyxl.utils.get_column_letter(col)
        assert any(dv.type == "list" and str(dv.sqref).startswith(f"{letter}2:") for dv in ws.data_validations.dataValidation)
        out = preview(w, data)
        assert out["summary"]["productsUnchanged"] == 3 and out["summary"]["errors"] == 0


# ── The menu-broadcast review ─────────────────────────────────────────────────


def test_the_broadcast_review_sees_a_channel_change_but_not_an_older_snapshot() -> None:
    from app.services import menu_broadcast as MB

    def snap(**fields):
        return {"products": {"p1": {"name": "במבה", "price": 1, "shopListed": True, **fields}}}

    changed = MB.diff(snap(salesChannel="all"), snap(salesChannel="kiosk_only"))["products"]
    assert [c["field"] for c in changed[0]["changes"]] == ["salesChannel"]
    assert MB.diff(snap(), snap(salesChannel="all"))["products"] == []
