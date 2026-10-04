"""
The journal as an xlsx (spec §2.5): fixed columns for Priority's load interface, and the
human-readable copy that goes in every export zip.

Columns: תאריך · אסמכתא · אסמכתא 2 · חשבון · חובה · זכות · פרטים · סניף · סוג תנועה.
Right-to-left sheet, real dates and numbers (never text), a totals row.

!! Column names must be checked with the client's Priority implementer (spec §7).
"""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.accounting.journal import JournalEntry

COLUMNS = (
    ("תאריך", 12),
    ("אסמכתא", 10),
    ("אסמכתא 2", 10),
    ("חשבון", 14),
    ("חובה", 14),
    ("זכות", 14),
    ("פרטים", 40),
    ("סניף", 8),
    ("סוג תנועה", 10),
)
MONEY = "#,##0.00"


def write_xlsx(entries: Sequence[JournalEntry], *, title: str = "פקודות יומן") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    ws.sheet_view.rightToLeft = True

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E5E7EB")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    row = 2
    total_d = total_c = Decimal("0")
    for entry in entries:
        for line in entry.lines:
            debit = line.amount if line.side == "D" else None
            credit = line.amount if line.side == "C" else None
            values = (
                entry.entry_date,
                entry.reference1,
                entry.reference2 or None,
                line.account,
                float(debit) if debit is not None else None,
                float(credit) if credit is not None else None,
                f"{entry.details} {line.label}".strip(),
                (line.branch if line.branch is not None else entry.branch) or None,
                entry.movement_type or None,
            )
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=1).number_format = "dd/mm/yyyy"
            ws.cell(row=row, column=5).number_format = MONEY
            ws.cell(row=row, column=6).number_format = MONEY
            total_d += debit or 0
            total_c += credit or 0
            row += 1

    ws.cell(row=row, column=4, value='סה"כ').font = bold
    for col, total in ((5, total_d), (6, total_c)):
        cell = ws.cell(row=row, column=col, value=float(total))
        cell.font = bold
        cell.number_format = MONEY

    wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
