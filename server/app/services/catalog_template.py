"""
Drawing the menu workbook ("טמפלט קליטת פריטים"): the blank template a dealer sends a
customer, and the same workbook filled with a company's current catalog for editing.

Only drawing. What goes in it - the company's categories, products, add-on groups and
their options, quick notes, catalog menus, printers and routing - is read from the
database by app/services/catalog_import.py (`template_view`), and the columns are
app/services/catalog_sheet.py's, so a downloaded file always reads back.

Every sheet is right-to-left with a styled, frozen header; required headers end with "*";
each header carries a note that explains it; the dropdowns (כן/לא, יחידה, מחלקות, מדפסות,
קבוצות תוספות, סוג קבוצה, שובר פריט, פריטים) and number / link checks are Excel data
validations; the grey rows marked "דוגמה" - one or more on every sheet - are skipped by
the import.
"""
from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.catalog_sheet import (
    CATEGORY_COLUMNS,
    CLEAR_WORD,
    COST_MAX,
    DIETARY_LABELS_HE,
    DIETARY_NONE,
    ENTRIES_MAX,
    EXAMPLE_MARK,
    GROUP_COLUMNS,
    GROUP_KIND_LABELS,
    GROUPS_INHERIT,
    GROUPS_NONE,
    LIMIT_NONE,
    MENU_PRICE_MAX,
    MENU_PRICE_PREFIX,
    NO,
    NOTE_COLUMNS,
    OPTION_COLUMNS,
    OPTION_PRICE_MAX,
    OPTION_PRICE_MIN,
    PRINTER_REFERENCE_HEADERS,
    PRINTERS_INHERIT,
    PRINTERS_NONE,
    PRODUCT_COLUMNS,
    SHEET_CATEGORIES,
    SHEET_GROUPS,
    SHEET_INSTRUCTIONS,
    SHEET_LISTS,
    SHEET_NOTES,
    SHEET_OPTIONS,
    SHEET_PRINTERS,
    SHEET_PRODUCTS,
    SKIP_MARK,
    SORT_MAX,
    SORT_MIN,
    TICKET_INHERIT,
    TICKET_LABELS,
    UNIT_CHOICES,
    UNIT_KG,
    UNIT_PIECE,
    YES,
    Column,
    menu_price_column,
)
from app.services.sales_channel import SALES_CHANNEL_LABELS_HE

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# iOS system colours, as the dashboard's pages use them.
BLUE = "007AFF"
GREY = "8E8E93"
INK = "1C1C1E"
GROUPED = "F2F2F7"
HAIRLINE = "D1D1D6"
GREEN = "34C759"
ORANGE = "FF9500"
PURPLE = "AF52DE"
PINK = "FF2D55"
TEAL = "30B0C7"

_HEADER_REQUIRED = PatternFill("solid", fgColor=BLUE)
_HEADER_OPTIONAL = PatternFill("solid", fgColor="5E5CE6")
_HEADER_MARKER = PatternFill("solid", fgColor="C7C7CC")
_EXAMPLE_FILL = PatternFill("solid", fgColor=GROUPED)
_EXAMPLE_FONT = Font(color=GREY, italic=True)
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_DARK_FONT = Font(color=INK, bold=True)
_THIN = Side(style="thin", color=HAIRLINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
_RIGHT = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)

_NUMBER_FORMATS = {"money": "#,##0.00", "signed_money": "#,##0.00", "menu_price": "#,##0.00", "code": "@", "int": "0"}

#: Empty rows below the data that still get the dropdowns and the text format.
SPARE_ROWS = 500
MIN_ROWS = 1000


@dataclass
class PrinterRef:
    name: str
    shop: str
    connection: str
    narrowed: str
    active: bool


@dataclass
class TemplateView:
    company_name: str
    generated_at: datetime
    #: The current catalog (True) or a blank template (False; examples, no products).
    with_data: bool
    #: Rows as {column key: value}; `notes` holds {column key: cell note}.
    categories: List[Dict[str, Any]] = field(default_factory=list)
    products: List[Dict[str, Any]] = field(default_factory=list)
    printers: List[PrinterRef] = field(default_factory=list)
    #: The printer names a "מדפסות" cell may hold, deduplicated.
    printer_choices: List[str] = field(default_factory=list)
    #: The routing of some rows differs between shops and was left blank.
    has_mixed_routing: bool = False
    #: The add-on groups and their options (pre-filled in the blank template too, like the
    #: categories), and - with data - the quick notes.
    groups: List[Dict[str, Any]] = field(default_factory=list)
    options: List[Dict[str, Any]] = field(default_factory=list)
    notes: List[Dict[str, Any]] = field(default_factory=list)
    #: The catalog menus the company sees: a "מחיר בתפריט: <name>" column each.
    menu_names: List[str] = field(default_factory=list)

    def product_columns(self) -> Tuple[Column, ...]:
        return PRODUCT_COLUMNS + tuple(menu_price_column(name) for name in self.menu_names)


# ── Helpers ───────────────────────────────────────────────────────────────────


def _rtl(ws) -> None:
    ws.sheet_view.rightToLeft = True


def _quoted(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _list_validation(formula: str, *, strict: bool, prompt: Optional[str] = None,
                     error: Optional[str] = None, free_text: bool = False) -> DataValidation:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if free_text:
        # The dropdown helps; anything else (several printers, a new name) is still allowed.
        dv.showErrorMessage = False
    else:
        dv.showErrorMessage = True
        dv.errorStyle = "stop" if strict else "warning"
        dv.errorTitle = "ערך לא מהרשימה"
        dv.error = error or "בחרו ערך מהרשימה"
    if prompt:
        dv.showInputMessage = True
        dv.prompt = prompt[:250]
    return dv


def _number_validation(kind: str, low, high, error: str) -> DataValidation:
    dv = DataValidation(type=kind, operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "stop"
    dv.errorTitle = "מספר לא תקין"
    dv.error = error
    return dv


def _length_validation(maximum: int) -> DataValidation:
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(maximum), allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "stop"
    dv.errorTitle = "טקסט ארוך מדי"
    dv.error = f"עד {maximum} תווים"
    return dv


def _custom_validation(formula: str, title: str, error: str, *, strict: bool = True) -> DataValidation:
    """A formula check, written for the column's first data cell (Excel shifts it per row)."""
    dv = DataValidation(type="custom", formula1=formula, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "stop" if strict else "warning"
    dv.errorTitle = title
    dv.error = error
    return dv


def _header(ws, columns: Sequence[Column], freeze: str = "C2") -> None:
    for index, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=col.header)
        if col.kind == "marker":
            cell.fill, cell.font = _HEADER_MARKER, _HEADER_DARK_FONT
        else:
            cell.fill = _HEADER_REQUIRED if col.required else _HEADER_OPTIONAL
            cell.font = _HEADER_FONT
        cell.alignment = _WRAP_CENTER
        cell.border = _BOX
        note = Comment(col.help + (f"\nלדוגמה: {col.example}" if col.example else ""), "POS")
        note.width, note.height = 260, 120
        cell.comment = note
        ws.column_dimensions[get_column_letter(index)].width = col.width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = freeze


def _write_rows(ws, columns: Sequence[Column], rows: Sequence[Dict[str, Any]], start: int, *, example: bool) -> int:
    """Write `rows` from row `start`; the next free row."""
    row_number = start
    for row in rows:
        notes = row.get("notes") or {}
        for index, col in enumerate(columns, start=1):
            value = row.get(col.key)
            cell = ws.cell(row=row_number, column=index, value=value if value not in ("",) else None)
            if col.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[col.kind]
            if example:
                cell.fill, cell.font = _EXAMPLE_FILL, _EXAMPLE_FONT
            if col.key in notes:
                note = Comment(notes[col.key], "POS")
                note.width, note.height = 240, 90
                cell.comment = note
        row_number += 1
    return row_number


def _column_formats(ws, columns: Sequence[Column]) -> None:
    """
    Each column's format for the cells typed into it later - text for a barcode / SKU, so
    0123 stays 0123 - set on the column itself. Formatting the empty cells instead would
    stretch the sheet's used range, and a row added below it would sit past the rows read.
    """
    for index, col in enumerate(columns, start=1):
        if col.kind in _NUMBER_FORMATS:
            ws.column_dimensions[get_column_letter(index)].number_format = _NUMBER_FORMATS[col.kind]


def _page_setup(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ── The data sheets ───────────────────────────────────────────────────────────


@dataclass
class _Refs:
    """The ranges the dropdowns list: the hidden lists, and the names typed in other sheets."""

    lists: Dict[str, str]
    categories: str
    groups: str
    products: str


def _validations(ws, columns: Sequence[Column], last_row: int, refs: _Refs) -> None:
    letter = {col.key: get_column_letter(i) for i, col in enumerate(columns, start=1)}
    lists = refs.lists

    def apply(dv: DataValidation, key: str) -> None:
        ws.add_data_validation(dv)
        dv.add(f"{letter[key]}2:{letter[key]}{last_row}")

    yes_no = _list_validation(lists["yes_no"], strict=True, error="כתבו כן או לא")
    ws.add_data_validation(yes_no)
    for col in columns:
        key = col.key
        first = f"{letter[key]}2"
        if col.kind == "bool":
            yes_no.add(f"{letter[key]}2:{letter[key]}{last_row}")
        elif col.kind == "marker":
            apply(_list_validation(f'"{EXAMPLE_MARK},{SKIP_MARK}"', strict=True,
                                   error="השאירו ריק, או 'דוגמה' / 'דלג' לשורה שלא תיקלט"), key)
        elif col.kind == "money":
            high = COST_MAX if key == "cost" else 99999999
            apply(_number_validation("decimal", 0, high, "הקלידו מספר (למשל 12.50), בלי טקסט"), key)
        elif col.kind == "signed_money":
            apply(_number_validation("decimal", OPTION_PRICE_MIN, OPTION_PRICE_MAX,
                                     f"הקלידו מספר בין {OPTION_PRICE_MIN} ל-{OPTION_PRICE_MAX} (למשל 4 או 2.50)"), key)
        elif col.kind == "menu_price":
            # IF, not AND: Excel's OR / AND pass an error on, and a word is no number.
            apply(_custom_validation(
                f'OR(ISBLANK({first}),{first}="{CLEAR_WORD}",'
                f'IF(ISNUMBER({first}),AND({first}>=0,{first}<={MENU_PRICE_MAX}),FALSE))',
                "מחיר לא תקין", f"הקלידו מספר (למשל 48), או '{CLEAR_WORD}' לביטול המחיר בתפריט"), key)
        elif col.kind == "limit":
            low, high = col.low or 1, col.high or 99
            apply(_custom_validation(
                f'OR(ISBLANK({first}),{first}="{LIMIT_NONE}",{first}="ללא",'
                f'IF(ISNUMBER({first}),AND({first}>={low},{first}<={high},INT({first})={first}),FALSE))',
                "מספר לא תקין", f"מספר שלם בין {low} ל-{high}, או '{LIMIT_NONE}'"), key)
        elif col.kind == "image":
            apply(_custom_validation(
                f'OR(ISBLANK({first}),{first}="{CLEAR_WORD}",LEFT({first},4)="http",LEFT({first},4)="www.")',
                "קישור לא תקין", "הדביקו קישור לתמונה שמתחיל ב-https:// (או 'ללא' להסרת התמונה). להמשיך?",
                strict=False), key)
        elif col.kind == "unit":
            apply(_list_validation(lists["units"], strict=False,
                                   error="יחידה שלא ברשימה - להמשיך?"), key)
        elif col.kind == "printers":
            apply(_list_validation(lists["printers"], strict=False, free_text=True,
                                   prompt=f"בחרו מהרשימה או כתבו כמה שמות מופרדים בפסיק. "
                                          f"'{PRINTERS_NONE}' = לא להדפיס, ריק = ללא שינוי"), key)
        elif col.kind == "groups":
            apply(_list_validation(refs.groups, strict=False, free_text=True,
                                   prompt=f"בחרו קבוצה מגיליון '{SHEET_GROUPS}', או כתבו כמה מופרדות בפסיק "
                                          f"(בסדר ההצגה). '{GROUPS_NONE}' = בלי תוספות, '{GROUPS_INHERIT}' = לפי "
                                          "המחלקה, ריק = ללא שינוי"), key)
        elif col.kind == "group":
            apply(_list_validation(refs.groups, strict=False,
                                   error=f"הקבוצה לא מופיעה בגיליון '{SHEET_GROUPS}' - הוסיפו אותה שם. להמשיך?"), key)
        elif col.kind == "kind":
            apply(_list_validation(lists["kinds"], strict=True, error="בחרו: תוספת, בחירה או הסרה"), key)
        elif col.kind == "category_list":
            apply(_list_validation(refs.categories, strict=False, free_text=True,
                                   prompt="בחרו מחלקה, או כתבו כמה מופרדות בפסיק"), key)
        elif col.kind == "product_list":
            apply(_list_validation(refs.products, strict=False, free_text=True,
                                   prompt="בחרו פריט, או כתבו כמה מופרדים בפסיק (שם, מק״ט או ברקוד)"), key)
        elif col.kind == "ticket":
            apply(_list_validation(lists["tickets"], strict=True, error="בחרו מהרשימה"), key)
        elif key == "entries":
            apply(_number_validation("whole", 0, ENTRIES_MAX, f"מספר שלם בין 1 ל-{ENTRIES_MAX} (0 = לבטל)"), key)
        elif col.kind == "int" and col.low is not None and col.high is not None:
            apply(_number_validation("whole", col.low, col.high, f"מספר שלם בין {col.low} ל-{col.high}"), key)
        elif key == "sort":
            apply(_number_validation("whole", SORT_MIN, SORT_MAX, "מספר שלם"), key)
        elif key in ("category", "parent"):
            apply(_list_validation(refs.categories, strict=False,
                                   error="המחלקה לא מופיעה בגיליון מחלקות - היא תיווצר בקליטה. להמשיך?"), key)
        elif col.kind == "dietary":
            apply(_list_validation(lists["dietary"], strict=False, free_text=True,
                                   prompt="בחרו מהרשימה או כתבו כמה מופרדים בפסיק (למשל: טבעוני, חריף). "
                                          "'ללא' = לנקות, ריק = ללא שינוי"), key)
        elif col.kind == "channel":
            apply(_list_validation(lists["channels"], strict=True, error="בחרו מהרשימה"), key)
        elif col.max_len:
            apply(_length_validation(col.max_len), key)


def _data_sheet(ws, columns: Sequence[Column], examples: Sequence[Dict[str, Any]], rows: Sequence[Dict[str, Any]],
                refs: _Refs, tab: str, *, freeze: str = "C2") -> None:
    _rtl(ws)
    ws.sheet_properties.tabColor = tab
    _header(ws, columns, freeze)
    next_row = _write_rows(ws, columns, examples, 2, example=True)
    next_row = _write_rows(ws, columns, rows, next_row, example=False)
    last = max(next_row - 1 + SPARE_ROWS, MIN_ROWS)
    _column_formats(ws, columns)
    _validations(ws, columns, last, refs)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(next_row - 1, 2)}"
    _page_setup(ws)


@dataclass
class _Examples:
    categories: List[Dict[str, Any]]
    products: List[Dict[str, Any]]
    groups: List[Dict[str, Any]]
    options: List[Dict[str, Any]]
    notes: List[Dict[str, Any]]


def _examples(view: TemplateView) -> _Examples:
    """A few rows on each sheet that show the shape of a menu, with the company's own printer names."""
    names = list(view.printer_choices)
    food = names[0] if names else "מטבח"
    drinks = names[1] if len(names) > 1 else ("בר" if not names else names[0])
    doneness, extras, changes = "מידת עשייה", "תוספות להמבורגר", "שינויים במנה"
    categories = [
        {"marker": EXAMPLE_MARK, "name": "מנות עיקריות", "printers": food, "sort": 1, "active": YES},
        {"marker": EXAMPLE_MARK, "name": "המבורגרים", "parent": "מנות עיקריות", "sort": 2, "active": YES,
         "groups": f"{doneness}, {extras}, {changes}"},
        {"marker": EXAMPLE_MARK, "name": "שתייה", "printers": drinks, "sort": 3, "active": YES},
    ]
    burger = {"marker": EXAMPLE_MARK, "name": "המבורגר קלאסי", "category": "המבורגרים", "price": 58,
              "cost": 18.5, "open_price": NO, "weighed": NO, "unit": UNIT_PIECE, "no_discount": NO,
              "active": YES, "description": "200 גרם, חסה, עגבנייה, בצל", "dietary": "בשרי",
              "image": "https://example.com/images/burger.jpg"}
    if view.menu_names:
        burger[f"menu:{view.menu_names[0]}"] = 52
    products = [
        burger,
        {"marker": EXAMPLE_MARK, "name": "קולה זכוכית", "category": "שתייה", "price": 12,
         "barcode": "7290001234567", "cost": 4.2, "open_price": NO, "weighed": NO, "unit": UNIT_PIECE,
         "no_discount": YES, "printers": drinks, "active": YES},
        {"marker": EXAMPLE_MARK, "name": "סלט חומוס", "category": "מנות עיקריות", "price": 49.9,
         "open_price": NO, "weighed": YES, "unit": UNIT_KG, "no_discount": NO,
         "ticket": TICKET_LABELS["per_line"], "active": YES, "description": "מחיר לק״ג",
         "dietary": "טבעוני, ללא גלוטן", "groups": changes},
    ]
    kinds = GROUP_KIND_LABELS
    groups = [
        {"marker": EXAMPLE_MARK, "name": doneness, "kind": kinds["choice"], "required": YES, "min": 1, "max": 1,
         "free": 0, "allow_quantity": NO, "allow_pre": NO, "sort": 1, "active": YES},
        {"marker": EXAMPLE_MARK, "name": extras, "kind": kinds["addon"], "required": NO, "min": 0, "max": 3,
         "free": 1, "allow_quantity": YES, "allow_pre": YES, "sort": 2, "active": YES},
        {"marker": EXAMPLE_MARK, "name": changes, "kind": kinds["removal"], "required": NO, "min": 0,
         "max": LIMIT_NONE, "free": 0, "allow_quantity": NO, "allow_pre": NO, "sort": 3, "active": YES},
    ]
    option = lambda group, name, price, sort, **kw: {  # noqa: E731
        "marker": EXAMPLE_MARK, "group": group, "name": name, "price": price, "sort": sort, "active": YES,
        "default": NO, **kw}
    options = [
        option(doneness, "מדיום", 0, 1, default=YES),
        option(doneness, "מדיום וול", 0, 2),
        option(doneness, "וול דאן", 0, 3),
        option(extras, "גבינה צהובה", 4, 1, kitchen_name="גבינה", max_qty=2),
        option(extras, "בצל מטוגן", 3, 2),
        option(extras, "ביצת עין", 5, 3, kitchen_name="ביצה"),
        option(changes, "בלי בצל", 0, 1),
        option(changes, "בלי עגבנייה", 0, 2),
    ]
    notes = [
        {"marker": EXAMPLE_MARK, "name": "רוטב בצד", "all": NO, "categories": "המבורגרים, מנות עיקריות",
         "important": NO, "sort": 1},
        {"marker": EXAMPLE_MARK, "name": "בלי קרח", "all": NO, "products": "קולה זכוכית", "important": NO},
        {"marker": EXAMPLE_MARK, "name": "אלרגיה לאגוזים", "all": YES, "important": YES},
    ]
    return _Examples(categories, products, groups, options, notes)


def _lists_sheet(wb: Workbook, view: TemplateView) -> Dict[str, str]:
    ws = wb.create_sheet(SHEET_LISTS)
    _rtl(ws)
    ws.sheet_state = "hidden"
    columns = {
        "printers": list(view.printer_choices) + [PRINTERS_NONE, PRINTERS_INHERIT],
        "yes_no": [YES, NO],
        "units": list(UNIT_CHOICES),
        "tickets": [label for mode, label in TICKET_LABELS.items()],
        "dietary": list(DIETARY_LABELS_HE.values()) + [DIETARY_NONE],
        "channels": list(SALES_CHANNEL_LABELS_HE.values()),
        "kinds": list(GROUP_KIND_LABELS.values()),
    }
    refs: Dict[str, str] = {}
    for index, (key, values) in enumerate(columns.items(), start=1):
        letter = get_column_letter(index)
        ws.cell(row=1, column=index, value=key)
        for r, value in enumerate(values, start=2):
            ws.cell(row=r, column=index, value=value)
        refs[key] = f"{_quoted(SHEET_LISTS)}!${letter}$2:${letter}${len(values) + 1}"
    return refs


def _printers_sheet(ws, view: TemplateView) -> None:
    _rtl(ws)
    ws.sheet_properties.tabColor = GREY
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(PRINTER_REFERENCE_HEADERS))
    intro = ws.cell(
        row=1, column=1,
        value="המדפסות שהוגדרו בחברה בעת הפקת הקובץ - לעיון בלבד. בעמודה 'מדפסות' בגיליונות "
              "'מחלקות' ו'פריטים' רשמו את השם מהעמודה הראשונה. השיוך חל בכל סניף שיש בו מדפסת בשם הזה.",
    )
    intro.alignment = _RIGHT
    intro.font = Font(color=INK)
    ws.row_dimensions[1].height = 48
    for index, title in enumerate(PRINTER_REFERENCE_HEADERS, start=1):
        cell = ws.cell(row=3, column=index, value=title)
        cell.fill, cell.font, cell.alignment, cell.border = _HEADER_OPTIONAL, _HEADER_FONT, _WRAP_CENTER, _BOX
    widths = (24, 22, 22, 24, 9)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    if not view.printers:
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(PRINTER_REFERENCE_HEADERS))
        empty = ws.cell(row=4, column=1, value="לא הוגדרו עדיין מדפסות בחברה. אפשר להגדיר אותן במסך 'מדפסות בונים' "
                                               "ולהוריד את הטמפלט מחדש, או לרשום עכשיו את השמות המתוכננים.")
        empty.alignment = _RIGHT
        empty.font = _EXAMPLE_FONT
        ws.row_dimensions[4].height = 32
    for r, printer in enumerate(view.printers, start=4):
        values = (printer.name, printer.shop, printer.connection, printer.narrowed, YES if printer.active else NO)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = _BOX
            if not printer.active:
                cell.font = _EXAMPLE_FONT
    ws.freeze_panes = "A4"
    # Read-only: locked (no password - it is a reference, not a secret).
    ws.protection.sheet = True


# ── Instructions ──────────────────────────────────────────────────────────────


def _title(view: TemplateView) -> str:
    return f"טמפלט לקליטת פריטים - {view.company_name}" if view.company_name else "טמפלט לקליטת התפריט"


def _instructions(ws, view: TemplateView) -> None:
    _rtl(ws)
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.showGridLines = False
    for letter, width in zip("ABCD", (24, 9, 70, 26)):
        ws.column_dimensions[letter].width = width

    row = 1

    def line(text: str, *, font: Optional[Font] = None, height: Optional[float] = None, gap: int = 0) -> None:
        nonlocal row
        row += gap
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=text)
        cell.alignment = _RIGHT
        if font is not None:
            cell.font = font
        lines = max(1, math.ceil(len(text) / 115))
        ws.row_dimensions[row].height = height or (16 * lines + 4)
        row += 1

    def section(title: str) -> None:
        line(title, font=Font(bold=True, size=13, color=BLUE), height=24, gap=1)

    line(_title(view), font=Font(bold=True, size=18, color=INK), height=30)
    counts = [f"{len(view.categories)} מחלקות", f"{len(view.groups)} קבוצות תוספות", f"{len(view.printers)} מדפסות"]
    if view.with_data:
        counts.insert(0, f"{len(view.products)} פריטים")
    if view.menu_names:
        counts.append(f"{len(view.menu_names)} תפריטים")
    stamp = view.generated_at.strftime("%d/%m/%Y %H:%M")
    # A company's file says what it holds; the generic one (no company) only when it was made.
    line(f"הופק ב-{stamp}" + (" · " + " · ".join(counts) if view.company_name else ""), font=Font(color=GREY))

    section("מה יש בקובץ")
    for text in (
        f"'{SHEET_CATEGORIES}' - מחלקות התפריט (מנות עיקריות, שתייה...), מחלקות משנה, מדפסות ותוספות לכל המחלקה.",
        f"'{SHEET_PRODUCTS}' - שורה לכל פריט: שם, מחלקה, מחיר, ברקוד, תוספות, תמונה ועוד.",
        f"'{SHEET_GROUPS}' - קבוצות התוספות: 'מידת עשייה', 'תוספות להמבורגר', 'בלי...' - כמה בוחרים, חובה או לא.",
        f"'{SHEET_OPTIONS}' - האפשרויות בכל קבוצה, עם מחיר לכל אחת (גבינה +4, בצל +3, מדיום).",
        f"'{SHEET_NOTES}' - כפתורי הערה למטבח ('בלי בצל', 'רוטב בצד') למחלקות, לפריטים או לכל המנות.",
        f"'{SHEET_PRINTERS}' - המדפסות שהוגדרו בחברה, לעיון בלבד.",
    ):
        line(text)

    section("איך ממלאים - בסדר הזה")
    for text in (
        "1. 'מחלקות': רשמו את מחלקות התפריט. למחלקת משנה רשמו בעמודה 'מחלקת אב' את המחלקה שמעליה.",
        f"2. '{SHEET_GROUPS}': שורה לכל קבוצת תוספות. '{SHEET_OPTIONS}': שורה לכל אפשרות - בעמודה 'קבוצה' "
        "בחרו את הקבוצה מהרשימה, וכתבו את שם האפשרות ואת המחיר שלה.",
        "3. 'פריטים': שורה לכל פריט. חובה: שם פריט, מחלקה ומחיר (העמודות המסומנות ב-* בכחול). כל השאר רשות. "
        "מחלקה שלא רשמתם בגיליון 'מחלקות' תיווצר אוטומטית.",
        "4. קבוצות תוספות לפריט: בעמודה 'קבוצות תוספות' בגיליון 'פריטים' (או לכל המחלקה - באותה עמודה "
        "בגיליון 'מחלקות'), שמות הקבוצות מופרדים בפסיק, בסדר שבו יוצגו.",
        f"5. '{SHEET_NOTES}': טקסט ההערה, ועל מה היא חלה - מחלקות, פריטים, או 'כן' בעמודה 'לכל המנות'.",
        "6. תמונות: בעמודה 'קישור לתמונה' הדביקו קישור לתמונה באינטרנט (https://...). אפשר גם קישור שיתוף "
        "של Google Drive או Dropbox (משותף לכל מי שיש לו את הקישור).",
        "7. בעמודות כן/לא, יחידה, סוג ושובר פריט בחרו מהרשימה שנפתחת בתא. תא ריק בשורה חדשה = ברירת המחדל.",
        f"8. השורות האפורות שמסומנות '{EXAMPLE_MARK}' לא ייקלטו - אפשר למחוק אותן או להשאיר. כדי לדלג על "
        f"שורה כתבו '{SKIP_MARK}' בעמודה 'סימון'.",
        "9. שמרו את הקובץ כ-xlsx והחזירו אותו, או העלו אותו במסך 'ייבוא פריטים מאקסל' בניהול. לפני הקליטה "
        "מוצגת בדיקה של כל השורות.",
        "העבירו את העכבר מעל כותרת של עמודה כדי לראות הסבר עליה.",
    ):
        line(text)

    section("מחירים")
    menu_line = (
        f"מחיר בתפריט: לכל תפריט בחברה יש בגיליון 'פריטים' עמודה '{MENU_PRICE_PREFIX}: <שם התפריט>' - המחיר "
        "בזמן שהתפריט פעיל (למשל 'צהריים'). ריק = מחיר הקטלוג; פריט שעוד לא בתפריט יתווסף אליו."
        if view.menu_names else
        "מחיר בתפריט: כשמוגדרים בחברה תפריטים (בוקר, צהריים...), יופיעו בטמפלט עמודות "
        f"'{MENU_PRICE_PREFIX}: <שם התפריט>' למחיר מיוחד בזמן שהתפריט פעיל."
    )
    for text in (
        "כל המחירים בשקלים וכוללים מע״מ - בדיוק כמו בטופס המוצר: 12.50 (אפשר גם ₪12.50). המע״מ הוא של הפריט.",
        "'עלות' היא לפני מע״מ ואינה מוצגת בקופה - לניתוח רווחיות בלבד.",
        "מחיר אפשרות (בגיליון 'אפשרויות') מתווסף למחיר המנה כשבוחרים בה. 0 = חינם; אפשר גם מינוס (הנחה).",
        menu_line,
    ):
        line(text)

    section("תוספות - איך זה עובד")
    for text in (
        "סוג הקבוצה: 'בחירה' - בוחרים אחת מכמה (מידת עשייה, גודל); 'תוספת' - מוסיפים למנה (+ גבינה); "
        "'הסרה' - מורידים רכיב (בלי בצל), מודפס מודגש בבון.",
        "'חובה' = כן: הקופאי (או הלקוח בקיוסק) חייב לבחור לפני שהמנה נכנסת להזמנה. 'מינימום' ו'מקסימום' - "
        "כמה בוחרים. 'כמה בחינם' - כמה מהבחירות לא יחויבו (הזולות ביותר).",
        "פריט מקבל את הקבוצות של המחלקה שלו, אלא אם רשמתם לו קבוצות משלו. מחלקת משנה מקבלת את של מחלקת האב. "
        f"'{GROUPS_NONE}' = בלי תוספות בכלל; '{GROUPS_INHERIT}' = לבטל הגדרה קודמת ולחזור לקבוצות של המחלקה.",
        "קבוצה שכבר קיימת במערכת מזוהה לפי השם, ואפשרות - לפי השם בתוך הקבוצה. אפשרות שלא מופיעה בקובץ לא "
        "נמחקת; כדי להסתיר אותה כתבו 'לא' בעמודה 'פעילה'.",
    ):
        line(text)

    section("הערות מהירות")
    for text in (
        "כל הערה היא כפתור שהקופאי לוחץ עליו כשהוא מוסיף מנה ('בלי בצל', 'רוטב בצד'), והיא מודפסת בבון למטבח. "
        "'חשובה' = כן: מודפסת בהדגשה.",
        "הערה של מחלקה מוצעת בכל הפריטים שבה. פריט שמקבל הערה משלו שומר גם את ההערות שקיבל מהמחלקה.",
        "בעמודה 'פריטים' רשמו את שמות הפריטים מופרדים בפסיק; כשיש שני פריטים באותו שם - רשמו את המק״ט.",
    ):
        line(text)

    section("מדפסות המטבח והבר")
    for text in (
        "בעמודה 'מדפסות' רשמו את שמות המדפסות בדיוק כפי שהם מופיעים בגיליון 'מדפסות', מופרדים בפסיק "
        "(למשל: מטבח, בר).",
        "מדפסת של מחלקה חלה על כל הפריטים שבה ועל מחלקות המשנה שלה. מדפסת שנרשמה לפריט גוברת על המחלקה.",
        "הניתוב נקבע לפי שם המדפסת בכל סניף של החברה: בסניף שיש בו מדפסת בשם הזה - ההדפסה תהיה בה; "
        "בסניף שאין בו מדפסת בשם הזה - לא משתנה דבר.",
        f"'{PRINTERS_NONE}' = לא להדפיס. '{PRINTERS_INHERIT}' = לבטל הגדרה קודמת (פריט - חוזר להגדרת המחלקה; "
        "מחלקה - חוזרת למחלקת האב). תא ריק = ללא שינוי.",
        "שם מדפסת שלא מוגדרת באף סניף - השיוך שלו יידלג, ותופיע על כך אזהרה בבדיקת הקובץ.",
    ):
        line(text)
    if view.has_mixed_routing:
        line("תא מדפסות ריק שיש עליו הערה: הניתוב שלו שונה בין הסניפים, ולכן לא נכתב. ערכו אותו במסך "
             "'מדפסות בונים'.", font=Font(color=ORANGE))

    section("עדכון פריטים קיימים")
    for text in (
        "פריט קיים מזוהה לפי הברקוד; אם אין ברקוד - לפי המק״ט; ואם אין גם מק״ט - לפי שם זהה באותה מחלקה. "
        "מחלקה, קבוצת תוספות ותפריט מזוהים לפי השם.",
        "פריט שזוהה יעודכן לפי השורה, ופריט שלא זוהה ייווצר. פריט חדש נמכר בכל הסניפים של החברה.",
        "תא ריק בשורה קיימת = הערך הקיים נשמר. פריטים, מחלקות, קבוצות, אפשרויות והערות שלא מופיעים בקובץ "
        "לא נמחקים.",
        "לפני הקליטה מוצגת בדיקה של הקובץ: מה ייווצר, מה יתעדכן, אזהרות ושגיאות (עם מספר השורה). "
        "ייבוא חוזר של אותו קובץ לא משנה דבר.",
    ):
        line(text)

    section("העמודות")
    head = row
    for c, text in enumerate(("עמודה", "חובה", "מה לרשום", "דוגמה"), start=1):
        cell = ws.cell(row=head, column=c, value=text)
        cell.fill, cell.font, cell.alignment, cell.border = _HEADER_REQUIRED, _HEADER_FONT, _WRAP_CENTER, _BOX
    row += 1
    sheets = ((SHEET_CATEGORIES, CATEGORY_COLUMNS), (SHEET_PRODUCTS, view.product_columns()),
              (SHEET_GROUPS, GROUP_COLUMNS), (SHEET_OPTIONS, OPTION_COLUMNS), (SHEET_NOTES, NOTE_COLUMNS))
    for sheet, columns in sheets:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        sub = ws.cell(row=row, column=1, value=f"גיליון '{sheet}'")
        sub.font = Font(bold=True, color=INK)
        sub.fill = _EXAMPLE_FILL
        row += 1
        for col in columns:
            values = (col.title, YES if col.required else "", col.help, col.example)
            for c, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=value or None)
                cell.alignment = _RIGHT
                cell.border = _BOX
                if c == 1:
                    cell.font = Font(bold=True)
            ws.row_dimensions[row].height = 16 * max(1, math.ceil(len(col.help) / 70)) + 4
            row += 1


# ── The workbook ──────────────────────────────────────────────────────────────


def build_workbook(view: TemplateView) -> bytes:
    wb = Workbook()
    help_ws = wb.active
    help_ws.title = SHEET_INSTRUCTIONS
    categories_ws = wb.create_sheet(SHEET_CATEGORIES)
    products_ws = wb.create_sheet(SHEET_PRODUCTS)
    groups_ws = wb.create_sheet(SHEET_GROUPS)
    options_ws = wb.create_sheet(SHEET_OPTIONS)
    notes_ws = wb.create_sheet(SHEET_NOTES)
    printers_ws = wb.create_sheet(SHEET_PRINTERS)
    lists = _lists_sheet(wb, view)

    examples = _Examples([], [], [], [], []) if view.with_data else _examples(view)

    def names_ref(sheet: str, rows: int) -> str:
        return f"{_quoted(sheet)}!$B$2:$B${max(rows + 1 + SPARE_ROWS, MIN_ROWS)}"

    refs = _Refs(
        lists=lists,
        categories=names_ref(SHEET_CATEGORIES, len(examples.categories) + len(view.categories)),
        groups=names_ref(SHEET_GROUPS, len(examples.groups) + len(view.groups)),
        products=names_ref(SHEET_PRODUCTS, len(examples.products) + len(view.products)),
    )

    _instructions(help_ws, view)
    _data_sheet(categories_ws, CATEGORY_COLUMNS, examples.categories, view.categories, refs, ORANGE)
    _data_sheet(products_ws, view.product_columns(), examples.products, view.products, refs, GREEN)
    _data_sheet(groups_ws, GROUP_COLUMNS, examples.groups, view.groups, refs, PURPLE)
    _data_sheet(options_ws, OPTION_COLUMNS, examples.options, view.options, refs, PINK, freeze="D2")
    _data_sheet(notes_ws, NOTE_COLUMNS, examples.notes, view.notes, refs, TEAL)
    _printers_sheet(printers_ws, view)

    wb.properties.title = _title(view)
    wb.properties.creator = "POS"
    wb.active = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_csv(view: TemplateView) -> bytes:
    """The products sheet alone, as UTF-8 CSV (with a BOM, so Excel reads the Hebrew)."""
    out = io.StringIO()
    writer = csv.writer(out)
    columns = view.product_columns()
    writer.writerow([col.header for col in columns])
    rows = view.products if view.with_data else _examples(view).products
    for row in rows:
        values = []
        for col in columns:
            value = row.get(col.key)
            if value is None:
                values.append("")
            elif isinstance(value, Decimal):
                values.append(f"{value:.2f}")
            else:
                values.append(str(value))
        writer.writerow(values)
    return ("﻿" + out.getvalue()).encode("utf-8")


def ticket_label(mode: Optional[str]) -> str:
    """A stored ticket mode as the sheet writes it; the category's (None) is blank."""
    if mode is None or mode == TICKET_INHERIT:
        return ""
    return TICKET_LABELS.get(mode, "")
