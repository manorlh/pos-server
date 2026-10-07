"""
The menu sheet ("טמפלט קליטת פריטים") as data: its sheets and columns, the words a cell
may hold, and reading a filled-in file - XLSX or CSV - back into raw rows.

The sheets: מחלקות, פריטים (with a "מחיר בתפריט: <menu>" column per catalog menu),
קבוצות תוספות (add-on groups), אפשרויות (their options, one row each), הערות מהירות
(quick-note chips); הוראות and מדפסות are for reading only. A CSV is the products sheet
(or the categories sheet) alone.

No database here. The rows are matched against a company's catalog in
app/services/catalog_import.py and the workbook is drawn in app/services/catalog_template.py,
both from the column lists below, so the three can never disagree about a header.

Reading is lenient about the things people do to a spreadsheet: columns in another order
or under a near name ("קטגוריה" for "מחלקה", 'מק"ט' for "מק״ט"), a header that is not on
row 1, a CSV saved by Hebrew Excel (Windows-1255, ";" separated), a price typed as
"₪1,234.50". Each such reading of a cell is reported back as a warning, never silently.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Sequence, Tuple

from app.schemas.menu import (  # noqa: F401 - NOTES_MAX / OPTIONS_MAX are read through this module
    LINKS_MAX, MONEY_MAX, NAME_MAX as MENU_NAME_MAX, NOTES_MAX, OPTIONS_MAX, SELECT_MAX,
)
from app.services import dietary
from app.services import sales_channel
from app.services.dietary import DIETARY_LABELS_HE
from app.services.sales_channel import SALES_CHANNEL_LABELS_HE

# ── Sheets ────────────────────────────────────────────────────────────────────

SHEET_INSTRUCTIONS = "הוראות"
SHEET_CATEGORIES = "מחלקות"
SHEET_PRODUCTS = "פריטים"
#: The add-on groups ("מידת עשייה", "תוספות להמבורגר") and their options, one row each.
SHEET_GROUPS = "קבוצות תוספות"
SHEET_OPTIONS = "אפשרויות"
#: The quick-note chips ("בלי בצל") the till and the kiosk offer on a dish.
SHEET_NOTES = "הערות מהירות"
SHEET_PRINTERS = "מדפסות"
#: Hidden: the sources of the dropdowns that are not a sheet of their own.
SHEET_LISTS = "רשימות"

# ── Words ─────────────────────────────────────────────────────────────────────

YES = "כן"
NO = "לא"
#: In the "סימון" column: a row the import skips.
EXAMPLE_MARK = "דוגמה"
SKIP_MARK = "דלג"
#: In a "מדפסות" cell: print nothing ("no ticket"), or drop the row's own setting.
PRINTERS_NONE = "ללא"
PRINTERS_INHERIT = "ירושה"
#: In "מחלקת אב": move an existing category to the top level.
PARENT_NONE = "ללא"
#: Separates a parent from its child where a category name alone is ambiguous.
PATH_SEPARATOR = " > "

UNIT_PIECE = "יח׳"
UNIT_KG = "ק״ג"
UNIT_CHOICES = (UNIT_PIECE, UNIT_KG, "ליטר")

#: Item-ticket ("שובר פריט") modes as the sheet writes them; "" is the category's.
TICKET_INHERIT = "inherit"
TICKET_LABELS: Dict[str, str] = {
    TICKET_INHERIT: "לפי המחלקה",
    "off": "ללא שובר",
    "per_unit": "לכל יחידה",
    "per_line": "לכל פריט",
    "per_sale": "אחד לעסקה",
}

#: Add-on group kinds (app/models/menu.py GROUP_KINDS) as the sheet writes them, as the
#: dashboard's group editor names them.
GROUP_KIND_LABELS: Dict[str, str] = {"addon": "תוספת", "choice": "בחירה", "removal": "הסרה"}
#: In a "קבוצות תוספות" cell of a product / category: no add-ons at all, or drop the row's
#: own list and follow the category again - the words of the "מדפסות" cell.
GROUPS_NONE = "ללא"
GROUPS_INHERIT = "ירושה"
#: In "מקסימום בחירות" / "מקסימום לאפשרות": no upper limit.
LIMIT_NONE = "ללא הגבלה"
#: In "קישור לתמונה": remove the picture. In "שם למטבח" / a menu price: clear it.
CLEAR_WORD = "ללא"
#: Before a menu's name in a products column: "מחיר בתפריט: צהריים".
MENU_PRICE_PREFIX = "מחיר בתפריט"

#: Limits of the columns they land in (app/models/product.py, category.py, menu.py).
NAME_MAX = 255
DESCRIPTION_MAX = 1000
CODE_MAX = 100
UNIT_MAX = 16
PRICE_MAX = Decimal("99999999.99")
#: As the insights' own cost endpoint allows (app/services/insights/service.py).
COST_MAX = Decimal("100000")
ENTRIES_MAX = 50
#: The add-on layer's limits, as its own editor enforces them (app/schemas/menu.py).
GROUP_NAME_MAX = MENU_NAME_MAX
OPTION_NAME_MAX = MENU_NAME_MAX
KITCHEN_NAME_MAX = 60
NOTE_TEXT_MAX = 60
OPTION_PRICE_MIN = Decimal("-1000")
OPTION_PRICE_MAX = MONEY_MAX
#: A catalog menu's price (app/schemas/catalog_menu.py).
MENU_PRICE_MAX = MONEY_MAX
IMAGE_URL_MAX = 2000
SORT_MIN, SORT_MAX = -100000, 100000

#: Bounds on what one upload may make us read.
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_COLUMNS = 60
#: A sheet formatted to its last row reports a million rows; physical rows past this are
#: not read (and the preview says so). Empty rows in between cost nothing.
MAX_SHEET_ROWS = 20000
HEADER_SEARCH_ROWS = 10

CENT = Decimal("0.01")


# ── Columns ───────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Column:
    key: str
    #: As written in the template; a required column ends with "*".
    header: str
    width: float
    #: text | code | money | signed_money | menu_price | bool | int | limit | printers |
    #: groups | group | kind | ticket | unit | dietary | channel | image | category_list |
    #: product_list | marker
    kind: str
    help: str
    example: str = ""
    #: Other headers read as this column (normalized like `normalize_header`).
    aliases: Tuple[str, ...] = ()
    #: A text column's length limit, checked in Excel too.
    max_len: Optional[int] = None
    #: A number column's range, checked in Excel too.
    low: Optional[int] = None
    high: Optional[int] = None

    @property
    def required(self) -> bool:
        return self.header.endswith("*")

    @property
    def title(self) -> str:
        """The header without its "*"."""
        return self.header.rstrip("*").strip()


MARKER = Column(
    "marker", "סימון", 9, "marker",
    "שורות שמסומנות 'דוגמה' (או 'דלג') לא ייקלטו. בשורות שלכם השאירו ריק.",
    EXAMPLE_MARK, ("דוגמה", "דוגמא", "marker"),
)

PRODUCT_COLUMNS: Tuple[Column, ...] = (
    MARKER,
    Column("name", "שם פריט*", 30, "text",
           "השם שיופיע בקופה ובקבלה (עד 255 תווים).", "המבורגר קלאסי",
           ("שם פריט", "שם", "שם מוצר", "מוצר", "פריט", "name", "product", "item"), max_len=NAME_MAX),
    Column("category", "מחלקה*", 22, "text",
           "שם המחלקה מגיליון 'מחלקות'. מחלקה שלא קיימת תיווצר אוטומטית.", "המבורגרים",
           ("מחלקה", "קטגוריה", "מחלקה ראשית", "category", "department")),
    Column("price", "מחיר*", 11, "money",
           "מחיר מכירה בשקלים כולל מע״מ: 58 או 12.50 (אפשר גם ₪12.50).", "58",
           ("מחיר", "מחיר מכירה", "מחיר לצרכן", "price")),
    Column("barcode", "ברקוד", 17, "code",
           "הברקוד שעל המוצר. משמש גם לזיהוי פריט קיים.", "7290001234567",
           ("ברקוד", "barcode", "ean", "upc")),
    Column("sku", "מק״ט", 13, "code",
           "קוד פנימי. ריק בפריט חדש = יוקצה אוטומטית. משמש לזיהוי פריט קיים.", "1001",
           ("מקט", "קוד פריט", "קוד", "sku", "code")),
    Column("cost", "עלות (ללא מע״מ)", 13, "money",
           "כמה עולה לכם יחידה, לפני מע״מ - לניתוח רווחיות. לא מוצג בקופה.", "18.50",
           ("עלות", "עלות ללא מעמ", "מחיר עלות", "מחיר קנייה", "cost")),
    Column("open_price", "מחיר פתוח", 11, "bool",
           "כן = הקופאי מקליד את הסכום בכל מכירה; המחיר בעמודה 'מחיר' הוא הצעה בלבד.", NO,
           ("מחיר פתוח", "open price")),
    Column("weighed", "נמכר במשקל", 12, "bool",
           "כן = הכמות נשקלת (למשל 0.75) והמחיר הוא ליחידה שבעמודה 'יחידה'.", NO,
           ("נמכר במשקל", "שקיל", "במשקל", "משקל", "weighed")),
    Column("unit", "יחידה", 9, "unit",
           "יח׳ או ק״ג (או ליטר). למוצר שנמכר במשקל - בדרך כלל ק״ג.", UNIT_PIECE,
           ("יחידה", "יחידת מידה", "יחידת מכירה", "unit")),
    Column("no_discount", "ללא הנחה", 10, "bool",
           "כן = הקופה לא תיתן על הפריט שום הנחה, ומבצעים לא יחולו עליו.", NO,
           ("ללא הנחה", "לא מקבל הנחות", "בלי הנחה", "no discount")),
    Column("printers", "מדפסות", 22, "printers",
           "לאן יודפס הפריט במטבח/בבר: שמות מגיליון 'מדפסות' מופרדים בפסיק. ריק = לפי "
           "המחלקה. 'ללא' = לא להדפיס. 'ירושה' = לבטל הגדרה קודמת ולחזור למחלקה.", "בר",
           ("מדפסות", "מדפסת", "מדפסות בונים", "תחנה", "תחנות", "printers", "printer")),
    Column("ticket", "שובר פריט", 15, "ticket",
           "שובר להחלפה בפריט (למשל בדוכן): ריק = לפי המחלקה, 'ללא שובר', 'לכל יחידה', "
           "'לכל פריט' או 'אחד לעסקה'.", "",
           ("שובר פריט", "שובר", "הדפסת שובר פריט", "ticket")),
    Column("entries", "כרטיס כניסה", 12, "int",
           "לכרטיס כניסה: כמה כניסות יחידה אחת מקנה (1-50). ריק = פריט רגיל; 0 = לבטל.", "",
           ("כרטיס כניסה", "כניסות", "מספר כניסות", "entries")),
    Column("active", "פעיל", 8, "bool",
           "לא = הפריט מוצג בקופה אבל חסום למכירה. ריק בפריט חדש = כן.", YES,
           ("פעיל", "זמין", "active")),
    Column("description", "תיאור", 40, "text",
           "טקסט חופשי (לא חובה). מוצג בקיוסק ליד הפריט - עד 300 תווים מומלץ, 1000 לכל היותר.",
           "200 גרם, חסה, עגבנייה",
           ("תיאור", "תיאור הפריט", "הערות", "description"), max_len=DESCRIPTION_MAX),
    Column("dietary", "סימוני תזונה", 22, "dietary",
           "מופרדים בפסיק: " + ", ".join(DIETARY_LABELS_HE.values()) + " (או הקודים באנגלית). "
           "טבעוני מסומן גם כצמחוני; בשרי וחלבי לא יחד. ריק = ללא שינוי; 'ללא' = לנקות.",
           "בשרי, חריף",
           ("סימוני תזונה", "סימון תזונה", "תזונה", "סימונים", "dietary", "dietary tags", "diet")),
    Column("channel", "ערוץ מכירה", 15, "channel",
           "היכן הפריט נמכר: " + ", ".join(SALES_CHANNEL_LABELS_HE.values()) + ". "
           "ריק בפריט חדש = קופות וקיוסק; ריק בפריט קיים = ללא שינוי.",
           SALES_CHANNEL_LABELS_HE[sales_channel.ALL],
           ("ערוץ מכירה", "ערוץ", "היכן נמכר", "היכן הפריט נמכר", "נמכר ב", "channel", "sales channel")),
    Column("groups", "קבוצות תוספות", 28, "groups",
           "קבוצות התוספות של הפריט: שמות מגיליון 'קבוצות תוספות' מופרדים בפסיק, בסדר שבו יוצגו "
           "בקופה ובקיוסק. ריק = לפי המחלקה (בפריט קיים: ללא שינוי). 'ללא' = בלי תוספות בכלל; "
           "'ירושה' = לבטל הגדרה קודמת ולחזור לקבוצות של המחלקה.",
           "מידת עשייה, תוספות להמבורגר",
           ("קבוצות תוספות", "קבוצות", "תוספות", "קבוצת תוספות", "modifiers", "modifier groups", "addons",
            "add-ons")),
    Column("image", "קישור לתמונה", 34, "image",
           "כתובת אינטרנט של תמונת הפריט (מתחילה ב-https://), JPG / PNG / WEBP עד 5MB. התמונה "
           "מורדת ונשמרת במערכת בזמן הייבוא. ריק = ללא שינוי; 'ללא' = להסיר את התמונה.",
           "https://example.com/images/burger.jpg",
           ("קישור לתמונה", "תמונה", "קישור תמונה", "כתובת תמונה", "תמונה (קישור)", "image", "image url",
            "photo", "picture"), max_len=IMAGE_URL_MAX),
)

CATEGORY_COLUMNS: Tuple[Column, ...] = (
    MARKER,
    Column("name", "שם מחלקה*", 26, "text",
           "שם המחלקה כפי שיופיע בקופה.", "מנות עיקריות",
           ("שם מחלקה", "מחלקה", "שם", "קטגוריה", "name", "category"), max_len=NAME_MAX),
    Column("parent", "מחלקת אב", 22, "text",
           "למחלקת משנה - שם המחלקה שמעליה. ריק = מחלקה ראשית (במחלקה קיימת: ללא שינוי; "
           "'ללא' = להעביר לראשית).", "",
           ("מחלקת אב", "אב", "מחלקה ראשית", "parent")),
    Column("printers", "מדפסות", 24, "printers",
           "לאן יודפסו כל פריטי המחלקה ומחלקות המשנה שלה: שמות מגיליון 'מדפסות' מופרדים "
           "בפסיק. 'ללא' = לא להדפיס; 'ירושה' = לפי מחלקת האב.", "מטבח",
           ("מדפסות", "מדפסת", "מדפסות בונים", "תחנה", "תחנות", "printers", "printer")),
    Column("sort", "סדר", 8, "int",
           "סדר ההצגה בקופה (מספר קטן קודם).", "1",
           ("סדר", "מיון", "סדר הצגה", "sort", "order")),
    Column("active", "פעיל", 8, "bool",
           "לא = המחלקה מוסתרת בקופה. ריק במחלקה חדשה = כן.", YES,
           ("פעיל", "זמין", "active")),
    Column("groups", "קבוצות תוספות", 28, "groups",
           "קבוצות התוספות שכל פריטי המחלקה (ומחלקות המשנה שלה) מקבלים, אלא אם נקבעו לפריט קבוצות "
           "משלו: שמות מגיליון 'קבוצות תוספות' מופרדים בפסיק. ריק = ללא שינוי; 'ללא' = בלי תוספות; "
           "'ירושה' = לפי מחלקת האב.",
           "תוספות להמבורגר",
           ("קבוצות תוספות", "קבוצות", "תוספות", "קבוצת תוספות", "modifiers", "modifier groups", "addons",
            "add-ons")),
    Column("image", "קישור לתמונה", 34, "image",
           "כתובת אינטרנט של תמונת המחלקה (https://...). מורדת ונשמרת בזמן הייבוא. ריק = ללא שינוי; "
           "'ללא' = להסיר את התמונה.",
           "",
           ("קישור לתמונה", "תמונה", "קישור תמונה", "כתובת תמונה", "image", "image url", "photo", "picture"),
           max_len=IMAGE_URL_MAX),
)

GROUP_COLUMNS: Tuple[Column, ...] = (
    MARKER,
    Column("name", "שם קבוצה*", 26, "text",
           "שם הקבוצה כפי שיופיע לקופאי וללקוח בקיוסק (עד 100 תווים), למשל 'מידת עשייה' או "
           "'תוספות להמבורגר'. קבוצה קיימת מזוהה לפי השם.", "תוספות להמבורגר",
           ("שם קבוצה", "קבוצה", "שם הקבוצה", "קבוצת תוספות", "group", "name"), max_len=GROUP_NAME_MAX),
    Column("kind", "סוג", 11, "kind",
           "תוספת = מוסיפים למנה (+ גבינה); בחירה = בוחרים אחת מכמה (מידת עשייה, גודל); "
           "הסרה = מורידים רכיב (בלי בצל, מודגש בבון). ריק בקבוצה חדשה = תוספת.",
           GROUP_KIND_LABELS["addon"], ("סוג", "סוג קבוצה", "סוג הקבוצה", "kind", "type")),
    Column("required", "חובה", 8, "bool",
           "כן = חייבים לבחור לפני שהמנה נכנסת להזמנה (לפחות אחת, או המינימום שבעמודה הבאה). "
           "ריק בקבוצה חדשה = לא.", NO, ("חובה", "חובה לבחור", "required")),
    Column("min", "מינימום בחירות", 11, "int",
           "כמה חייבים לבחור לכל הפחות (0-99). ריק = לפי 'חובה' (כן = 1, לא = 0).", "",
           ("מינימום", "מינימום בחירות", "לפחות", "min"), low=0, high=SELECT_MAX),
    Column("max", "מקסימום בחירות", 12, "limit",
           f"כמה אפשר לבחור לכל היותר (1-99), או '{LIMIT_NONE}'. בקבוצת 'בחירה' בדרך כלל 1. "
           "ריק בקבוצה חדשה = ללא הגבלה.", "3",
           ("מקסימום", "מקסימום בחירות", "לכל היותר", "max"), low=1, high=SELECT_MAX),
    Column("free", "כמה בחינם", 10, "int",
           "כמה מהבחירות לא יחויבו - הזולות ביותר, בלי תלות בסדר הבחירה (0 = הכל בתשלום).", "1",
           ("כמה בחינם", "בחינם", "חינם", "כמות חינם", "free"), low=0, high=SELECT_MAX),
    Column("allow_quantity", "כמות לאפשרות", 11, "bool",
           "כן = אפשר לבחור את אותה אפשרות יותר מפעם אחת (ביצה ×2).", NO,
           ("כמות לאפשרות", "כמות", "אפשר כמה פעמים", "allow quantity")),
    Column("allow_pre", "מעט / הרבה / בצד", 13, "bool",
           "כן = אפשר לסמן לכל אפשרות 'מעט', 'הרבה' (במחיר כפול) או 'בצד'.", NO,
           ("מעט / הרבה / בצד", "מעט הרבה בצד", "מעט/הרבה/בצד", "pre modifiers")),
    Column("sort", "סדר", 8, "int",
           "סדר הקבוצה בחלון התוספות (מספר קטן קודם). ריק בקבוצה חדשה = בסוף.", "1",
           ("סדר", "מיון", "סדר הצגה", "sort", "order"), low=SORT_MIN, high=SORT_MAX),
    Column("active", "פעילה", 8, "bool",
           "לא = הקבוצה לא מוצגת בקופה ובקיוסק. ריק בקבוצה חדשה = כן.", YES,
           ("פעילה", "פעיל", "active")),
)

OPTION_COLUMNS: Tuple[Column, ...] = (
    MARKER,
    Column("group", "קבוצה*", 24, "group",
           "שם הקבוצה מגיליון 'קבוצות תוספות' (או קבוצה שכבר קיימת במערכת).", "תוספות להמבורגר",
           ("קבוצה", "שם קבוצה", "שם הקבוצה", "קבוצת תוספות", "group")),
    Column("name", "אפשרות*", 24, "text",
           "שם האפשרות כפי שיופיע בקופה ובקבלה (עד 100 תווים): 'גבינה צהובה', 'בלי בצל', 'מדיום'. "
           "אפשרות קיימת מזוהה לפי השם בתוך הקבוצה.", "גבינה צהובה",
           ("אפשרות", "שם אפשרות", "שם האפשרות", "תוספת", "שם תוספת", "option", "name"), max_len=OPTION_NAME_MAX),
    Column("price", "מחיר", 10, "signed_money",
           "כמה מתווסף למחיר המנה, בשקלים כולל מע״מ (המע״מ של המנה). 0 או ריק באפשרות חדשה = חינם. "
           "אפשר גם מינוס, עד -1000.", "4",
           ("מחיר", "מחיר אפשרות", "מחיר תוספת", "תוספת מחיר", "price")),
    Column("kitchen_name", "שם למטבח", 16, "text",
           f"שם קצר לבון במטבח (לא חובה, עד 60 תווים). ריק = ללא שינוי; '{CLEAR_WORD}' = לפי השם הרגיל.", "",
           ("שם למטבח", "שם בבון", "שם במטבח", "kitchen name"), max_len=KITCHEN_NAME_MAX),
    Column("default", "מסומנת מראש", 11, "bool",
           "כן = האפשרות מסומנת אוטומטית כשמוסיפים את המנה (הקופאי יכול לבטל).", NO,
           ("מסומנת מראש", "ברירת מחדל", "default")),
    Column("max_qty", "מקסימום לאפשרות", 12, "limit",
           "כמה פעמים לכל היותר אפשר לבחור את האפשרות במנה אחת (מעל 1 דורש 'כמות לאפשרות' = כן "
           f"בקבוצה). ריק = ללא שינוי; '{LIMIT_NONE}' = לפי הקבוצה.", "",
           ("מקסימום לאפשרות", "כמות מקסימלית", "מקסימום", "max qty"), low=1, high=SELECT_MAX),
    Column("sort", "סדר", 8, "int",
           "סדר האפשרות בתוך הקבוצה (מספר קטן קודם). ריק באפשרות חדשה = בסוף.", "1",
           ("סדר", "מיון", "סדר הצגה", "sort", "order"), low=SORT_MIN, high=SORT_MAX),
    Column("active", "פעילה", 8, "bool",
           "לא = האפשרות לא מוצגת. ריק באפשרות חדשה = כן.", YES,
           ("פעילה", "פעיל", "active")),
)

NOTE_COLUMNS: Tuple[Column, ...] = (
    MARKER,
    Column("name", "הערה*", 26, "text",
           "טקסט ההערה, כפי שיופיע ככפתור בקופה ובקיוסק (עד 60 תווים): 'בלי בצל', 'רוטב בצד'.", "בלי בצל",
           ("הערה", "טקסט", "טקסט ההערה", "הערה מהירה", "הערות", "note", "text"), max_len=NOTE_TEXT_MAX),
    Column("all", "לכל המנות", 10, "bool",
           "כן = ההערה מוצעת בכל המנות של החברה.", NO,
           ("לכל המנות", "לכל הפריטים", "כל המנות", "כל הפריטים", "all")),
    Column("categories", "מחלקות", 28, "category_list",
           "המחלקות שההערה מוצעת בהן (לכל הפריטים שבהן ובמחלקות המשנה), מופרדות בפסיק.", "המבורגרים",
           ("מחלקות", "מחלקה", "categories")),
    Column("products", "פריטים", 30, "product_list",
           "פריטים מסוימים שההערה מוצעת בהם, מופרדים בפסיק - לפי השם (או המק״ט / הברקוד, כשיש שני "
           "פריטים באותו שם).", "",
           ("פריטים", "פריט", "מוצרים", "products")),
    Column("important", "חשובה", 8, "bool",
           "כן = כשהקופאי בוחר בהערה היא מודפסת בהדגשה בבון למטבח.", NO,
           ("חשובה", "חשוב", "important")),
    Column("sort", "סדר", 8, "int",
           "מקום ההערה בין הכפתורים (מספר קטן קודם). ריק בהערה חדשה = בסוף.", "",
           ("סדר", "מיון", "sort", "order"), low=SORT_MIN, high=SORT_MAX),
)


def menu_price_column(menu_name: str) -> Column:
    """The products column of one catalog menu's price ("מחיר בתפריט: צהריים")."""
    return Column(
        f"menu:{menu_name}", f"{MENU_PRICE_PREFIX}: {menu_name}", max(14, min(24, len(menu_name) + 14)),
        "menu_price",
        f"המחיר בזמן שהתפריט '{menu_name}' פעיל, בשקלים כולל מע״מ (המע״מ של הפריט). ריק = ללא שינוי "
        f"(פריט חדש: מחיר הקטלוג). '{CLEAR_WORD}' = לבטל את המחיר המיוחד. פריט שעוד לא בתפריט יתווסף אליו.",
        "",
    )


_MENU_HEADER = re.compile(r"^\s*מחיר\s+(?:ב)?תפריט\s*[:：\-–]?\s*(?P<name>.+?)\s*\*?\s*$")


def menu_name_of_header(text: Any) -> Optional[str]:
    """The menu a "מחיר בתפריט: X" header names, or None for any other header."""
    match = _MENU_HEADER.match(clean_text(text))
    if not match:
        return None
    name = match.group("name").strip(" :-–")
    return name or None


def _menu_key(text: Any) -> Optional[str]:
    name = menu_name_of_header(text)
    return f"menu:{name}" if name else None


PRINTER_REFERENCE_HEADERS = ("שם מדפסת", "סניף", "חיבור", "נקודת מכירה / קופה", "פעילה")


def column(columns: Sequence[Column], key: str) -> Column:
    return next(c for c in columns if c.key == key)


# ── Normalizing ───────────────────────────────────────────────────────────────

_BIDI_MARKS = re.compile("[‎‏‪-‮⁦-⁩﻿]")
_SPACES = re.compile(r"\s+")
#: Hebrew points and cantillation.
_NIQQUD = re.compile("[֑-ׇ]")
_QUOTES = "\"'`׳״‘’“”"
_QUOTES_RE = re.compile("[" + re.escape(_QUOTES) + "]")


def clean_text(value: Any) -> str:
    """A cell as one line of text: no bidi marks, single spaces, trimmed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return YES if value else NO
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.strftime("%d/%m/%Y %H:%M") if (value.hour or value.minute) else value.strftime("%d/%m/%Y")
    elif isinstance(value, date):
        value = value.strftime("%d/%m/%Y")
    text = _BIDI_MARKS.sub("", str(value)).replace(" ", " ")
    return _SPACES.sub(" ", text).strip()


def normalize_name(text: Any) -> str:
    """How two names are compared: case, spacing and the kind of quote mark do not count."""
    out = unicodedata.normalize("NFKC", clean_text(text))
    out = out.replace("״", '"').replace("“", '"').replace("”", '"')
    out = out.replace("׳", "'").replace("‘", "'").replace("’", "'").replace("`", "'")
    return out.casefold()


def normalize_header(text: Any) -> str:
    """A header as it is matched: no "*", no quote marks, no "(₪)", single spaces."""
    out = unicodedata.normalize("NFKC", clean_text(text))
    out = _NIQQUD.sub("", out)
    out = out.replace("*", " ").replace("₪", " ").replace("(", " ").replace(")", " ")
    out = _QUOTES_RE.sub("", out)
    return _SPACES.sub(" ", out).strip().casefold()


def _header_keys(columns: Sequence[Column]) -> Dict[str, str]:
    keys: Dict[str, str] = {}
    for col in columns:
        for name in (col.header, col.title) + col.aliases:
            keys.setdefault(normalize_header(name), col.key)
    return keys


# ── Cells ─────────────────────────────────────────────────────────────────────


@dataclass
class Parsed:
    """One cell read: its value (None = empty), and what reading it had to say."""

    value: Any = None
    warning: Optional[str] = None
    error: Optional[str] = None


_YES_WORDS = {"כן", "v", "✓", "✔", "true", "yes", "y", "1", "פעיל", "on"}
_NO_WORDS = {"לא", "false", "no", "n", "0", "-", "לא פעיל", "off"}


def parse_bool(value: Any, title: str) -> Parsed:
    if value is None:
        return Parsed()
    if isinstance(value, bool):
        return Parsed(value)
    if isinstance(value, (int, float)) and value in (0, 1):
        return Parsed(bool(value))
    text = normalize_name(value)
    if not text:
        return Parsed()
    if text in _YES_WORDS:
        return Parsed(True)
    if text in _NO_WORDS:
        return Parsed(False)
    return Parsed(error=f"בעמודה '{title}' צריך לכתוב כן או לא (נכתב '{clean_text(value)}')")


_CURRENCY = re.compile(r"(₪|ש\"ח|ש״ח|שח|nis|ils)", re.IGNORECASE)
_THOUSANDS = re.compile(r"^\d{1,3}(,\d{3})+(\.\d+)?$")
_DECIMAL_COMMA = re.compile(r"^\d+,\d{1,2}$")
_PLAIN_NUMBER = re.compile(r"^\d+(\.\d+)?$")


def parse_money(value: Any, title: str, maximum: Decimal = PRICE_MAX, minimum: Decimal = Decimal("0")) -> Parsed:
    """
    A sum of shekels. A number cell is taken as it is; text may carry "₪", thousands
    commas ("1,234.50") or a decimal comma ("12,50") - read, and reported as read.
    Below `minimum` (0: no negative sums) is an error.
    """
    if value is None or (isinstance(value, str) and not clean_text(value)):
        return Parsed()
    note = None
    if isinstance(value, bool):
        return Parsed(error=f"{title} לא תקין ('{clean_text(value)}')")
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
        except InvalidOperation:
            return Parsed(error=f"{title} לא תקין ('{value}')")
    else:
        raw = clean_text(value)
        text = _CURRENCY.sub("", raw).replace(" ", "")
        negative = text.startswith("-") or (text.startswith("(") and text.endswith(")"))
        text = text.strip("-()")
        if _THOUSANDS.match(text):
            text = text.replace(",", "")
        elif _DECIMAL_COMMA.match(text):
            text = text.replace(",", ".")
        if not _PLAIN_NUMBER.match(text):
            return Parsed(error=f"{title} לא תקין ('{raw}')")
        amount = Decimal(text)
        if negative:
            amount = -amount
        if text != raw.replace(" ", ""):
            note = f"{title} '{raw}' נקרא כ-{_money_text(amount)}"
    if amount.is_nan() or amount.is_infinite():
        return Parsed(error=f"{title} לא תקין ('{clean_text(value)}')")
    if amount < minimum:
        if minimum == 0:
            return Parsed(error=f"{title} לא יכול להיות שלילי ({_money_text(amount)})")
        return Parsed(error=f"{title} נמוך מדי ({_money_text(amount)}, לכל הפחות {_money_text(minimum)})")
    rounded = amount.quantize(CENT, rounding=ROUND_HALF_UP)
    if rounded != amount:
        note = f"{title} {amount} עוגל ל-{_money_text(rounded)}"
    if rounded > maximum:
        return Parsed(error=f"{title} גבוה מדי ({_money_text(rounded)})")
    return Parsed(rounded, warning=note)


def _money_text(amount: Decimal) -> str:
    return f"{amount:.2f}"


_SCIENTIFIC = re.compile(r"^\d+(\.\d+)?e\+?\d+$", re.IGNORECASE)


def parse_code(value: Any, title: str, *, strip_spaces: bool = False) -> Parsed:
    """A barcode or SKU: text, never a number Excel has rounded or shortened."""
    if value is None:
        return Parsed()
    if isinstance(value, bool):
        return Parsed(error=f"{title} לא תקין ('{clean_text(value)}')")
    if isinstance(value, float):
        if not value.is_integer():
            return Parsed(error=f"{title} לא תקין ('{value}')")
        if abs(value) >= 1e15:
            # Past 15 digits Excel has already dropped the end of the number.
            return Parsed(error=f"{title} ארוך מדי לשדה מספרי - עצבו את העמודה כטקסט והקלידו שוב")
        value = int(value)
    text = clean_text(value)
    if not text:
        return Parsed()
    if _SCIENTIFIC.match(text):
        return Parsed(error=f"{title} בפורמט מדעי ({text}) - עצבו את העמודה כטקסט והקלידו שוב")
    if len(text) > CODE_MAX:
        return Parsed(error=f"{title} ארוך מדי (מעל {CODE_MAX} תווים)")
    if strip_spaces:
        text = text.replace(" ", "")
    return Parsed(text)


def parse_int(value: Any, title: str, low: int, high: int, *, allow_clear: bool = False) -> Parsed:
    """A whole number in [low, high]; with `allow_clear`, "ללא" / 0 is an explicit clear ("")."""
    if value is None:
        return Parsed()
    if isinstance(value, bool):
        return Parsed(error=f"{title}: צריך מספר שלם")
    text = clean_text(value)
    if not text:
        return Parsed()
    if allow_clear and normalize_name(text) in (normalize_name(PRINTERS_NONE), "0"):
        return Parsed("")
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return Parsed(error=f"{title}: צריך מספר שלם (נכתב '{text}')")
    if number != number.to_integral_value():
        return Parsed(error=f"{title}: צריך מספר שלם (נכתב '{text}')")
    number = int(number)
    if number < low or number > high:
        return Parsed(error=f"{title}: מספר בין {low} ל-{high} (נכתב {number})")
    return Parsed(number)


def parse_ticket(value: Any, title: str) -> Parsed:
    """An item-ticket mode, or `TICKET_INHERIT`; empty is "no change"."""
    text = normalize_name(value)
    if not text:
        return Parsed()
    for mode, label in TICKET_LABELS.items():
        if text in (normalize_name(label), mode, normalize_name("שובר " + label)):
            return Parsed(mode)
    words = {
        "ללא": "off", "לא": "off", "off": "off", "inherit": TICKET_INHERIT,
        "ירושה": TICKET_INHERIT, "לפי הקטגוריה": TICKET_INHERIT,
        "שובר לכל יחידה": "per_unit", "שובר לכל פריט": "per_line",
        "שובר אחד לעסקה": "per_sale", "לכל שורה": "per_line", "לכל מכירה": "per_sale",
    }
    if text in words:
        return Parsed(words[text])
    choices = ", ".join(TICKET_LABELS.values())
    return Parsed(error=f"{title} לא מוכר ('{clean_text(value)}') - אפשר: {choices}")


@dataclass(frozen=True)
class RoutingSpec:
    """A "מדפסות" cell: keep as is, no ticket, back to inherited, or these printer names."""

    mode: str  # "none" | "inherit" | "names"
    names: Tuple[str, ...] = ()

    def text(self) -> str:
        if self.mode == "none":
            return PRINTERS_NONE
        if self.mode == "inherit":
            return PRINTERS_INHERIT
        return ", ".join(self.names)


_LIST_SPLIT = re.compile(r"[,;\n،|]+")


def split_names(value: Any) -> List[str]:
    """A cell of names separated by commas: each once (as first written), in order."""
    text = clean_text(value) if not isinstance(value, str) else value.strip()
    names: List[str] = []
    seen = set()
    for part in (clean_text(p) for p in _LIST_SPLIT.split(text or "")):
        key = normalize_name(part)
        if key and key not in seen:
            seen.add(key)
            names.append(part)
    return names


def _parse_spec(value: Any, none_word: str, inherit_word: str, what: str) -> Parsed:
    """Names, or alone "ללא" (none) / "ירושה" (inherit); empty is "no change"."""
    names = split_names(value)
    if not names:
        return Parsed()
    keys = [normalize_name(n) for n in names]
    special = {normalize_name(none_word): "none", normalize_name(inherit_word): "inherit"}
    modes = [special[k] for k in keys if k in special]
    if modes:
        if len(names) > 1:
            return Parsed(error=f"'{none_word}' או '{inherit_word}' נכתבים לבד, בלי שמות {what}")
        return Parsed(RoutingSpec(modes[0]))
    return Parsed(RoutingSpec("names", tuple(names)))


def parse_printers(value: Any) -> Parsed:
    return _parse_spec(value, PRINTERS_NONE, PRINTERS_INHERIT, "מדפסות")


def parse_groups(value: Any) -> Parsed:
    """A "קבוצות תוספות" cell: group names in order, or "ללא" / "ירושה" (a `RoutingSpec`)."""
    parsed = _parse_spec(value, GROUPS_NONE, GROUPS_INHERIT, "קבוצות")
    if parsed.value is not None and len(parsed.value.names) > LINKS_MAX:
        return Parsed(error=f"יותר מ-{LINKS_MAX} קבוצות תוספות")
    return parsed


def parse_names(value: Any, title: str, maximum: int = NAME_MAX * 2) -> Parsed:
    """A list cell ("מחלקות", "פריטים"): a tuple of names; empty is None."""
    names = split_names(value)
    if not names:
        return Parsed()
    too_long = [n for n in names if len(n) > maximum]
    if too_long:
        return Parsed(error=f"{title}: '{too_long[0][:40]}…' ארוך מדי")
    return Parsed(tuple(names))


def parse_kind(value: Any) -> Parsed:
    """An add-on group's kind: a Hebrew label or a code; empty is "no change"."""
    text = normalize_name(value)
    if not text:
        return Parsed()
    words = {"addon": "addon", "add-on": "addon", "תוספות": "addon", "choice": "choice", "בחירה אחת": "choice",
             "removal": "removal", "הסרות": "removal", "בלי": "removal"}
    for code, label in GROUP_KIND_LABELS.items():
        words[normalize_name(label)] = code
    if text in words:
        return Parsed(words[text])
    choices = ", ".join(GROUP_KIND_LABELS.values())
    return Parsed(error=f"סוג קבוצה לא מוכר ('{clean_text(value)}') - אפשר: {choices}")


_LIMIT_NONE_WORDS = frozenset(
    normalize_name(w) for w in (LIMIT_NONE, "ללא", "בלי הגבלה", "אין הגבלה", "אין", "unlimited", "none")
)


def parse_limit(value: Any, title: str, low: int, high: int) -> Parsed:
    """
    An upper limit: a whole number in [low, high], or "ללא הגבלה" - read as "" (clear,
    no limit). Empty is "no change".
    """
    text = clean_text(value)
    if not text:
        return Parsed()
    if normalize_name(text) in _LIMIT_NONE_WORDS:
        return Parsed("")
    return parse_int(value, title, low, high)


def parse_clearable_text(value: Any, title: str, maximum: int) -> Parsed:
    """Text, or "ללא" - read as "" (clear). Empty is "no change"."""
    text = clean_text(value)
    if text and normalize_name(text) == normalize_name(CLEAR_WORD):
        return Parsed("")
    return parse_text(value, title, maximum)


def parse_menu_price(value: Any, title: str) -> Parsed:
    """A catalog menu's price, or "ללא" - read as "" (no menu price: the catalog's)."""
    text = clean_text(value)
    if text and normalize_name(text) == normalize_name(CLEAR_WORD):
        return Parsed("")
    return parse_money(value, title, MENU_PRICE_MAX)


_URL = re.compile(r"^https?://[^\s/$.?#][^\s]*$", re.IGNORECASE)


def parse_image(value: Any) -> Parsed:
    """
    A "קישור לתמונה" cell: an http(s) address, or "ללא" - read as "" (remove the picture).
    Empty is "no change".
    """
    text = clean_text(value)
    if not text:
        return Parsed()
    if normalize_name(text) == normalize_name(CLEAR_WORD):
        return Parsed("")
    if text.lower().startswith("www."):
        text = "https://" + text
    if len(text) > IMAGE_URL_MAX:
        return Parsed(error=f"הקישור לתמונה ארוך מדי (מעל {IMAGE_URL_MAX} תווים)")
    if not _URL.match(text):
        return Parsed(error=f"הקישור לתמונה לא תקין ('{text[:60]}') - צריך כתובת שמתחילה ב-https://")
    return Parsed(text)


def parse_text(value: Any, title: str, maximum: int) -> Parsed:
    text = clean_text(value)
    if not text:
        return Parsed()
    if len(text) > maximum:
        return Parsed(error=f"{title} ארוך מדי (מעל {maximum} תווים)")
    return Parsed(text)


#: In "סימוני תזונה": clear the product's tags.
DIETARY_NONE = "ללא"


def parse_dietary(value: Any) -> Parsed:
    """
    A "סימוני תזונה" cell: Hebrew labels or codes, separated by commas, in any order. The
    value is a tuple of codes in the fixed order (app/services/dietary.py); `()` is an
    explicit clear ("ללא"); empty is "no change".
    """
    text = clean_text(value)
    if not text:
        return Parsed()
    parts = [clean_text(p) for p in _LIST_SPLIT.split(text) if clean_text(p)]
    if [normalize_name(p) for p in parts] in ([normalize_name(DIETARY_NONE)], ["none"]):
        return Parsed(())
    codes: List[str] = []
    unknown: List[str] = []
    for part in parts:
        code = dietary.code_of(part)
        if code is None:
            unknown.append(part)
        elif code not in codes:
            codes.append(code)
    if unknown:
        choices = ", ".join(DIETARY_LABELS_HE.values())
        return Parsed(error=f"סימון תזונה לא מוכר ('{unknown[0]}') - אפשר: {choices}")
    pair = dietary.conflict_of(codes + ([dietary.VEGETARIAN] if dietary.VEGAN in codes else []))
    if pair is not None:
        return Parsed(error=f"סימוני תזונה סותרים: {dietary.label(pair[0])} ו{dietary.label(pair[1])}")
    return Parsed(tuple(dietary.clean(codes)))


def parse_channel(value: Any) -> Parsed:
    """A "ערוץ מכירה" cell: a Hebrew label or a code (app/services/sales_channel.py); empty is "no change"."""
    text = clean_text(value)
    if not text:
        return Parsed()
    code = sales_channel.code_of(text)
    if code is None:
        choices = ", ".join(SALES_CHANNEL_LABELS_HE.values())
        return Parsed(error=f"ערוץ מכירה לא מוכר ('{text}') - אפשר: {choices}")
    return Parsed(code)


def is_marked(value: Any) -> bool:
    """Is this "סימון" a row the import skips?"""
    text = normalize_name(value)
    return bool(text) and (text.startswith(normalize_name(EXAMPLE_MARK)) or text.startswith("דוגמא")
                           or text == normalize_name(SKIP_MARK))


_WEIGHT_UNITS = {normalize_name(u) for u in ("ק״ג", "קג", "קילו", "קילוגרם", "גרם", "100 גרם",
                                               "ליטר", "ל׳", "מ״ל", "kg", "g", "l", "lb")}


def is_weight_unit(unit: Optional[str]) -> bool:
    return bool(unit) and normalize_name(unit) in _WEIGHT_UNITS


# ── Reading a file ────────────────────────────────────────────────────────────


class SheetError(Exception):
    """The file as a whole cannot be read; the message is the user's (Hebrew)."""


@dataclass
class RawRow:
    #: The row number the user sees: Excel's row, or the CSV line.
    number: int
    cells: Dict[str, Any]


@dataclass
class RawSheet:
    title: str
    #: column key → its header as written in the file.
    headers: Dict[str, str]
    rows: List[RawRow]
    unknown_headers: List[str] = field(default_factory=list)


@dataclass
class RawFile:
    kind: str  # "xlsx" | "csv"
    products: Optional[RawSheet]
    categories: Optional[RawSheet]
    #: What reading the file had to say (not the "הערות מהירות" sheet: `quick_notes`).
    notes: List[str] = field(default_factory=list)
    groups: Optional[RawSheet] = None
    options: Optional[RawSheet] = None
    quick_notes: Optional[RawSheet] = None

    def sheets(self) -> List[Tuple[str, RawSheet]]:
        """(the sheet's Hebrew name, the sheet) of every sheet read."""
        pairs = (("הפריטים", self.products), ("המחלקות", self.categories), ("קבוצות התוספות", self.groups),
                 ("האפשרויות", self.options), ("ההערות המהירות", self.quick_notes))
        return [(title, sheet) for title, sheet in pairs if sheet is not None]


def read_file(data: bytes, filename: str = "") -> RawFile:
    """XLSX by its zip signature, anything else as CSV text. Raises `SheetError`."""
    if not data:
        raise SheetError("הקובץ ריק")
    if len(data) > MAX_FILE_BYTES:
        raise SheetError("הקובץ גדול מדי (עד 5MB)")
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04":
        return _read_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or name.endswith(".xls"):
        raise SheetError("זה קובץ Excel ישן (xls). פתחו אותו ושמרו בשם כ-'חוברת עבודה של Excel (xlsx)'")
    return _read_csv(data)


Dynamic = Optional[Any]  # Callable[[Any], Optional[str]]: a header no column names → a key


def _find_header(rows: List[Sequence[Any]], columns: Sequence[Column],
                 dynamic: Dynamic = None) -> Optional[Tuple[int, Dict[int, str], List[str]]]:
    """(row index, {column index: key}, unknown headers) of the first row that names the sheet's columns."""
    keys = _header_keys(columns)
    name_key = "name"
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping: Dict[int, str] = {}
        unknown: List[str] = []
        for position, cell in enumerate(list(row)[:MAX_COLUMNS]):
            header = normalize_header(cell)
            if not header:
                continue
            key = keys.get(header)
            if key is None and dynamic is not None:
                key = dynamic(cell)
            if key and key not in mapping.values():
                mapping[position] = key
            else:
                unknown.append(clean_text(cell))
        if name_key in mapping.values() and len(mapping) >= 2:
            return index, mapping, unknown
    return None


def _sheet_from_rows(title: str, rows: List[Sequence[Any]], columns: Sequence[Column], first_number: int = 1,
                     dynamic: Dynamic = None) -> Optional[RawSheet]:
    found = _find_header(rows, columns, dynamic)
    if found is None:
        return None
    header_index, mapping, unknown = found
    header_row = list(rows[header_index])
    out: List[RawRow] = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 1):
        values = list(row)
        cells = {key: (values[pos] if pos < len(values) else None) for pos, key in mapping.items()}
        if all(clean_text(v) == "" for v in cells.values()):
            continue
        out.append(RawRow(number=offset + first_number, cells=cells))
        if len(out) > MAX_DATA_ROWS:
            raise SheetError(f"בגיליון '{title}' יותר מ-{MAX_DATA_ROWS} שורות - פצלו לכמה קבצים")
    return RawSheet(
        title=title,
        headers={key: clean_text(header_row[pos]) for pos, key in mapping.items()},
        rows=out,
        unknown_headers=[u for u in unknown if u],
    )


def _xlsx_rows(ws, notes: List[str]) -> List[Tuple[Any, ...]]:
    """The sheet's rows (an empty one as `()`, so numbering holds), at most MAX_SHEET_ROWS."""
    rows: List[Tuple[Any, ...]] = []
    filled = 0
    for row in ws.iter_rows(values_only=True, max_col=MAX_COLUMNS, max_row=MAX_SHEET_ROWS):
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            rows.append(())
            continue
        rows.append(tuple(row))
        filled += 1
        if filled > MAX_DATA_ROWS + HEADER_SEARCH_ROWS:
            raise SheetError(f"בגיליון '{ws.title}' יותר מ-{MAX_DATA_ROWS} שורות - פצלו לכמה קבצים")
    while rows and rows[-1] == ():
        rows.pop()
    if (ws.max_row or 0) > MAX_SHEET_ROWS:
        notes.append(f"בגיליון '{ws.title}' נקראו רק {MAX_SHEET_ROWS:,} השורות הראשונות")
    return rows


def _read_xlsx(data: bytes) -> RawFile:
    import openpyxl

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            # A zip bomb is refused before openpyxl inflates it.
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise SheetError("הקובץ גדול מדי")
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except SheetError:
        raise
    except Exception:  # noqa: BLE001 - a broken upload is the user's file, not our bug
        raise SheetError("לא הצלחנו לפתוח את הקובץ - ודאו שזה קובץ Excel (xlsx) תקין")
    try:
        sheets = {normalize_header(ws.title): ws for ws in workbook.worksheets}
        notes: List[str] = []

        def read(title: str, columns: Sequence[Column], dynamic: Dynamic = None) -> Optional[RawSheet]:
            ws = sheets.get(normalize_header(title))
            if ws is None:
                return None
            return _sheet_from_rows(title, _xlsx_rows(ws, notes), columns, dynamic=dynamic)

        products_ws = sheets.get(normalize_header(SHEET_PRODUCTS))
        products = read(SHEET_PRODUCTS, PRODUCT_COLUMNS, _menu_key)
        categories = read(SHEET_CATEGORIES, CATEGORY_COLUMNS)
        groups = read(SHEET_GROUPS, GROUP_COLUMNS)
        options = read(SHEET_OPTIONS, OPTION_COLUMNS)
        quick_notes = read(SHEET_NOTES, NOTE_COLUMNS)
        if products is None:
            # No "פריטים" sheet (or no header in it): any sheet whose header is a product's.
            skip = {normalize_header(n) for n in (SHEET_CATEGORIES, SHEET_INSTRUCTIONS, SHEET_PRINTERS, SHEET_LISTS,
                                                  SHEET_GROUPS, SHEET_OPTIONS, SHEET_NOTES)}
            for title, ws in sheets.items():
                if title in skip or ws is products_ws:
                    continue
                candidate = _sheet_from_rows(ws.title, _xlsx_rows(ws, notes), PRODUCT_COLUMNS, dynamic=_menu_key)
                if candidate is not None and "price" in candidate.headers:
                    products = candidate
                    break
        if all(s is None for s in (products, categories, groups, options, quick_notes)):
            raise SheetError(
                f"לא נמצא גיליון '{SHEET_PRODUCTS}' עם שורת כותרות (שם פריט, מחלקה, מחיר) - "
                "השתמשו בטמפלט שהורדתם"
            )
        return RawFile(kind="xlsx", products=products, categories=categories, notes=notes,
                       groups=groups, options=options, quick_notes=quick_notes)
    finally:
        workbook.close()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1255"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SheetError("לא הצלחנו לקרוא את הקובץ - שמרו אותו כ-CSV UTF-8 או כ-xlsx")


def _read_csv(data: bytes) -> RawFile:
    text = _decode(data)
    if "\x00" in text:
        raise SheetError("הקובץ אינו קובץ CSV או Excel")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: List[List[str]] = []
    for row in reader:
        rows.append(row[:MAX_COLUMNS])
        if len(rows) > MAX_DATA_ROWS + HEADER_SEARCH_ROWS:
            raise SheetError(f"בקובץ יותר מ-{MAX_DATA_ROWS} שורות - פצלו לכמה קבצים")
    # CSV cells are text; an empty one is None, as an empty Excel cell is.
    cleaned = [[cell if cell.strip() else None for cell in row] for row in rows]
    products = _sheet_from_rows(SHEET_PRODUCTS, cleaned, PRODUCT_COLUMNS, dynamic=_menu_key)
    if products is not None and "price" in products.headers:
        return RawFile(kind="csv", products=products, categories=None)
    categories = _sheet_from_rows(SHEET_CATEGORIES, cleaned, CATEGORY_COLUMNS)
    if categories is not None and "price" not in categories.headers:
        return RawFile(kind="csv", products=None, categories=categories)
    raise SheetError("לא נמצאה שורת כותרות (שם פריט, מחלקה, מחיר) - השתמשו בכותרות של הטמפלט")
