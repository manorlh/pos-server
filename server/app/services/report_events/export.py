"""
The event report as an Excel workbook for the producer (docs/SPEC_EVENTS.md §6.1).

Sheets: סיכום, תובנות, קופות, פריטים, פילוח, ציר זמן, משמרות, התאמות — every one
right-to-left, with the blue frozen header of the catalog template
(app/services/catalog_template.py), money as `#,##0.00`, data bars on the takings, the
insights coloured by level, and native Excel charts (the timeline, the tills, the tenders).

Drawn from the report dict alone — the live one for a draft, the frozen snapshot for a
confirmed event — so an export after confirmation is exactly what was confirmed.
"""
from __future__ import annotations

import io
from collections import OrderedDict
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.catalog_template import BLUE, GREEN, GREY, GROUPED, HAIRLINE, INK, ORANGE, XLSX_MEDIA_TYPE

from .report import zone
from .rules import ils

__all__ = ["build_workbook", "XLSX_MEDIA_TYPE", "file_name"]

RED = "FF3B30"
PURPLE = "5E5CE6"
_HEADER_FILL = PatternFill("solid", fgColor=BLUE)
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(color=INK, bold=True, size=18)
_SUBTITLE_FONT = Font(color=GREY, size=11)
_SECTION_FONT = Font(color=BLUE, bold=True, size=13)
_THIN = Side(style="thin", color=HAIRLINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
_RIGHT = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
_TOTAL_FILL = PatternFill("solid", fgColor=GROUPED)
_LEVEL_FILL = {
    "alert": PatternFill("solid", fgColor="FFE5E3"),
    "warning": PatternFill("solid", fgColor="FFF3E0"),
    "info": PatternFill("solid", fgColor="E8F2FF"),
}
_LEVEL_TEXT = {"alert": "התראה", "warning": "אזהרה", "info": "מידע"}
_LEVEL_FONT = {"alert": Font(color=RED, bold=True), "warning": Font(color="C93400", bold=True), "info": Font(color=BLUE, bold=True)}
_FORMATS = {"money": "#,##0.00", "int": "#,##0", "qty": "#,##0.###", "pct": "0.0\"%\"", "time": "dd/mm/yyyy hh:mm", "text": "@"}
_METHOD = {"cash": "מזומן", "card": "אשראי", "other": "אחר", "exchange": "החלפה (סל מעורב)"}
_STATUS = {"draft": "טיוטה", "confirmed": "אושר", "open": "פתוחה", "closed": "סגורה",
           "match": "תואם", "mismatch": "אי-התאמה", "legacy": "לא ניתן להשוות", "pending": "ממתין",
           "none": "אין נתונים", "success": "הצליח", "failed": "נכשל", "unknown": "לא ידוע", "n/a": "—"}

Col = Tuple[str, str, int]  # header, kind, width


def file_name(report: Dict[str, Any]) -> str:
    ev = report["event"]
    safe = "".join(ch for ch in (ev.get("name") or "event") if ch not in '\\/:*?"<>|').strip() or "event"
    return f"דוח אירוע - {safe} - {ev.get('startDate')}.xlsx"


class _Clock:
    def __init__(self, tz_name: Optional[str]):
        self.tz = zone(tz_name)

    def local(self, value: Optional[str]) -> Optional[datetime]:
        if not value:
            return None
        return datetime.fromisoformat(value).astimezone(self.tz).replace(tzinfo=None)

    def text(self, value: Optional[str], fmt: str = "%d/%m/%Y %H:%M") -> str:
        moment = self.local(value)
        return moment.strftime(fmt) if moment else "—"


def _sheet(wb: Workbook, title: str, tab: str):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _table(ws, row: int, columns: Sequence[Col], rows: Sequence[Sequence[Any]], *, title: Optional[str] = None,
           col: int = 1, totals: Optional[Sequence[Any]] = None, freeze: bool = False,
           fills: Optional[Sequence[Optional[PatternFill]]] = None) -> Tuple[int, int, int]:
    """Write a titled table at `row`; (header row, last data row, next free row)."""
    if title:
        ws.cell(row=row, column=col, value=title).font = _SECTION_FONT
        row += 1
    header = row
    for i, (name, _kind, width) in enumerate(columns):
        c = ws.cell(row=row, column=col + i, value=name)
        c.fill, c.font, c.alignment, c.border = _HEADER_FILL, _HEADER_FONT, _CENTER, _BOX
        letter = get_column_letter(col + i)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
    ws.row_dimensions[row].height = 30
    for r_index, values in enumerate(rows):
        row += 1
        for i, (_name, kind, _w) in enumerate(columns):
            value = values[i] if i < len(values) else None
            c = ws.cell(row=row, column=col + i, value=value)
            c.border = _BOX
            if kind in _FORMATS and kind != "text":
                c.number_format = _FORMATS[kind]
            if kind == "text":
                c.alignment = _RIGHT
            if fills and fills[r_index] is not None:
                c.fill = fills[r_index]
    last = row
    if totals is not None:
        row += 1
        for i, (_name, kind, _w) in enumerate(columns):
            value = totals[i] if i < len(totals) else None
            c = ws.cell(row=row, column=col + i, value=value)
            c.fill, c.font, c.border = _TOTAL_FILL, Font(bold=True), _BOX
            if kind in _FORMATS and kind != "text":
                c.number_format = _FORMATS[kind]
    if freeze:
        ws.freeze_panes = ws.cell(row=header + 1, column=col)
        if rows:
            ws.auto_filter.ref = f"{get_column_letter(col)}{header}:{get_column_letter(col + len(columns) - 1)}{last}"
    return header, last, row + 2


def _data_bar(ws, column: int, first: int, last: int, color: str = BLUE) -> None:
    if last < first:
        return
    letter = get_column_letter(column)
    ws.conditional_formatting.add(
        f"{letter}{first}:{letter}{last}",
        DataBarRule(start_type="min", end_type="max", color=color, showValue=True),
    )


def _bar_chart(title: str, y_title: str, data_ref, cats_ref, color: str = BLUE, width: float = 22, height: float = 9) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = y_title
    chart.legend = None
    chart.width, chart.height = width, height
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.series[0].graphicalProperties.solidFill = color
    chart.series[0].graphicalProperties.line.solidFill = color
    chart.gapWidth = 40
    return chart


# ── Sheets ────────────────────────────────────────────────────────────────────


def _summary(wb: Workbook, report: Dict[str, Any], clock: _Clock) -> None:
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = BLUE
    ev, k = report["event"], report["kpis"]
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 22
    ws["A1"] = ev["name"]
    ws["A1"].font = _TITLE_FONT
    ws.row_dimensions[1].height = 30
    status = _STATUS.get(ev["status"], ev["status"])
    if ev.get("confirmedAt"):
        status += f" · {clock.text(ev['confirmedAt'])}" + (f" ע״י {ev['confirmedBy']}" if ev.get("confirmedBy") else "")
    lines = [
        ("סניף", ev.get("shopName") or "—"),
        ("מפיק", ev.get("producerName") or "—"),
        ("חלון האירוע", f"{clock.text(ev['startsAt'])} — {clock.text(ev['endsAt'])}"),
        ("קופות", ", ".join(m["name"] for m in ev.get("machines", [])) or "—"),
        ("סטטוס", status),
        ("הופק", clock.text(report.get("generatedAt"))),
    ]
    if ev.get("confirmNote"):
        lines.append(("הערת אישור", ev["confirmNote"]))
    if ev.get("notes"):
        lines.append(("הערות", ev["notes"]))
    row = 2
    for label, value in lines:
        ws.cell(row=row, column=1, value=label).font = _SUBTITLE_FONT
        c = ws.cell(row=row, column=2, value=value)
        c.alignment = _RIGHT
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        row += 1
    if report.get("frozen"):
        c = ws.cell(row=row, column=1, value="✔ דוח מאושר — נתונים מוקפאים ברגע האישור")
        c.font = Font(color=GREEN, bold=True)
        row += 1
    row += 1

    peak = k.get("peakHour")
    kpis = [
        ("נטו (אחרי זיכויים)", k["net"], "money"),
        ("נטו ללא מע״מ", k["netExVat"], "money"),
        ("מע״מ" + (" (חלקו מוערך)" if k.get("vatEstimatedCount") else ""), k["vat"], "money"),
        ("מכירות ברוטו", k["gross"], "money"),
        ("הנחות", k["discounts"], "money"),
        ("נגבה במכירות", k["sales"], "money"),
        ("זיכויים", k["refunds"], "money"),
        ("מספר זיכויים", k["refundsCount"], "int"),
        ("מספר מכירות", k["salesCount"], "int"),
        ("ממוצע לעסקה", k["avgTicket"], "money"),
        ("פריטים שנמכרו", k["itemsSold"], "qty"),
        ("פריטים לעסקה", k.get("itemsPerSale"), "qty"),
        ("טיפים", k["tips"], "money"),
        ("אחוז טיפ", k.get("tipPct"), "pct"),
        ("חריגות", k["exceptionsCount"], "int"),
    ]
    kpis2 = [
        ("מזומן", k["cash"], "money"),
        ("אשראי", k["card"], "money"),
        ("אחר", k["other"], "money"),
        ("מכירות לשעה (חלון)", k["salesPerHour"], "money"),
        ("ממוצע לשעה פעילה", k["avgPerActiveHour"], "money"),
        ("שעות פעילות", k["activeHours"], "int"),
        ("שעת השיא", f"{clock.text(peak['start'], '%d/%m %H:%M')} · {ils(peak['net'])} · {peak['sharePct'] or 0:.1f}%" if peak else "—", "text"),
        ("קופות פעילות", f"{k['activeTills']} מתוך {k['tillsCount']}", "text"),
        ("מכירה ראשונה", clock.local(k.get("firstSaleAt")), "time"),
        ("מכירה אחרונה", clock.local(k.get("lastSaleAt")), "time"),
        ("ממוצע לעסקה רגיל בסניף (28 יום)", k.get("baselineAvgTicket"), "money"),
    ]
    _, _, after = _table(ws, row, [("מדד", "text", 30), ("ערך", "money", 26)],
                         [(label, value) for label, value, _kind in kpis], title="מדדי ראש")
    for i, (_l, _v, kind) in enumerate(kpis):
        ws.cell(row=row + 2 + i, column=2).number_format = _FORMATS.get(kind, "General")
    _table(ws, row, [("מדד", "text", 30), ("ערך", "money", 22)],
           [(label, value) for label, value, _kind in kpis2], title="תשלום וזמן", col=4)
    for i, (_l, _v, kind) in enumerate(kpis2):
        ws.cell(row=row + 2 + i, column=5).number_format = _FORMATS.get(kind, "General")

    counts = {lvl: sum(1 for i in report["insights"] if i["level"] == lvl) for lvl in ("alert", "warning", "info")}
    row = after
    ws.cell(row=row, column=1, value="תובנות").font = _SECTION_FONT
    for i, lvl in enumerate(("alert", "warning", "info")):
        c = ws.cell(row=row + 1 + i, column=1, value=_LEVEL_TEXT[lvl])
        c.fill, c.font = _LEVEL_FILL[lvl], _LEVEL_FONT[lvl]
        ws.cell(row=row + 1 + i, column=2, value=counts[lvl])
    row += 5
    rec = report.get("reconciliation") or {}
    ws.cell(row=row, column=1, value="התאמה מול Z").font = _SUBTITLE_FONT
    ws.cell(row=row, column=2, value=_STATUS.get((rec.get("z") or {}).get("status"), "—"))
    ws.cell(row=row + 1, column=1, value="התאמה מול שידורי אשראי").font = _SUBTITLE_FONT
    ws.cell(row=row + 1, column=2, value=_STATUS.get((rec.get("transmissions") or {}).get("status"), "—"))
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _insights(wb: Workbook, report: Dict[str, Any]) -> None:
    ws = _sheet(wb, "תובנות", RED)
    rows, fills = [], []
    for i in report["insights"]:
        ref = i.get("ref") or {}
        rows.append((_LEVEL_TEXT.get(i["level"], i["level"]), i["text"], ref.get("name") or "", i["code"]))
        fills.append(_LEVEL_FILL.get(i["level"]))
    header, last, _ = _table(ws, 1, [("רמה", "text", 10), ("תובנה", "text", 110), ("הפניה", "text", 22), ("קוד", "text", 18)],
                             rows, freeze=True, fills=fills)
    for r in range(header + 1, last + 1):
        lvl = next((k for k, v in _LEVEL_TEXT.items() if v == ws.cell(row=r, column=1).value), None)
        if lvl:
            ws.cell(row=r, column=1).font = _LEVEL_FONT[lvl]
        ws.row_dimensions[r].height = 32


def _tills(wb: Workbook, report: Dict[str, Any], clock: _Clock) -> None:
    ws = _sheet(wb, "קופות", GREEN)
    cols: List[Col] = [
        ("קופה", "text", 20), ("מס׳", "text", 6), ("אזור", "text", 12), ("נטו", "money", 14), ("נתח %", "pct", 9),
        ("מכירות", "int", 9), ("ממוצע לעסקה", "money", 12), ("פריטים", "qty", 9), ("זיכויים", "int", 9),
        ("זיכויים ₪", "money", 12), ("טיפים", "money", 11), ("% טיפ", "pct", 8), ("מכירות לשעה", "money", 13),
        ("ממוצע לשעה פעילה", "money", 13), ("דקות פעילות", "int", 10), ("מכירה ראשונה", "time", 16),
        ("מכירה אחרונה", "time", 16), ("דקות שקט", "int", 9), ("חריגות", "int", 8), ("ביטולי שורה", "int", 9),
        ("סלים מבוטלים", "int", 9), ("פתיחות מגירה", "int", 9), ("מזומן", "money", 12), ("אשראי", "money", 12),
        ("אחר", "money", 11), ("סימון", "text", 18),
    ]
    rows = []
    for t in report["tills"]:
        marks = [m for m, on in (("ללא מכירות", t.get("noSales")), ("חלשה", t.get("weak")), ("שקטה", t.get("idle"))) if on]
        rows.append((
            t["name"], t.get("posNumber"), t.get("areaName"), t["net"], t.get("sharePct"), t["salesCount"], t["avgTicket"],
            t["items"], t["refundsCount"], t["refunds"], t["tips"], t.get("tipPct"), t["salesPerHour"],
            t["avgPerActiveHour"], t["activeMinutes"], clock.local(t.get("firstSaleAt")), clock.local(t.get("lastSaleAt")),
            t.get("idleMinutes", 0), t["exceptions"], t["voids"], t["cancels"], t["drawerOpens"], t["cash"], t["card"],
            t["other"], " · ".join(marks),
        ))
    k = report["kpis"]
    totals = ["סה״כ", None, None, k["net"], 100.0 if k["net"] else None, k["salesCount"], k["avgTicket"], k["itemsSold"],
              k["refundsCount"], k["refunds"], k["tips"], k.get("tipPct"), k["salesPerHour"], k["avgPerActiveHour"],
              None, clock.local(k.get("firstSaleAt")), clock.local(k.get("lastSaleAt")), None, k["exceptionsCount"],
              None, None, None, k["cash"], k["card"], k["other"], None]
    header, last, after = _table(ws, 1, cols, rows, totals=totals, freeze=True)
    _data_bar(ws, 4, header + 1, last, GREEN)
    _data_bar(ws, 13, header + 1, last, PURPLE)
    if rows:
        chart = _bar_chart("נטו לפי קופה", "₪", Reference(ws, min_col=4, min_row=header, max_row=last),
                           Reference(ws, min_col=1, min_row=header + 1, max_row=last), GREEN)
        ws.add_chart(chart, f"A{after + 1}")


def _items(wb: Workbook, report: Dict[str, Any]) -> None:
    ws = _sheet(wb, "פריטים", ORANGE)
    items = report["items"]
    till_names = {t["machineId"]: t["name"] for t in report["tills"]}
    till_ids = [m for m in items["matrix"]["machineIds"] if m in till_names]
    cols: List[Col] = [("דירוג", "int", 7), ("פריט", "text", 28), ("קטגוריה", "text", 16), ("כמות נטו", "qty", 10),
                       ("נמכרו", "qty", 9), ("הוחזרו", "qty", 9), ("הכנסה נטו", "money", 13), ("נתח %", "pct", 9)]
    cols += [(f"כמות · {till_names[m]}", "qty", 12) for m in till_ids]
    rows = [
        (r["rank"], r["name"], r.get("categoryName"), r["quantity"], r["sold"], r["refunded"], r["revenue"], r.get("sharePct"),
         *[r["byTill"].get(m) for m in till_ids])
        for r in items["rows"]
    ]
    totals = ["סה״כ", None, None, items["itemsSold"], None, None, items["totalRevenue"], 100.0 if items["rows"] else None]
    header, last, _ = _table(ws, 1, cols, rows, totals=totals, freeze=True)
    _data_bar(ws, 7, header + 1, last, ORANGE)


def _segments(wb: Workbook, report: Dict[str, Any], clock: _Clock) -> None:
    ws = _sheet(wb, "פילוח", PURPLE)
    seg = report["segments"]
    row = 1
    h, last, row = _table(ws, row, [("קטגוריה", "text", 26), ("כמות", "qty", 12), ("הכנסה", "money", 14), ("נתח %", "pct", 10)],
                          [(c["name"], c["quantity"], c["revenue"], c.get("sharePct")) for c in seg["byCategory"]],
                          title="לפי קטגוריה")
    _data_bar(ws, 3, h + 1, last, ORANGE)
    pay_header, pay_last, row = _table(
        ws, row, [("אמצעי תשלום", "text", 26), ("מסמכים", "int", 12), ("סכום", "money", 14), ("נתח %", "pct", 10)],
        [(_METHOD.get(p["method"], p["method"]), p["count"], p["amount"], p.get("sharePct")) for p in seg["byPayment"]],
        title="לפי אמצעי תשלום")
    if seg["byPayment"]:
        pie = PieChart()
        pie.title = "אמצעי תשלום"
        pie.add_data(Reference(ws, min_col=3, min_row=pay_header, max_row=pay_last), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=pay_header + 1, max_row=pay_last))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.width, pie.height = 12, 8
        ws.add_chart(pie, f"G{pay_header}")
    if seg.get("byCardBrand"):
        _, _, row = _table(ws, row, [("מותג כרטיס", "text", 26), ("עסקאות", "int", 12), ("סכום", "money", 14)],
                           [(b["brand"], b["count"], b["amount"]) for b in seg["byCardBrand"]], title="אשראי לפי מותג")
    h, last, row = _table(ws, row, [("שעה", "text", 26), ("מכירות", "int", 12), ("נטו", "money", 14), ("נתח %", "pct", 10)],
                          [(f"{b['hour']:02d}:00–{(b['hour'] + 1) % 24:02d}:00", b["count"], b["net"], b.get("sharePct"))
                           for b in seg["byHour"]], title="לפי שעה ביום")
    _data_bar(ws, 3, h + 1, last, BLUE)
    _, _, row = _table(ws, row, [("קופה", "text", 26), ("מכירות", "int", 12), ("נטו", "money", 14), ("נתח %", "pct", 10)],
                       [(t["name"], t["count"], t["net"], t.get("sharePct")) for t in seg["byTill"]], title="לפי קופה")
    h, last, row = _table(
        ws, row,
        [("קופאי / מלצר", "text", 26), ("מכירות", "int", 12), ("נטו", "money", 14), ("נתח %", "pct", 10),
         ("ממוצע לעסקה", "money", 13), ("זיכויים", "int", 10), ("זיכויים ₪", "money", 12), ("טיפים", "money", 12), ("% טיפ", "pct", 9)],
        [(c["name"] or "ללא שם", c["salesCount"], c["net"], c.get("sharePct"), c["avgTicket"], c["refundsCount"],
          c["refunds"], c["tips"], c.get("tipPct")) for c in seg["byCashier"]],
        title="לפי קופאי / מלצר")
    _data_bar(ws, 3, h + 1, last, PURPLE)
    if seg.get("byArea"):
        _, _, row = _table(ws, row, [("אזור", "text", 26), ("מכירות", "int", 12), ("נטו", "money", 14), ("נתח %", "pct", 10)],
                           [(a["name"], a["count"], a["net"], a.get("sharePct")) for a in seg["byArea"]], title="לפי אזור")
    mods = report["items"].get("modifiers") or []
    if mods:
        _, _, row = _table(ws, row, [("תוספת / שינוי", "text", 26), ("קבוצה", "text", 16), ("כמות", "qty", 12), ("הכנסה", "money", 14)],
                           [(m["name"], m.get("groupName"), m["quantity"], m["revenue"]) for m in mods], title="תוספות מובילות")
    ex = report["exceptions"]["byType"]
    if ex:
        _table(ws, row, [("סוג חריגה", "text", 26), ("כמות", "int", 12), ("סכום", "money", 14), ("קופות", "text", 40)],
               [(e["label"], e["count"], e["amount"], ", ".join(f"{n} ({c})" for n, c in e["tills"])) for e in ex],
               title="חריגות לפי סוג")


def rebucket(timeline: Sequence[Dict[str, Any]], minutes: int, base: int = 15) -> List[Dict[str, Any]]:
    """15-minute buckets → `minutes` (a multiple of 15), aligned to the clock."""
    if minutes <= base:
        return list(timeline)
    out: "OrderedDict[int, Dict[str, Any]]" = OrderedDict()
    seconds = minutes * 60
    for b in timeline:
        epoch = int(datetime.fromisoformat(b["start"]).timestamp())
        key = epoch - epoch % seconds
        agg = out.setdefault(key, {"start": datetime.fromtimestamp(key, tz=datetime.fromisoformat(b["start"]).tzinfo).isoformat(),
                                   "net": 0.0, "count": 0, "byTill": {}})
        agg["net"] = round(agg["net"] + b["net"], 2)
        agg["count"] += b["count"]
        for mid, v in b["byTill"].items():
            agg["byTill"][mid] = round(agg["byTill"].get(mid, 0.0) + v, 2)
    return list(out.values())


def _timeline(wb: Workbook, report: Dict[str, Any], clock: _Clock, bucket: int) -> None:
    ws = _sheet(wb, "ציר זמן", BLUE)
    buckets = rebucket(report["timeline"], bucket, report.get("bucketMinutes", 15))
    tills = [t for t in report["tills"] if t["documentsCount"]]
    cols: List[Col] = [("התחלת דלי", "time", 18), ("נטו", "money", 13), ("מכירות", "int", 9)]
    cols += [(t["name"], "money", 12) for t in tills]
    rows = [
        (clock.local(b["start"]), b["net"], b["count"], *[b["byTill"].get(t["machineId"], 0) for t in tills])
        for b in buckets
    ]
    ws.cell(row=1, column=1, value=f"ציר זמן — דליים של {bucket} דקות").font = _SECTION_FONT
    header, last, after = _table(ws, 2, cols, rows, freeze=True,
                                 totals=["סה״כ", report["kpis"]["net"], report["kpis"]["salesCount"],
                                         *[t["net"] for t in tills]])
    _data_bar(ws, 2, header + 1, last, BLUE)
    if rows:
        chart = _bar_chart(f"נטו לכל {bucket} דקות", "₪", Reference(ws, min_col=2, min_row=header, max_row=last),
                           Reference(ws, min_col=1, min_row=header + 1, max_row=last), BLUE, width=30)
        chart.x_axis.number_format = "hh:mm"
        ws.add_chart(chart, f"E{after}")


def _shifts(wb: Workbook, report: Dict[str, Any], clock: _Clock) -> None:
    ws = _sheet(wb, "משמרות", GREY)
    cols: List[Col] = [
        ("קופה", "text", 18), ("משמרת", "int", 8), ("סטטוס", "text", 9), ("נפתחה", "time", 16), ("נפתחה ע״י", "text", 14),
        ("נסגרה", "time", 16), ("נסגרה ע״י", "text", 14), ("מזומן פתיחה", "money", 12), ("מזומן צפוי", "money", 12),
        ("מזומן נספר", "money", 12), ("הפרש", "money", 10), ("Z", "text", 14), ("מסמכים בחלון", "int", 11),
        ("נטו בחלון", "money", 12), ("הערה", "text", 34),
    ]
    rows, fills = [], []
    for s in report["shifts"]:
        notes = []
        if s["status"] == "open":
            notes.append("משמרת פתוחה")
        if s.get("openedBeforeWindow"):
            notes.append("נפתחה לפני האירוע")
        if s.get("closesAfterWindow"):
            notes.append("נסגרה אחרי האירוע")
        if not s.get("zReportId") and s["status"] != "open":
            notes.append("טרם נכללה ב-Z")
        rows.append((s["machineName"], s.get("sequence"), _STATUS.get(s["status"], s["status"]), clock.local(s["openedAt"]),
                     s.get("openedBy"), clock.local(s.get("closedAt")), s.get("closedBy"), s.get("openingCash"),
                     s.get("expectedCash"), s.get("countedCash"), s.get("discrepancy"), s.get("zLabel"),
                     s["documentsInWindow"], s["netInWindow"], " · ".join(notes)))
        diff = s.get("discrepancy")
        fills.append(_LEVEL_FILL["alert"] if diff and abs(diff) >= 20 else _LEVEL_FILL["warning"] if s["status"] == "open" else None)
    _table(ws, 1, cols, rows, freeze=True, fills=fills)


def _reconciliation(wb: Workbook, report: Dict[str, Any], clock: _Clock) -> None:
    ws = _sheet(wb, "התאמות", "AF52DE")
    rec = report.get("reconciliation") or {}
    z = rec.get("z") or {"rows": [], "tills": [], "zReports": []}
    tx = rec.get("transmissions") or {"tills": []}
    row = 1
    ws.cell(row=row, column=1, value=f"התאמה מול Z — {_STATUS.get(z.get('status'), '—')}").font = _SECTION_FONT
    row += 1
    _, _, row = _table(
        ws, row,
        [("קופה", "text", 20), ("נטו באירוע", "money", 13), ("מסמכים ב-Z", "int", 11), ("נטו ב-Z", "money", 13),
         ("ממתינים ל-Z", "int", 11), ("נטו ממתין", "money", 13), ("Zים", "text", 30)],
        [(t["name"], t["eventNet"], t["inZ"]["count"], t["inZ"]["net"], t["pendingZ"]["count"], t["pendingZ"]["net"],
          ", ".join(t["zLabels"])) for t in z["tills"]],
        title="לפי קופה")
    if z.get("zReports"):
        _, _, row = _table(
            ws, row,
            [("Z", "text", 20), ("נסגר", "time", 16), ("תקופה מ-", "time", 16), ("עד", "time", 16), ("קופות באירוע", "text", 30),
             ("קופות נוספות ב-Z", "text", 24), ("מסמכים מאוחרים", "int", 11), ("מסמכים ששונו", "int", 11)],
            [(r["label"], clock.local(r.get("closedAt")), clock.local(r.get("periodStart")), clock.local(r.get("periodEnd")),
              ", ".join(r["tills"]), ", ".join(r["otherTills"]), r["lateDocuments"], r["amendedDocuments"]) for r in z["zReports"]],
            title="דוחות Z שחופפים לאירוע")
    labels = (("sales", "מכירות"), ("refunds", "זיכויים"), ("cash", "מזומן"), ("card", "אשראי"), ("other", "אחר"),
              ("tips", "טיפים"), ("count", "מסמכים"))
    rows, fills = [], []
    for r in z["rows"]:
        for key, label in labels:
            zv = (r.get("z") or {}).get(key)
            dv = (r.get("diff") or {}).get(key)
            rows.append((r["label"], r["machineName"], label, zv, r["inWindow"][key], r["outsideWindow"][key], dv,
                         _STATUS.get(r["status"], r["status"])))
            fills.append(_LEVEL_FILL["alert"] if dv not in (None, 0, 0.0) else None)
    if rows:
        _, _, row = _table(ws, row, [("Z", "text", 20), ("קופה", "text", 18), ("נתון", "text", 10), ("ב-Z", "money", 13),
                                     ("בחלון", "money", 13), ("מחוץ לחלון", "money", 13), ("הפרש", "money", 11), ("תוצאה", "text", 14)],
                           rows, title="Z מול העסקאות שלו (Z = בחלון + מחוץ לחלון)", fills=fills)

    ws.cell(row=row, column=1, value=f"התאמה מול שידורי אשראי — {_STATUS.get(tx.get('status'), '—')}").font = _SECTION_FONT
    row += 1
    _, _, row = _table(
        ws, row,
        [("קופה", "text", 20), ("רגלי אשראי", "int", 11), ("סכום", "money", 13), ("שודרו", "int", 9), ("סכום שודר", "money", 13),
         ("טרם שודרו", "int", 10), ("סכום טרם שודר", "money", 13), ("לא ניתנות להתאמה", "int", 12), ("ממתין לפי הקופה", "text", 22)],
        [(t["name"], t["cardLegs"]["count"], t["cardLegs"]["amount"], t["transmitted"]["count"], t["transmitted"]["amount"],
          t["untransmitted"]["count"], t["untransmitted"]["amount"], t["untracked"]["count"],
          (f"{t['tillPending']['count']} · {clock.text(t['tillPending']['reportedAt'], '%d/%m %H:%M')}"
           if t["tillPending"].get("count") is not None else "—"))
         for t in tx["tills"]],
        title="לפי קופה")
    rows, fills = [], []
    for t in tx["tills"]:
        for b in t["batches"]:
            rows.append((t["name"], b.get("batchNumber"), _STATUS.get(b["status"], b["status"]), clock.local(b["startedAt"]),
                         b.get("transactionCount"), b.get("amount"),
                         "ברמת עסקה" if b["level"] == "transaction" else "ברמת סכומים",
                         b.get("legsCompared"), b.get("amountCompared"), b["legsInWindow"], b["amountInWindow"],
                         _STATUS.get(b.get("compare") or "n/a", "—")))
            bad = b["status"] != "success" or b.get("compare") == "mismatch"
            fills.append(_LEVEL_FILL["alert"] if bad else None)
    if rows:
        _table(ws, row,
               [("קופה", "text", 18), ("אצווה", "text", 12), ("סטטוס", "text", 9), ("זמן", "time", 16), ("עסקאות (קופה)", "int", 11),
                ("סכום (קופה)", "money", 13), ("רמת השוואה", "text", 13), ("עסקאות (אצלנו)", "int", 11), ("סכום (אצלנו)", "money", 13),
                ("מהחלון", "int", 9), ("סכום מהחלון", "money", 13), ("תוצאה", "text", 12)],
               rows, title="אצוות שידור", fills=fills)


def build_workbook(report: Dict[str, Any], bucket: int = 30) -> bytes:
    clock = _Clock(report.get("timezone"))
    wb = Workbook()
    _summary(wb, report, clock)
    _insights(wb, report)
    _tills(wb, report, clock)
    _items(wb, report)
    _segments(wb, report, clock)
    _timeline(wb, report, clock, bucket)
    _shifts(wb, report, clock)
    _reconciliation(wb, report, clock)
    wb.properties.title = f"דוח אירוע - {report['event']['name']}"
    wb.properties.creator = "R2M POS"
    wb.active = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
